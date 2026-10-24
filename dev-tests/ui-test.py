# Copyright (c) 2026 IOTC (Indus Overseas Tech. Corp.). All rights reserved. GrievIQ: proprietary and confidential. Unauthorised copying, use or distribution is prohibited. See LICENSE.txt.
"""Browser tests for the Departments page (Playwright + Chromium)."""
import os, sys, re, json
from playwright.sync_api import sync_playwright

B = "http://localhost:8788"
SHOTS = os.environ.get("GRIEVIQ_TEST_DIR", "/tmp/grieviq-tests") + "/shots"
os.makedirs(SHOTS, exist_ok=True)
passed = failed = 0
def ok(c, name, extra=""):
    global passed, failed
    if c: passed += 1
    else: failed += 1; print("FAIL:", name, str(extra)[:300])

CONTRAST_JS = r"""
() => {
  function rgb(s){const m=s.match(/rgba?\(([^)]+)\)/); if(!m) return null; const p=m[1].split(',').map(Number); return {r:p[0],g:p[1],b:p[2],a:p.length>3?p[3]:1};}
  function lum(c){const f=v=>{v/=255;return v<=0.03928?v/12.92:Math.pow((v+0.055)/1.055,2.4)};return 0.2126*f(c.r)+0.7152*f(c.g)+0.0722*f(c.b);}
  function bg(el){ let cur=el; const layers=[]; while(cur && cur.nodeType===1){const c=rgb(getComputedStyle(cur).backgroundColor); if(c&&c.a>0){layers.push(c); if(c.a>=1) break;} cur=cur.parentElement;} let base={r:255,g:255,b:255}; if(!layers.length||layers[layers.length-1].a<1){base=rgb(getComputedStyle(document.querySelector('body')).backgroundColor)||base;} for(let i=layers.length-1;i>=0;i--){const c=layers[i]; base={r:c.r*c.a+base.r*(1-c.a),g:c.g*c.a+base.g*(1-c.a),b:c.b*c.a+base.b*(1-c.a)};} return base;}
  const bad=[];
  const walker=document.createTreeWalker(document.body,NodeFilter.SHOW_TEXT);
  const seen=new Set();
  while(walker.nextNode()){
    const t=walker.currentNode; if(!t.textContent.trim()) continue; const el=t.parentElement; if(!el||seen.has(el)) continue; seen.add(el);
    const cs=getComputedStyle(el); if(cs.visibility==='hidden'||cs.display==='none') continue;
    const r=el.getBoundingClientRect(); if(r.width<2||r.height<2) continue;
    if(el.closest('.sr-only,.visually-hidden,[hidden],.skip-link')) continue;
    let o=1, e=el; while(e){o*=Number(getComputedStyle(e).opacity); e=e.parentElement;}
    const fg=rgb(cs.color); const b=bg(el); const f={r:fg.r*fg.a*o+b.r*(1-fg.a*o),g:fg.g*fg.a*o+b.g*(1-fg.a*o),b:fg.b*fg.a*o+b.b*(1-fg.a*o)};
    const L1=lum(f),L2=lum(b); const ratio=(Math.max(L1,L2)+0.05)/(Math.min(L1,L2)+0.05);
    const size=parseFloat(cs.fontSize); const bold=Number(cs.fontWeight)>=700; const large=size>=24||(bold&&size>=18.66);
    const need=large?3:4.5;
    if(ratio<need) bad.push(el.tagName+'.'+el.className+' "'+t.textContent.trim().slice(0,40)+'" '+ratio.toFixed(2));
  }
  return bad;
}
"""
TARGET_JS = r"""
() => Array.from(document.querySelectorAll('main button, main a, main input, main select, main summary, header button')).filter(e=>{const r=e.getBoundingClientRect(); const cs=getComputedStyle(e); return r.width>0&&r.height>0&&cs.visibility!=='hidden' && !e.closest('dd, .meta, ol, .note, .msg, p, td')}).filter(e=>{const r=e.getBoundingClientRect(); return (r.height<24||r.width<24) && !(e.type==='checkbox'||e.type==='radio') }).map(e=>e.tagName+' '+(e.id||e.textContent.trim().slice(0,30))+' '+Math.round(e.getBoundingClientRect().width)+'x'+Math.round(e.getBoundingClientRect().height))
"""

def page_for(browser, email, width=1280, theme=None, lang=None):
    ctx = browser.new_context(viewport={"width": width, "height": 900})
    ctx.add_cookies([{"name": "test_email", "value": email, "url": B}])
    init = ""
    if theme == "light": init += "try{localStorage.setItem('griq-admin-theme','light')}catch(e){};"
    if lang: init += "try{localStorage.setItem('" + "giq-lang" + "','" + lang + "')}catch(e){};"
    if init: ctx.add_init_script(init)
    p = ctx.new_page()
    errs = []
    p.on("pageerror", lambda e: errs.append(str(e)))
    p.on("console", lambda m: errs.append(m.text) if m.type == "error" else None)
    return ctx, p, errs

def save(p, sel):
    with p.expect_response(lambda r: r.request.method == "POST"):
        p.click(sel)
    p.wait_for_timeout(150)

def shot(p, name, full=True):
    p.screenshot(path=os.path.join(SHOTS, name + ".png"), full_page=full)

def fill_otp(p):
    # The code boxes move focus as you type; fill until all six hold a digit
    # (a fast fill can land while focus is moving).
    for _ in range(5):
        for i, d in enumerate("123456"):
            box = p.locator(".otp-digit").nth(i)
            if box.input_value() != d: box.fill(d)
        if all(p.locator(".otp-digit").nth(i).input_value() == d for i, d in enumerate("123456")): return
        p.wait_for_timeout(50)

def no_hscroll(p):
    return p.evaluate("() => document.documentElement.scrollWidth <= window.innerWidth + 1")

with sync_playwright() as pw:
    br = pw.chromium.launch()
    # language storage key used by i18n.js
    ctx, p, errs = page_for(br, "super@test.in")
    p.goto(B + "/admin-departments.html?area=lucknow"); p.wait_for_selector(".office")
    store = p.evaluate("() => Object.keys(localStorage)")
    ctx.close()
    LANGKEY = None
    src = open(os.path.join(sys.argv[1], "public/i18n.js"), encoding="utf-8").read()
    m = re.search(r"STORE\s*=\s*['\"]([^'\"]+)['\"]", src); LANGKEY = m.group(1) if m else "giq-lang"

    def P(email, width=1280, theme=None, lang=None):
        ctx = br.new_context(viewport={"width": width, "height": 900})
        ctx.add_cookies([{"name": "test_email", "value": email, "url": B}])
        init = ""
        if theme == "light": init += "try{localStorage.setItem('griq-admin-theme','light')}catch(e){};"
        if lang: init += "try{localStorage.setItem('%s','%s')}catch(e){};" % (LANGKEY, lang)
        if init: ctx.add_init_script(init)
        p = ctx.new_page(); errs = []
        p.on("pageerror", lambda e: errs.append(str(e)))
        p.on("console", lambda m: errs.append(m.text) if m.type == "error" and "favicon" not in m.text and "ERR_TUNNEL_CONNECTION_FAILED" not in m.text and "status of 400" not in m.text and "status of 409" not in m.text else None)
        p.on("dialog", lambda d: d.accept())
        return ctx, p, errs

    # ---------- super admin: whole flow ----------
    ctx, p, errs = P("super@test.in")
    p.goto(B + "/admin-departments.html?area=lucknow"); p.wait_for_selector(".office")
    ok("Departments" in p.title(), "page title")
    p.click('button[data-key="adm.nav_g_data"]')
    ok(p.is_visible('#adm-drop-data a[href="/admin-departments.html"]'), "menu: Departments under Areas & data")
    links = p.eval_on_selector_all("#adm-drop-data a", "els => els.map(e => e.textContent.trim())")
    ok(links[:2] == ["Areas", "Departments"], "Departments right after Areas", links)
    p.keyboard.press("Escape")
    shot(p, "01-list-dark")
    ok(p.locator(".card .num").first.inner_text() in ("6", "7"), "offices card")
    # filters
    p.select_option("#f-dept", "Electricity")
    ok(p.locator(".office").count() == 1 and "MVVNL" in p.locator(".office h3").first.inner_text(), "filter by type")
    p.select_option("#f-dept", "")
    p.check("#f-stale"); ok(p.locator(".office").count() == 1, "needs-checking filter"); p.uncheck("#f-stale")
    p.fill("#f-q", "cmo"); ok(p.locator(".office").count() == 1, "search"); ok(p.evaluate("document.activeElement.id") == "f-q", "focus stays in search"); p.fill("#f-q", "")
    p.check("#f-retired"); ok(p.locator(".office.retired").count() >= 1 and p.locator(".office:not(.retired)").count() == 0, "retired list"); p.uncheck("#f-retired")
    # coverage
    p.click('a[href="#coverage"]')
    ok(p.locator(".cov td.gap").count() > 0, "coverage shows gaps")
    p.check("#cov-gaps")
    shot(p, "02-coverage", full=True)
    p.uncheck("#cov-gaps")

    # add office: empty submit -> error summary
    p.click("#add-office"); p.wait_for_selector("#office-form")
    ok(p.evaluate("document.activeElement.id") == "form-h", "focus moves to form heading")
    p.fill("#o-source", ""); p.click("#o-save")
    p.wait_for_selector("#o-summary")
    ok(p.evaluate("document.activeElement.id") == "o-summary", "focus on error summary")
    items = p.locator("#o-summary li").all_inner_texts()
    ok(len(items) >= 4, "summary lists problems", items)
    shot(p, "03-form-errors")
    p.click("#o-summary li >> nth=0 >> a")
    ok(p.evaluate("document.activeElement.id") == "o-name-en", "summary link goes to field")
    ok(p.get_attribute("#o-name-en", "aria-invalid") == "true" and "o-name-en-err" in p.get_attribute("#o-name-en", "aria-describedby"), "field marked invalid + described")
    # fill properly, some wards
    p.fill("#o-name-en", "Jal Kal Zone 3 Office"); p.fill("#o-name-hi", "जलकल ज़ोन 3 कार्यालय")
    p.check('[data-dept="Water Supply"]')
    p.check('input[name="o-wards"][value="LIST"]')
    p.fill("#o-ward-q", "nagar")
    ok(p.locator("[data-ward]").count() == 2, "ward search filters list")
    p.click("#o-ward-all")
    ok("2 of 5" in p.inner_text("#o-ward-count"), "select shown", p.inner_text("#o-ward-count"))
    p.fill("#o-ward-q", "")
    p.fill("#o-phone", "0522 2612345"); p.fill("#o-whatsapp", "12345")
    p.fill("#o-source", "RTI reply from Jal Kal, 3 Oct 2026")
    save(p, "#o-save")
    ok("10-digit" in p.inner_text("#o-summary"), "server error for WhatsApp shown", p.inner_text("#o-summary"))
    ok(p.input_value("#o-name-hi") == "जलकल ज़ोन 3 कार्यालय", "typed values kept after error")
    p.fill("#o-whatsapp", "")
    p.click("details.officer summary"); p.fill("#o-off-phone", "9876543210"); p.click("#o-save"); p.wait_for_timeout(150)
    ok(p.locator("#o-summary li").count() == 3, "officer number needs name, date, how")
    p.fill("#o-off-phone", "")
    save(p, "#o-save")
    ok("was added" in p.inner_text("#top-msg") and p.evaluate("document.activeElement.id") == "top-msg", "added message + focus")
    ok(p.locator('.office h3:has-text("Jal Kal Zone 3 Office")').count() == 1, "new office listed")
    card = p.locator('.office:has(h3:has-text("Jal Kal Zone 3 Office"))')
    ok("Gomti Nagar" in card.inner_text() and "Tilak Nagar" in card.inner_text(), "coverage text on card")
    # edit
    card.locator("[data-edit]").click(); p.wait_for_selector("#office-form")
    ok(p.input_value("#o-name-en") == "Jal Kal Zone 3 Office" and p.is_checked('[data-dept="Water Supply"]') and p.is_checked('input[name="o-wards"][value="LIST"]'), "edit form prefilled")
    save(p, "#o-save")
    ok("Nothing has changed" in p.inner_text("#top-msg"), "unchanged edit explained")
    p.fill("#o-hours", "Mon–Sat, 10 am – 5 pm"); save(p, "#o-save")
    ok("saved" in p.inner_text("#top-msg"), "edit saved")
    # cancel returns focus
    card = p.locator('.office:has(h3:has-text("Jal Kal Zone 3 Office"))')
    card.locator("[data-edit]").click(); p.click("#o-cancel")
    ok("Jal Kal Zone 3" in (p.evaluate("document.activeElement.textContent") or ""), "cancel returns focus to Edit")
    # retire
    card = p.locator('.office:has(h3:has-text("Jal Kal Zone 3 Office"))')
    card.locator("[data-retire]").click()
    ok(p.evaluate("document.activeElement.id") == "r-reason", "focus in retire reason")
    p.fill("#r-reason", "short"); p.click("#r-go")
    ok(p.is_visible("#r-reason-err"), "short reason error")
    p.fill("#r-reason", "Office moved into the main Jal Kal building"); save(p, "#r-go")
    ok("retired" in p.inner_text("#top-msg"), "retired")
    p.check("#f-retired")
    rc = p.locator('.office:has(h3:has-text("Jal Kal Zone 3 Office"))')
    ok("main Jal Kal building" in rc.inner_text(), "reason shown on retired card")
    save(p, '.office:has(h3:has-text("Jal Kal Zone 3 Office")) [data-restore]')
    ok("back in use" in p.inner_text("#top-msg") and not p.is_checked("#f-retired"), "brought back")

    # upload
    p.click("#open-upload"); p.wait_for_selector("#upload-panel")
    ok(p.is_disabled("#up-check"), "check disabled until a file is chosen")
    csvp = os.path.join(SHOTS, "test-upload.csv")
    open(csvp, "w", encoding="utf-8").write("office_name_en,office_name_hi,handles,wards,helpline,office_phone,office_email,whatsapp,website,address,hours,source,last_checked,notes\n"
        "Health Helpline 104,,Health,ALL,104,,,,,,,https://up.gov.in,2026-10-01,\n"
        "Broken,,Water,Atlantis,,,,,,,,,,\n")
    p.set_input_files("#up-file", csvp); p.click("#up-check"); p.wait_for_selector("#up-sum")
    ok("1 of 2 rows need fixing" in p.inner_text("#up-sum"), "file report", p.inner_text("#up-sum"))
    ok("Atlantis" in p.inner_text(".report") and p.locator("#up-save").count() == 0, "bad row explained, no save button")
    shot(p, "04-upload-report")
    open(csvp, "w", encoding="utf-8").write("office_name_en,office_name_hi,handles,wards,helpline,office_phone,office_email,whatsapp,website,address,hours,source,last_checked,notes\n"
        "Health Helpline 104,,Health,ALL,104,,,,,,,https://up.gov.in,2026-10-01,\n")
    p.set_input_files("#up-file", csvp); p.click("#up-check"); p.wait_for_selector("#up-save")
    save(p, "#up-save")
    ok("1 offices were added" in p.inner_text("#top-msg") or "were added" in p.inner_text("#top-msg"), "file saved", p.inner_text("#top-msg"))
    # template link
    with p.expect_download() as d:
        p.click("#open-upload"); p.click("text=Download template (CSV)")
    ok(d.value.suggested_filename == "grieviq-departments-template.csv", "template downloads")
    p.click("#up-close")
    # area switch
    p.select_option("#f-area", "kanpur"); p.wait_for_function("document.querySelector('#f-area').value==='kanpur' && document.querySelectorAll('.office').length===1")
    ok("area=kanpur" in p.url, "area in address")
    ok(not errs, "no script errors (super)", errs)
    ctx.close()


    # ---------- grieviq-25: area-specific examples, department types ----------
    ctx, p, errs = P("super@test.in")
    p.goto(B + "/admin-departments.html?area=kanpur"); p.wait_for_selector("#f-area")
    ok(p.locator("#f-area option:checked").inner_text() == "Kanpur", "area name shown with a capital letter")
    p.click("#add-office"); p.wait_for_selector("#office-form")
    ok('"Jal Kal Vibhag, Kanpur"' in p.inner_text("#o-name-en-hint"), "English example names Kanpur", p.inner_text("#o-name-en-hint"))
    ok("कानपुर" in p.inner_text("#o-name-hi-hint"), "Hindi example names कानपुर")
    ok("0512" in p.inner_text("#o-phone-hint"), "phone example uses Kanpur's STD code")
    ok("Whole of Kanpur" in p.inner_text("#office-form"), "Whole of Kanpur")
    ok("add it here" in p.inner_text("#o-dept-fs"), "note under Other")
    # "+ Add a new type" inside the form: typed values kept, new type ticked
    p.fill("#o-name-en", "KESCO Street Light Cell"); p.check('[data-dept="Electricity"]')
    p.click("#nt-open"); ok(p.evaluate("document.activeElement.id") == "nt-h", "new-type box opens, focus moves")
    p.click("#nt-save"); p.wait_for_timeout(150)
    ok(p.is_visible("#nt-en-err") and p.is_visible("#nt-hi-err"), "new-type names required")
    p.fill("#nt-en", "Public Toilets"); p.fill("#nt-hi", "सार्वजनिक शौचालय विभाग"); p.fill("#nt-desc", "Dirty or broken public toilets")
    shot(p, "12-inline-type")
    save(p, "#nt-save"); p.wait_for_timeout(400)
    ok(p.is_checked('[data-dept="Public Toilets"]') and p.is_checked('[data-dept="Electricity"]'), "new type ticked, earlier tick kept")
    ok(p.input_value("#o-name-en") == "KESCO Street Light Cell", "typed office name kept")
    ok(p.evaluate("document.activeElement.dataset.dept") == "Public Toilets", "focus on the new tick")
    p.fill("#o-helpline", "1912"); p.fill("#o-source", "https://kesco.co.in")
    save(p, "#o-save")
    ok("was added" in p.inner_text("#top-msg") and p.locator('.office:has-text("KESCO Street Light Cell") .tag:has-text("Public Toilets")').count() == 1, "office saved with the new type")
    p.click("#add-office"); p.wait_for_selector("#office-form")
    p.click("#o-cancel")
    p.goto(B + "/admin-departments.html?area=lucknow"); p.wait_for_selector(".office")
    p.click('a[href="#types"]')
    ok(p.locator("table.types tbody tr").count() >= 8, "types table")
    other_row = p.locator('table.types tr:has(th:has-text("Other"))').first
    ok(other_row.locator("[data-tretire]").count() == 0, "Other has no Retire button")
    p.click("#t-add"); p.wait_for_selector("#type-form")
    ok(p.evaluate("document.activeElement.id") == "type-form-h", "focus on type form heading")
    p.click("#t-save"); p.wait_for_selector("#t-summary")
    ok(p.locator("#t-summary li").count() == 2 and p.evaluate("document.activeElement.id") == "t-summary", "type form errors summarised")
    p.fill("#t-name-en", "Pollution / Noise"); p.fill("#t-name-hi", "Pollution")
    save(p, "#t-save")
    ok("Hindi (Devanagari)" in p.inner_text("#t-summary"), "server says Hindi name must be Hindi")
    p.fill("#t-name-hi", "प्रदूषण / शोर विभाग"); p.fill("#t-desc", "Air, water and noise pollution")
    shot(p, "10-type-form")
    save(p, "#t-save")
    ok("was added" in p.inner_text("#top-msg"), "type added message")
    row = p.locator('table.types tr:has(th:has-text("Pollution / Noise"))')
    ok(row.count() == 1 and "Not used yet" in row.inner_text(), "new type listed, unused")
    ok(p.locator('.cov th:has-text("Pollution / Noise")').count() == 1, "new coverage column")
    # it can be ticked on an office
    p.click("#add-office"); p.wait_for_selector("#office-form")
    ok(p.locator('[data-dept="Pollution / Noise"]').count() == 1, "new type tickable on offices")
    p.click("#o-cancel")
    # in use -> retire blocked with names
    row = p.locator('table.types tr:has(th:has-text("Water Supply"))')
    row.locator("[data-tretire]").click(); p.wait_for_selector("#tr-msg")
    ok("Jal Kal" in p.inner_text("#tr-msg") and "Issue type" in p.inner_text("#tr-msg"), "in-use type: users named", p.inner_text("#tr-msg"))
    p.click("#tr-cancel")
    # rename then retire the unused new type
    row = p.locator('table.types tr:has(th:has-text("Pollution / Noise"))')
    row.locator("[data-tedit]").click(); p.wait_for_selector("#type-form")
    ok(p.input_value("#t-name-hi") == "प्रदूषण / शोर विभाग", "edit form prefilled")
    p.fill("#t-name-en", "Pollution and Noise"); save(p, "#t-save")
    ok("were saved" in p.inner_text("#top-msg"), "type renamed")
    row = p.locator('table.types tr:has(th:has-text("Pollution and Noise"))')
    row.locator("[data-tretire]").click()
    p.fill("#tr-reason", "short"); p.click("#tr-go"); ok(p.is_visible("#tr-reason-err"), "short reason refused")
    p.fill("#tr-reason", "Handled by the Pollution Control Board, not here"); save(p, "#tr-go")
    ok("was retired" in p.inner_text("#top-msg"), "type retired")
    row = p.locator('table.types tr:has(th:has-text("Pollution and Noise"))')
    ok("Retired" in row.inner_text() and p.locator('.cov th:has-text("Pollution")').count() == 0, "retired: marked, no coverage column")
    save(p, 'table.types tr:has(th:has-text("Pollution and Noise")) [data-trestore]')
    ok("back in use" in p.inner_text("#top-msg"), "type brought back")
    # Hindi view of the types table
    p.click("[data-lang-toggle]"); p.wait_for_timeout(300)
    ok(p.locator('table.types th:has-text("प्रदूषण / शोर विभाग")').count() == 1, "Hindi type name shown in Hindi")
    shot(p, "11-types-hi", full=True)
    p.click("[data-lang-toggle]")
    ok(not errs, "no script errors (types)", errs)
    ctx.close()

    # Issue types page lists the new type, in Hindi too
    ctx, p, errs = P("super@test.in", lang="hi")
    p.goto(B + "/admin-issue-types.html"); p.wait_for_selector(".dept-select")
    opts = p.eval_on_selector_all(".dept-select >> nth=0 >> option", "els => els.map(e => e.textContent)")
    ok("प्रदूषण / शोर विभाग" in opts, "Issue types: new type in Hindi", opts)
    ok(not errs, "no script errors (issue types)", errs)
    ctx.close()

    # Rep console and Track page show the new type (functions checked in the page itself)
    for lang in ("en", "hi"):
        ctx, p, errs = P("nobody@test.in", lang=lang)
        p.route("**/api/**", lambda route: route.fulfill(status=401, body='{"error":"SIGNED_OUT"}', headers={"content-type": "application/json"}))
        p.goto(B + "/rep.html"); p.wait_for_timeout(800)
        html = p.evaluate("""() => { takeDepartments({ departments: ['Water Supply','Pollution / Noise','Other'], deptNames: { 'Pollution / Noise': { en: 'Pollution and Noise', hi: 'प्रदूषण / शोर विभाग' } } });
          return renderFollowupControl({ id: 'c1', status: 'OPEN', suggestedDepartment: 'Pollution / Noise', currentDepartment: 'Pollution / Noise', followupHistory: [] }); }""")
        want = "प्रदूषण / शोर विभाग" if lang == "hi" else "Pollution and Noise"
        seg = html[html.index('class="followup-dept"'):]
        first_opt = re.findall(r'<option value="([^"]*)">([^<]*)</option>', seg[:seg.index('</select>')])
        ok(first_opt[1][0] == "Pollution / Noise" and want in first_opt[1][1], "rep: suggested new type listed first, named " + lang, first_opt[:3])
        ok(len(first_opt) == 4, "rep: only the current list offered " + lang)
        ctx.close()
        ctx, p, errs = P("nobody@test.in", lang=lang)
        p.goto(B + "/status.html"); p.wait_for_timeout(500)
        name = p.evaluate("() => { GIQ.addDepts({ 'Pollution / Noise': { en: 'Pollution and Noise', hi: 'प्रदूषण / शोर विभाग' } }); return GIQ.dept('Pollution / Noise') + '|' + GIQ.dept('Water Supply'); }")
        ok(name == (want + "|" + ("जलापूर्ति विभाग" if lang == "hi" else "Water Supply")), "Track: department names " + lang, name)
        ctx.close()


    # ---------- Departments stage 2: "Who handles this" card in the rep console ----------
    DIR = {"units": {"u1": "lucknow", "u2": "lucknow"}, "areas": {"lucknow": {"name": "Lucknow", "state": "Uttar Pradesh"}},
           "offices": [
             {"id": "o1", "areaId": "lucknow", "nameEn": "Lucknow Nagar Nigam Control Room", "nameHi": "लखनऊ नगर निगम कंट्रोल रूम", "departments": ["Water Supply", "Sanitation / Garbage"], "wards": None,
              "helpline": "1533", "officePhone": None, "officeEmail": "nnlko@nic.in", "whatsapp": "9219902911", "website": "https://lmc.up.nic.in/helpline.aspx", "address": None, "hours": "24 hours", "officerName": None, "officerPhone": None, "lastChecked": "2026-10-07"},
             {"id": "o2", "areaId": "lucknow", "nameEn": "Jal Kal Zone 3", "nameHi": None, "departments": ["Water Supply"], "wards": ["u1"],
              "helpline": None, "officePhone": "0522 2612345", "officeEmail": None, "whatsapp": None, "website": None, "address": "Aishbagh", "hours": None, "officerName": "R. K. Singh", "officerPhone": "9876543210", "lastChecked": "2026-10-03"}]}
    CASE = {"id": "c1", "status": "OPEN", "suggestedDepartment": "Water Supply", "localUnit": {"id": "u1"}, "followupHistory": [], "myRole": "FIELD_WORKER"}
    told = []
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                ctx = br.new_context(viewport={"width": width, "height": 900})
                ctx.add_init_script("try{localStorage.setItem('griq-theme','%s');localStorage.setItem('%s','%s')}catch(e){}" % (theme, LANGKEY, lang))
                p = ctx.new_page(); errs = []
                p.on("pageerror", lambda e: errs.append(str(e)))
                def route(rt):
                    if "/api/dept-gap" in rt.request.url:
                        told.append(rt.request.post_data); rt.fulfill(status=200, body='{"ok":true}', headers={"content-type": "application/json"})
                    else:
                        rt.fulfill(status=401, body='{"error":"SIGNED_OUT"}', headers={"content-type": "application/json"})
                p.route("**/api/**", route)
                p.goto(B + "/rep.html"); p.wait_for_timeout(600)
                p.evaluate("""([dir, c]) => { takeDepartments({ departments: ['Water Supply','Electricity','Sanitation / Garbage','Health','Other'], deptDirectory: dir });
                  currentCases = [c]; const m = document.querySelector('main') || document.body;
                  const box = document.createElement('div'); box.id = 'dc-test'; box.className = 'case-card'; box.style.maxWidth = '640px'; box.style.padding = '12px';
                  box.innerHTML = renderFollowupControl(c); m.prepend(box); window.scrollTo(0, 0); }""", [DIR, CASE])
                tag = "%s-%s-%d" % (theme, lang, width)
                card = p.locator("#dc-c1")
                ok(card.count() == 1, "card shown to a field worker " + tag)
                names = p.locator("#dc-c1 .dc-name").all_inner_texts()
                ok(names[0] == "Jal Kal Zone 3" and len(names) == 2, "ward office first, then whole-area office " + tag, names)
                if lang == "hi":
                    ok("लखनऊ नगर निगम कंट्रोल रूम" in names[1] and "यह कौन देखता है" in card.inner_text(), "card in Hindi " + tag)
                else:
                    hrefs = p.eval_on_selector_all("#dc-c1 .dc-link", "els => els.map(e => e.getAttribute('href'))")
                    ok("tel:05222612345" in hrefs and "tel:9876543210" in hrefs and "tel:1533" in hrefs and "mailto:nnlko@nic.in" in hrefs and "https://wa.me/919219902911" in hrefs, "tap-to-call, email, WhatsApp links " + tag, hrefs)
                    ok("Call helpline 1533" in card.inner_text() and "Jansunwai" not in card.inner_text(), "link text says what it does; no Jansunwai line " + tag)
                ok(no_hscroll(p), "card: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS.replace("document.body", "document.getElementById('dc-test')", 1))
                ok(not bad, "card contrast AA " + tag, bad[:5])
                small = p.evaluate("() => Array.from(document.querySelectorAll('#dc-c1 a.dc-link, #dc-c1 select, #dc-c1 button')).filter(e => { const r = e.getBoundingClientRect(); return r.height < 24; }).length")
                ok(small == 0, "card targets >= 24px " + tag)
                shot(p, "13-card-" + tag, full=False)
                # Not the right department? -> Health: no contact -> Tell GrievIQ
                p.select_option("#dc-sel-c1", "Health")
                ok(p.evaluate("document.activeElement.id") == "dc-sel-c1", "focus stays on the department list " + tag)
                ok(p.locator("#dc-c1 .dc-tell").count() == 1, "no contact: Tell GrievIQ button " + tag)
                p.click("#dc-c1 .dc-tell"); p.wait_for_selector("#dc-c1 .dc-told")
                ok(p.evaluate("document.activeElement.className") == "dc-told", "thank-you message focused " + tag)
                ok(not errs, "no script errors (card) " + tag, errs)
                ctx.close()
    ok(len(told) == 8 and '"department":"Health"' in told[0] and '"grievanceId":"c1"' in told[0], "Tell GrievIQ sends case and department", told[:1])


    # ---------- Departments stage 3: department progress, reply, forward, field check (rep console) ----------
    import json as _json
    def steps_case(kind_list, phase, role="REPRESENTATIVE", extra=None):
        st = [{"id": "s%d" % i, "kind": k, "department": "Water Supply", "officeName": "Jal Kal Zone 3", "channel": "PHONE" if k == "FORWARDED" else None,
               "expectedDate": "2026-10-12" if k == "SCHEDULED" else None, "suggestedDepartment": None, "note": None, "actor": "r1@x.in", "actorRole": "REPRESENTATIVE",
               "at": "2026-10-0%dT06:30:00Z" % (i + 1), "photos": [{"url": "/logo-mark.svg", "thumbUrl": "/logo-mark.svg"}] if k.startswith("CHECK_N") else None} for i, k in enumerate(kind_list)]
        c = {"id": "c9", "status": "ACKNOWLEDGED", "suggestedDepartment": "Water Supply", "localUnit": {"id": "u1", "name": "Aishbagh"}, "followupHistory": [], "myRole": role,
             "deptSteps": st, "deptState": dict({"phase": phase, "department": "Water Supply", "officeName": "Jal Kal Zone 3", "days": 7, "dayOf": 3, "overdue": False}, **(extra or {}))}
        return c
    posts = []
    def rep_page(lang="en", theme="light", width=1280):
        ctx = br.new_context(viewport={"width": width, "height": 900})
        ctx.add_init_script("try{localStorage.setItem('griq-theme','%s');localStorage.setItem('%s','%s')}catch(e){}" % (theme, LANGKEY, lang))
        p = ctx.new_page(); errs = []
        p.on("pageerror", lambda e: errs.append(str(e)))
        p.on("dialog", lambda d: d.accept())
        def route(rt):
            u = rt.request.url
            if "/dept-step" in u or "/add-followup" in u:
                posts.append((u.split("/api/")[1], _json.loads(rt.request.post_data or "{}")))
                rt.fulfill(status=200, body='{"ok":true}', headers={"content-type": "application/json"})
            else:
                rt.fulfill(status=401, body='{"error":"SIGNED_OUT"}', headers={"content-type": "application/json"})
        p.route("**/api/**", route)
        p.goto(B + "/rep.html"); p.wait_for_timeout(500)
        # Show the content area and stop the real reload after saving.
        p.evaluate("""() => { window.init = async () => {}; activeTab = 'cases';
          if (!document.getElementById('content')) { const d = document.createElement('div'); d.id = 'content'; document.body.prepend(d); }
          const m = document.getElementById('content'); m.hidden = false; m.style.display = 'block';
          let p = m.parentElement; while (p && p !== document.body) { p.hidden = false; p.style.display = ''; p = p.parentElement; } }""")
        return ctx, p, errs
    def show(p, case):
        p.evaluate("""([dir, c]) => { takeDepartments({ departments: ['Water Supply','Electricity','Health','Other'], deptDirectory: dir }); currentCases = [c];
          const box = document.getElementById('content'); box.innerHTML = '<div id="ds-test" style="max-width:640px;padding:12px">' + renderFollowupControl(c) + renderResolveControl(c) + '<p id="resolve-status" role="status"></p></div>';
          renderContent = () => { const cc = currentCases[0]; document.getElementById('ds-test').innerHTML = renderFollowupControl(cc) + renderResolveControl(cc) + '<p id="resolve-status" role="status"></p>'; attachFollowupHandlers(); attachResolveHandlers(); };
          attachFollowupHandlers(); attachResolveHandlers(); window.scrollTo(0, 0); }""", [DIR, case])
    # With the department, day 3 of 7; reply form
    ctx, p, errs = rep_page()
    show(p, steps_case(["FORWARDED", "SCHEDULED"], "WITH_DEPT", extra={"expectedDate": "2026-10-12"}))
    ok("With Jal Kal Zone 3: day 3 of 7." in p.inner_text("#ds-c9 .ds-status"), "status: day 3 of 7", p.inner_text("#ds-c9"))
    ok("Forwarded to Jal Kal Zone 3 (Phone call)" in p.inner_text("#ds-c9") and "scheduled the work for 12 Oct 2026" in p.inner_text("#ds-c9"), "steps listed")
    p.click(".ds-open"); ok(p.evaluate("document.activeElement.name") == "ds-kind-c9", "reply form opens, focus on the first choice")
    p.click(".ds-save"); ok("Choose what the department said." in p.inner_text("#ds-form-c9"), "reply: choose one")
    p.check('input.ds-kind[value="SCHEDULED"]'); ok(p.is_visible("#ds-date-c9"), "date field for scheduled")
    p.check('input.ds-kind[value="CANT_DO"]'); p.click(".ds-save")
    ok("reason the department gave" in p.inner_text("#ds-form-c9") and p.get_attribute("#ds-note-c9", "aria-invalid") == "true", "can't do needs the reason")
    p.fill("#ds-note-c9", "No budget for new pipeline this year"); p.click(".ds-save"); p.wait_for_timeout(300)
    ok(posts and posts[-1][0] == "grievances/c9/dept-step" and posts[-1][1]["kind"] == "CANT_DO" and "budget" in posts[-1][1]["note"], "reply sent", posts[-1:])
    shot(p, "14-ds-with-dept", full=False)
    # Forward form: office list follows the department, "not in the list" asks for a name, how contacted required
    p.select_option("#followup-dept-c9", "Water Supply")
    opts = p.eval_on_selector_all("#followup-office-c9 option", "els => els.map(e => e.textContent)")
    ok("Jal Kal Zone 3" in opts and "Office not in the list" in opts, "office list for the department", opts)
    p.select_option("#followup-office-c9", "__other"); ok(p.is_visible("#followup-oname-c9"), "office name box for an office not listed")
    p.click(".followup-btn"); p.wait_for_timeout(100)
    ok("Choose how you contacted the office." in p.inner_text("#fwd-form-c9") and "office's name" in p.inner_text("#fwd-form-c9"), "forward: how + name required")
    p.fill("#followup-oname-c9", "Jal Kal sub-division 2"); p.check('input.followup-ch[value="WHATSAPP"]'); p.click(".followup-btn"); p.wait_for_timeout(300)
    ok(posts[-1][0] == "grievances/c9/add-followup" and posts[-1][1]["channel"] == "WHATSAPP" and posts[-1][1]["officeName"] == "Jal Kal sub-division 2", "forward sent with office and channel", posts[-1:])
    ok(not errs, "no script errors (department progress)", errs)
    ctx.close()
    # Overdue and not-ours wording
    ctx, p, errs = rep_page()
    show(p, steps_case(["FORWARDED"], "WITH_DEPT", extra={"overdue": True, "overdueDays": 2}))
    ok("past its target by 2 day(s)" in p.inner_text("#ds-c9 .ds-status.late"), "overdue shown")
    show(p, steps_case(["FORWARDED", "NOT_OURS"], "NEEDS_FORWARD", extra={"suggestedDepartment": "Electricity"}))
    ok("They suggested Electricity" in p.inner_text("#ds-c9") and p.locator(".ds-open").count() == 0, "not ours: forward again, no reply button")
    ctx.close()
    # Field check (all themes/languages/widths): outcome, photos required, send back
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx, p, errs = rep_page(lang, theme, width)
                show(p, steps_case(["FORWARDED", "DONE_CLAIMED"], "NEEDS_CHECK", role="FIELD_WORKER"))
                btn = p.locator(".mark-resolved-btn")
                ok(btn.inner_text() == ("Record field check" if lang == "en" else "मौके पर जाँच दर्ज करें"), "field check button " + tag)
                btn.click(); p.wait_for_timeout(100)
                ok(p.locator(".r-outcome").count() == 3 and p.locator(".r-nophoto").count() == 0, "outcome choices, no 'no photo' option " + tag)
                p.fill("#rnote-c9", "Pipe still leaking at the market gate"); p.click(".r-submit"); p.wait_for_timeout(100)
                ok(p.locator("#routcome-err-c9").inner_text() != "", "choose what you found " + tag)
                p.check('input.r-outcome[value="NOT_FIXED"]'); p.wait_for_timeout(100)
                ok(p.inner_text(".r-submit") == ("Send back to the department" if lang == "en" else "विभाग को वापस भेजें"), "button says send back " + tag)
                p.click(".r-submit"); p.wait_for_timeout(100)
                ok(p.get_attribute("#rphoto-c9", "aria-invalid") == "true", "photo required " + tag)
                p.evaluate("() => { drafts['c9'].photos = [{ id: 'ph9', url: '/logo-mark.svg', thumbUrl: '/logo-mark.svg', warnings: [] }]; renderContent(); }")
                p.click(".r-submit"); p.wait_for_timeout(300)
                ok(posts[-1][0] == "grievances/c9/dept-step" and posts[-1][1]["kind"] == "CHECK_NOT_FIXED" and posts[-1][1]["photoIds"] == ["ph9"], "failed check sent " + tag, posts[-1:])
                show(p, steps_case(["FORWARDED", "DONE_CLAIMED", "CHECK_NOT_FIXED"], "WITH_DEPT", extra={"sentBack": True}))
                ok(no_hscroll(p), "no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS.replace("document.body", "document.getElementById('ds-test')", 1))
                ok(not bad, "contrast AA " + tag, bad[:5])
                shot(p, "15-ds-" + tag, full=False)
                ok(not errs, "no script errors (field check) " + tag, errs)
                ctx.close()

    # Track page: the department's steps for the citizen
    for lang in ("en", "hi"):
        ctx = br.new_context(viewport={"width": 390, "height": 900})
        ctx.add_init_script("try{localStorage.setItem('%s','%s')}catch(e){}" % (LANGKEY, lang))
        p = ctx.new_page(); errs = []
        p.on("pageerror", lambda e: errs.append(str(e)))
        case = {"trackingRef": "GRV-TEST02", "description": "Leaking pipe", "localUnitName": "Aishbagh", "createdAt": "2026-10-08T10:00:00Z", "status": "ACKNOWLEDGED",
                "tiers": [{"label": "Corporator", "visible": True, "current": True, "slaBreached": False}], "currentTierIndex": 0, "reopenStatus": {"can": False},
                "currentDepartment": "Water Supply", "followupHistory": [], "deptNames": {},
                "deptSteps": [{"kind": "FORWARDED", "department": "Water Supply", "officeName": "Jal Kal Zone 3", "at": "2026-10-08T11:00:00Z"},
                              {"kind": "DONE_CLAIMED", "department": "Water Supply", "officeName": "Jal Kal Zone 3", "at": "2026-10-09T11:00:00Z"},
                              {"kind": "CHECK_NOT_FIXED", "department": "Water Supply", "officeName": "Jal Kal Zone 3", "at": "2026-10-09T15:00:00Z"}]}
        def make_troute(case):
          def troute(rt):
            u = rt.request.url
            if u.endswith("/api/otp/request"): rt.fulfill(status=200, body='{"ok":true}', headers={"content-type": "application/json"})
            elif u.endswith("/api/otp/verify"): rt.fulfill(status=200, body=_json.dumps({"reports": [{"trackingRef": "GRV-TEST02", "status": "ACKNOWLEDGED"}], "case": case}), headers={"content-type": "application/json"})
            else: rt.fulfill(status=404, body='{}', headers={"content-type": "application/json"})
          return troute
        p.route("**/api/**", make_troute(case))
        p.goto(B + "/status.html?ref=GRV-TEST02"); p.wait_for_timeout(400)
        p.fill("#email", "c@x.in"); p.click("#entry-submit"); p.wait_for_selector(".otp-digit")
        fill_otp(p)
        p.click("#otp-submit"); p.wait_for_selector(".ds-list")
        txt = p.inner_text("#status-followup")
        if lang == "en":
            ok("forwarded it to Jal Kal Zone 3" in txt and "reported the work done" in txt and "not fixed yet. Sent back to Jal Kal Zone 3" in txt, "Track: steps in plain words", txt)
        else:
            ok("Jal Kal Zone 3 को भेजा" in txt and "वापस भेजा गया" in txt, "Track: steps in Hindi", txt)
        ok("r1@x.in" not in txt, "Track: no staff names")
        ok(no_hscroll(p), "Track: no sideways scroll " + lang)
        bad = p.evaluate(CONTRAST_JS.replace("document.body", "document.getElementById('status-followup')", 1))
        ok(not bad, "Track contrast " + lang, bad[:4])
        p.locator("#status-followup").screenshot(path=os.path.join(SHOTS, "16-track-" + lang + ".png"))
        ok(not errs, "no script errors (Track) " + lang, errs)
        ctx.close()

    # ---------- Citizen ratings (grieviq-30) ----------
    import json as _json2
    # Track page: the rating card (mocked server), every theme/language, phone and laptop
    def rated_case(rating=None, status=None):
        return {"trackingRef": "GRV-TEST02", "description": "Leaking pipe", "localUnitName": "Aishbagh", "createdAt": "2026-10-08T10:00:00Z", "status": "RESOLVED",
                "resolutionKind": "CONFIRMED", "tiers": [{"label": "Corporator", "visible": True, "current": True, "slaBreached": False}], "currentTierIndex": 0,
                "reopenStatus": {"can": True, "code": "OK", "until": "2026-11-08T10:00:00Z"}, "followupHistory": [], "deptNames": {"Electricity": {"en": "Electricity", "hi": "विद्युत विभाग"}}, "deptSteps": [],
                "rating": rating, "ratingStatus": status or {"can": True, "code": "OK", "until": "2026-11-08T10:00:00Z", "mode": "NEW"},
                "ratingDept": {"key": "Electricity", "officeName": "MVVNL sub-station Hazratganj"}}
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            for width in (375, 1280):
                if width == 1280 and (theme, lang) != ("light", "en"): continue
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx = br.new_context(viewport={"width": width, "height": 900}, color_scheme=theme)
                ctx.add_init_script("try{localStorage.setItem('%s','%s')}catch(e){}" % (LANGKEY, lang))
                p = ctx.new_page(); errs = []; sent = []
                p.on("pageerror", lambda e: errs.append(str(e)))
                case = rated_case()
                def make_r(case, sent):
                    def r(rt):
                        u = rt.request.url
                        if u.endswith("/api/otp/request"): rt.fulfill(status=200, body='{"ok":true}', headers={"content-type": "application/json"})
                        elif u.endswith("/api/otp/verify"): rt.fulfill(status=200, body=_json2.dumps({"reports": [{"trackingRef": "GRV-TEST02", "status": "RESOLVED"}], "case": case}), headers={"content-type": "application/json"})
                        elif u.endswith("/api/grievances/rate"):
                            b = _json2.loads(rt.request.post_data or "{}"); sent.append(b)
                            rt.fulfill(status=200, body=_json2.dumps({"ok": True, "rating": {"office": b["office"], "dept": b.get("dept"), "department": "Electricity", "comment": b.get("comment") or "", "submittedAt": "2026-10-09T10:00:00Z", "updatedAt": "2026-10-09T10:00:00Z"},
                                "ratingStatus": {"can": True, "code": "OK", "until": "2026-10-16T10:00:00Z", "mode": "EDIT"}}), headers={"content-type": "application/json"})
                        else: rt.fulfill(status=404, body='{}', headers={"content-type": "application/json"})
                    return r
                p.route("**/api/**", make_r(case, sent))
                p.goto(B + "/status.html?ref=GRV-TEST02"); p.wait_for_timeout(300)
                p.fill("#email", "c@x.in"); p.click("#entry-submit"); p.wait_for_selector(".otp-digit")
                fill_otp(p)
                p.click("#otp-submit"); p.wait_for_selector("#rate-card:not([hidden])")
                card = p.locator("#rate-card")
                ok(p.locator("input[name=rate-office]").count() == 5 and p.locator("input[name=rate-dept]").count() == 5, "five answers for each question " + tag)
                txt = card.inner_text()
                if lang == "en":
                    ok("How did we do?" in txt and "Neither satisfied nor dissatisfied" in txt and "MVVNL sub-station Hazratganj" in txt and "never shown publicly" in txt, "card wording " + tag, txt[:300])
                else:
                    ok("आपका अनुभव कैसा रहा?" in txt and "न संतुष्ट, न असंतुष्ट" in txt and "सार्वजनिक" in txt, "card in Hindi " + tag, txt[:300])
                    ok(not re.findall(r"status\.rate_\w+", txt), "no raw keys " + tag)
                ok(no_hscroll(p), "rating card: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS.replace("document.body", "document.getElementById('rate-card')", 1))
                ok(not bad, "rating card contrast AA " + tag, bad[:5])
                small = p.evaluate("() => Array.from(document.querySelectorAll('#rate-card .rate-opt, #rate-card button')).filter(e => e.getBoundingClientRect().height < 44).length")
                ok(small == 0, "answers and buttons at least 44px tall " + tag)
                card.screenshot(path=os.path.join(SHOTS, "20-rate-form-" + tag + ".png"))
                # nothing chosen
                p.click("#rate-send"); p.wait_for_timeout(100)
                ok(p.get_attribute("#rate-q-office", "aria-invalid") == "true" and p.inner_text("#rate-office-err") != "" and p.evaluate("document.activeElement.name") == "rate-office", "first question required, focus moves to it " + tag)
                ok(not sent, "nothing sent with an error " + tag)
                p.check("input[name=rate-office][value=DISSATISFIED]")
                ok(p.inner_text("#rate-office-err") == "", "error clears on choosing " + tag)
                p.check("input[name=rate-dept][value=SATISFIED]")
                p.fill("#rate-comment", "Took two visits")
                if width == 375 and theme == "light":
                    # a language switch keeps what was chosen
                    p.evaluate("() => GIQ.setLang(GIQ.lang === 'hi' ? 'en' : 'hi')"); p.wait_for_timeout(150)
                    ok(p.is_checked("input[name=rate-office][value=DISSATISFIED]") and p.input_value("#rate-comment") == "Took two visits", "language switch keeps the answers " + tag)
                    p.evaluate("() => GIQ.setLang(GIQ.lang === 'hi' ? 'en' : 'hi')"); p.wait_for_timeout(150)
                p.click("#rate-send"); p.wait_for_selector("#rate-msg")
                ok(sent and sent[-1]["office"] == "DISSATISFIED" and sent[-1]["dept"] == "SATISFIED" and sent[-1]["comment"] == "Took two visits" and sent[-1]["trackingRef"] == "GRV-TEST02", "rating sent " + tag, sent[-1:])
                ok(p.evaluate("document.activeElement.id") == "rate-msg", "thank-you focused " + tag)
                t2 = card.inner_text()
                ok(("Dissatisfied" in t2 and "16 October 2026" in t2) if lang == "en" else ("असंतुष्ट" in t2), "summary with change-until date " + tag, t2[:300])
                bad = p.evaluate(CONTRAST_JS.replace("document.body", "document.getElementById('rate-card')", 1))
                ok(not bad, "thank-you contrast AA " + tag, bad[:5])
                card.screenshot(path=os.path.join(SHOTS, "21-rate-thanks-" + tag + ".png"))
                p.click("#rate-change"); p.wait_for_timeout(100)
                ok(p.is_checked("input[name=rate-office][value=DISSATISFIED]") and p.input_value("#rate-comment") == "Took two visits", "change: form filled in " + tag)
                p.click("#rate-skip"); p.wait_for_timeout(100)
                ok(p.locator("#rate-change").count() == 1, "cancel goes back to the summary " + tag)
                ok(not errs, "no script errors (rating card) " + tag, errs)
                ctx.close()
    # Not now hides it; locked rating shows without a change button; an open case shows nothing
    for case, check in ((rated_case(), "skip"), (rated_case({"office": "SATISFIED", "dept": None, "department": None, "comment": "", "submittedAt": "2026-09-01T10:00:00Z", "updatedAt": "2026-09-01T10:00:00Z"}, {"can": False, "code": "LOCKED", "until": None, "mode": None}), "locked"),
                        (dict(rated_case(), status="ACKNOWLEDGED", resolutionKind=None), "open")):
        ctx = br.new_context(viewport={"width": 390, "height": 900})
        p = ctx.new_page(); errs = []; sent = []
        p.on("pageerror", lambda e: errs.append(str(e)))
        p.route("**/api/**", make_r(case, sent))
        p.goto(B + "/status.html?ref=GRV-TEST02"); p.wait_for_timeout(300)
        p.fill("#email", "c@x.in"); p.click("#entry-submit"); p.wait_for_selector(".otp-digit")
        fill_otp(p)
        p.click("#otp-submit"); p.wait_for_selector("#status-ref"); p.wait_for_timeout(300)
        if check == "skip":
            p.click("#rate-skip"); p.wait_for_timeout(100)
            ok(p.is_hidden("#rate-card"), "Not now hides the card")
        elif check == "locked":
            ok(p.is_visible("#rate-card") and p.locator("#rate-change").count() == 0 and "time to change this rating has ended" in p.inner_text("#rate-card"), "locked rating: shown, no change button", p.inner_text("#rate-card"))
        else:
            ok(p.is_hidden("#rate-card"), "open case: no rating card")
        ok(not errs, "no script errors (" + check + ")", errs)
        ctx.close()

    # Rep console: the citizen's rating on a case, and the overview tile
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx, p, errs = rep_page(lang, theme, width)
                p.evaluate("""() => { GIQ.addDepts({ Electricity: { en: 'Electricity', hi: 'विद्युत विभाग' } });
                  const c = { id: 'c1', rating: { office: 'SATISFIED', dept: 'VERY_DISSATISFIED', department: 'Electricity', comment: 'Road <b>left</b> dug up', submittedAt: '2026-10-09T05:00:00Z', updatedAt: '2026-10-10T05:00:00Z', changed: true } };
                  const few = { count: 3, needed: 5, average: null, word: null, byAnswer: { VERY_SATISFIED: 1, SATISFIED: 1, NEITHER: 0, DISSATISFIED: 1, VERY_DISSATISFIED: 0 } };
                  const many = { count: 23, needed: 5, average: 4.2, word: 'SATISFIED', byAnswer: { VERY_SATISFIED: 9, SATISFIED: 10, NEITHER: 2, DISSATISFIED: 1, VERY_DISSATISFIED: 1 } };
                  document.getElementById('content').innerHTML = '<div id="rt-test" style="max-width:720px;padding:12px">' + renderCitizenRating(c) + '<div id="rt-few">' + renderOvRatings(few) + '</div><div id="rt-many">' + renderOvRatings(many) + '</div></div>'; }""")
                txt = p.inner_text("#rt-test")
                if lang == "en":
                    ok("Citizen's rating" in txt and "Your office: Satisfied" in txt and "Electricity: Very dissatisfied" in txt and "changed by the citizen" in txt and "Only the representative and office managers" in txt, "rep: rating note " + tag, txt[:400])
                    ok("Not enough ratings yet" in p.inner_text("#rt-few") and "3 of 5" in p.inner_text("#rt-few"), "rep: no average below 5 " + tag)
                    ok("Satisfied" in p.inner_text("#rt-many .ov-rate") and "4.2 out of 5" in p.inner_text("#rt-many") and "23 ratings" in p.inner_text("#rt-many"), "rep: average shown in words and number " + tag)
                else:
                    ok("नागरिक की रेटिंग" in txt and "बहुत असंतुष्ट" in txt and "विद्युत विभाग" in txt, "rep: rating note in Hindi " + tag, txt[:300])
                    ok(not re.findall(r"rep\.(rate|ov_rate)\w*", txt), "no raw keys " + tag)
                ok(p.locator("#rt-test b").count() == 0, "comment is escaped " + tag)
                ok(no_hscroll(p), "rep rating: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS.replace("document.body", "document.getElementById('rt-test')", 1))
                ok(not bad, "rep rating contrast AA " + tag, bad[:5])
                shot(p, "22-rep-rating-" + tag, full=False)
                ok(not errs, "no script errors (rep rating) " + tag, errs)
                ctx.close()

    # Admin: Ratings page, dashboard card, case panel (real server)
    import sqlite3 as _sq
    _db = _sq.connect(os.environ.get("GRIEVIQ_TEST_DIR", "/tmp/grieviq-tests") + "/test.db", timeout=10)
    _now = __import__("datetime").datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%S.000Z")
    for i, (o, d, cm) in enumerate(((1, 2, "Nobody came for two weeks and the road is still dug up"), (2, None, None))):
        gid = "glow%d" % i
        _db.execute("INSERT INTO grievances (id, tracking_ref, citizen_phone, local_unit_id, category_id, status, description, current_tier, created_at, resolved_at, closed_at) VALUES (?, ?, '9', 'lu-hazratganj', 'water-sanitation', 'RESOLVED', 'x', 'LOCAL', ?, ?, ?)", (gid, "GRV-LOW%d" % i, _now, _now, _now))
        _db.execute("INSERT INTO case_ratings (id, grievance_id, office_score, dept_score, department, office_tier, office_id, comment, low, submitted_at, updated_at, round_closed_at) VALUES (?, ?, ?, ?, ?, 'LOCAL', 'lu-hazratganj', ?, 1, ?, ?, ?)",
                    ("crl%d" % i, gid, o, d, "Electricity" if d else None, cm, _now, _now, _now))
    _db.commit(); _db.close()
    for theme in (None, "light"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme or "dark", lang, width)
                ctx, p, errs = P("super@test.in", width=width, theme=theme, lang=lang)
                p.goto(B + "/admin-ratings.html"); p.wait_for_selector("article.fb")
                ok(p.locator("article.fb").count() == 2, "two low ratings waiting " + tag)
                body = p.inner_text("main")
                if lang == "en":
                    ok("Waiting for follow-up (2)" in body and "Very dissatisfied" in body and "Hazratganj-Ramtirth" in body and "Not enough ratings yet" in body, "ratings page wording " + tag, body[:500])
                else:
                    ok("नागरिकों की रेटिंग" in body and "बहुत असंतुष्ट" in body, "ratings page in Hindi " + tag)
                    ok(not re.findall(r"adm\.cr_\w+", body), "no raw keys " + tag)
                ok(no_hscroll(p), "ratings page: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS); ok(not bad, "ratings page contrast AA " + tag, bad[:6])
                small = p.evaluate(TARGET_JS); ok(not small, "ratings page targets >= 24px " + tag, small[:6])
                shot(p, "23-admin-ratings-" + tag, full=(width == 375))
                ok(not errs, "no script errors (ratings page) " + tag, errs)
                ctx.close()
    ctx, p, errs = P("super@test.in")
    p.goto(B + "/admin-ratings.html"); p.wait_for_selector("form.rt-follow")
    first = p.locator("form.rt-follow").first
    first.locator("button").click(); p.wait_for_timeout(150)
    ta = p.locator("form.rt-follow textarea").first
    ok(ta.get_attribute("aria-invalid") == "true" and p.evaluate("document.activeElement.tagName") == "TEXTAREA", "follow-up needs a note; focus on it")
    ta.fill("Called the office manager; repair booked for Monday")
    with p.expect_response(lambda r: r.request.method == "POST"): p.locator("form.rt-follow button").first.click()
    p.wait_for_selector("#top-msg")
    ok("Marked followed up." in p.inner_text("#top-msg") and p.locator("article.fb").count() == 1, "followed up: leaves the waiting list")
    p.click("[data-tab=DONE]"); p.wait_for_selector(".rt-done")
    ok("Called the office manager" in p.inner_text(".rt-done") and "super@test.in" in p.inner_text(".rt-done"), "followed-up tab shows what was done and who")
    p.goto(B + "/admin-dashboard.html"); p.wait_for_timeout(800)
    ok("Low ratings to follow up" in p.inner_text("main"), "dashboard card for low ratings")
    p.goto(B + "/admin-cases.html?case=glow0"); p.wait_for_selector("text=Nobody came for two weeks")
    mt = p.inner_text("main")
    ok("Citizen's rating" in mt and "Very dissatisfied" in mt and ("Low rating: waiting for follow-up." in mt or "By super@test.in" in mt), "admin case page: ratings panel", mt[-600:])
    shot(p, "24-admin-case-rating", full=True)
    ok(not errs, "no script errors (admin follow-up)", errs)
    ctx.close()
    ctx, p, errs = P("audit@test.in")
    p.goto(B + "/admin-ratings.html"); p.wait_for_selector("article.fb, .empty")
    ok(p.locator("form.rt-follow").count() == 0, "auditor: no follow-up form")
    ctx.close()
    ctx, p, errs = P("deo@test.in")
    p.goto(B + "/admin-dashboard.html"); p.wait_for_timeout(800)
    ok(p.locator('a[href*="admin-ratings"]:visible').count() == 0, "data entry operator: no Ratings link")
    ctx.close()

    # ---------- Charts (grieviq-31) ----------
    TREND = [{"month": "2025-%02d" % m if m <= 12 else "2026-%02d" % (m - 12), "received": r, "resolved": v} for m, r, v in
             zip(range(11, 23), [3, 5, 8, 6, 12, 9, 14, 11, 7, 10, 16, 4], [1, 4, 6, 7, 9, 10, 11, 12, 8, 9, 13, 2])]
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx, p, errs = rep_page(lang, theme, width)
                p.evaluate("""(trend) => {
                  const d = { groupBy: 'ward', trend, total: { pending: 31, ages: [9, 12, 7, 3] },
                    rows: [{ id: 'w1', name: 'Aishbagh', overdue: 4 }, { id: 'w2', name: 'Hazratganj-Ramtirth', overdue: 9 }, { id: 'w3', name: 'Gomti Nagar', overdue: 0 }]
                      .concat(Array.from({ length: 11 }, (x, i) => ({ id: 'x' + i, name: 'Ward ' + (i + 10), overdue: 1 }))) };
                  document.getElementById('content').innerHTML = '<div id="ch-test" style="max-width:820px;padding:12px">' + renderOvTrend(trend) + renderOvCharts(d) + '</div>';
                  GIQC.attach(document.getElementById('content'), chartText()); }""", TREND)
                txt = p.inner_text("#ch-test")
                if lang == "en":
                    ok("How long pending cases have waited" in txt and "31 cases are pending now; 3 of them for over 90 days." in txt and re.search(r"Over 90 days:\s*3", txt), "rep: age bar with counts " + tag, txt[:900])
                    ok("Most overdue: Hazratganj-Ramtirth (9)." in txt and "Showing the 10 highest of 13" in txt, "rep: overdue bars sorted, top 10 " + tag, txt[-600:])
                else:
                    ok("लंबित मामले कितने समय से रुके हैं" in txt and "90 दिन से अधिक" in txt, "rep: charts in Hindi " + tag, txt[:300])
                    ok(not re.findall(r"chart\.\w+", txt), "no raw keys " + tag)
                names = p.eval_on_selector_all("#ch-test .gc-bar .gc-name", "els => els.map(e => e.textContent)")
                ok(names[:2] == ["Hazratganj-Ramtirth", "Aishbagh"] and len(names) == 10, "bars largest first, 10 shown " + tag, names)
                ok(no_hscroll(p), "charts: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS.replace("document.body", "document.getElementById('ch-test')", 1))
                ok(not bad, "charts text contrast AA " + tag, bad[:5])
                # marks: lines and bars 3:1 against the card
                marks = p.evaluate("""() => { function rgb(s){const m=s.match(/\\d+(\\.\\d+)?/g).map(Number);return m;} function L(c){const f=v=>{v/=255;return v<=0.03928?v/12.92:Math.pow((v+0.055)/1.055,2.4)};return 0.2126*f(c[0])+0.7152*f(c[1])+0.0722*f(c[2]);}
                  const card = rgb(getComputedStyle(document.querySelector('.gc-box')).backgroundColor);
                  const cols = [getComputedStyle(document.querySelector('.gc-path.s1')).stroke, getComputedStyle(document.querySelector('.gc-path.s2')).stroke, getComputedStyle(document.querySelector('.gc-fill')).backgroundColor];
                  return cols.map(c => { const a=L(rgb(c)), b=L(card); return (Math.max(a,b)+0.05)/(Math.min(a,b)+0.05); }); }""")
                ok(all(m >= 3 for m in marks), "lines and bars 3:1 against the chart " + tag, marks)
                # keyboard: focus the line chart and move with arrows
                p.focus("#ov-trend"); p.wait_for_timeout(50)
                tip1 = p.inner_text("#ov-trend .gc-tip")
                p.keyboard.press("ArrowLeft"); p.wait_for_timeout(50)
                tip2 = p.inner_text("#ov-trend .gc-tip"); live = p.inner_text("#ov-trend .gc-sr")
                ok(p.is_visible("#ov-trend .gc-tip") and tip1 != tip2 and "16" in tip2 and "16" in live, "line chart readable with arrow keys " + tag, [tip1, tip2, live])
                p.keyboard.press("Home"); p.wait_for_timeout(50)
                ok(("2025" in p.inner_text("#ov-trend .gc-tip")), "Home goes to the first month " + tag)
                p.click("#ch-test details.gc-more >> nth=1"); p.wait_for_timeout(50)
                ok(p.locator("#ch-test details.gc-more[open] table").count() >= 1, "Show as table opens the numbers " + tag)
                small = p.evaluate("() => Array.from(document.querySelectorAll('#ch-test summary')).filter(e => e.getBoundingClientRect().height < 24).length")
                ok(small == 0, "table toggles at least 24px " + tag)
                shot(p, "25-rep-charts-" + tag, full=True)
                ok(not errs, "no script errors (rep charts) " + tag, errs)
                ctx.close()
    for theme in (None, "light"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme or "dark", lang, width)
                ctx, p, errs = P("super@test.in", width=width, theme=theme, lang=lang)
                p.goto(B + "/admin-dashboard.html"); p.wait_for_selector("#stats-area .st-tiles")
                body = p.inner_text("#stats-area")
                if lang == "en":
                    ok("How complaints are moving" in body and "Received this month" in body and "Last 12 months" in body and "How long pending cases have waited" in body, "admin charts section " + tag, body[:300])
                else:
                    ok("शिकायतों की प्रगति" in body and "इस महीने प्राप्त" in body, "admin charts in Hindi " + tag, body[:300])
                    ok(not re.findall(r"(adm\.cm_|chart\.)\w+", body), "no raw keys " + tag)
                ok(no_hscroll(p), "admin charts: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS.replace("document.body", "document.getElementById('stats-area')", 1))
                ok(not bad, "admin charts contrast AA " + tag, bad[:5])
                small = p.evaluate("() => Array.from(document.querySelectorAll('#stats-area summary, #stats-area select')).filter(e => e.getBoundingClientRect().height < 24).length")
                ok(small == 0, "admin chart controls >= 24px " + tag)
                p.locator("#stats-area").screenshot(path=os.path.join(SHOTS, "26-admin-charts-" + tag + ".png"))
                ok(not errs, "no script errors (admin charts) " + tag, errs)
                ctx.close()
    ctx, p, errs = P("super@test.in")
    p.goto(B + "/admin-dashboard.html"); p.wait_for_selector("#st-area")
    with p.expect_response(lambda r: "case-stats?area=kanpur" in r.url): p.select_option("#st-area", "kanpur")
    p.wait_for_timeout(200)
    ok(p.input_value("#st-area") == "kanpur" and p.evaluate("document.activeElement.id") == "st-area", "area filter reloads and keeps focus")
    p.focus("#st-trend"); p.keyboard.press("ArrowLeft"); p.wait_for_timeout(50)
    ok(p.is_visible("#st-trend .gc-tip"), "admin line chart works with the keyboard")
    ok(not errs, "no script errors (admin area filter)", errs)
    ctx.close()

    # ---------- Department dashboard (grieviq-32) ----------
    import hashlib as _hl
    WORK = os.environ.get("GRIEVIQ_TEST_DIR", "/tmp/grieviq-tests")
    def last_code(email):
        for line in reversed(open(WORK + "/mail.log", encoding="utf-8").read().strip().split("\n")):
            if not line: continue
            m = _json.loads(line)
            if m.get("to") and m["to"][0] == email and re.search(r"(\d{6})", m["subject"]):
                return re.search(r"(\d{6})", m["subject"]).group(1)
        return None
    # 1. Real sign-in with the emailed code (the phone shares its location: 50 km away)
    ctx = br.new_context(viewport={"width": 390, "height": 900}, permissions=["geolocation"], geolocation={"latitude": 26.4499, "longitude": 80.3318, "accuracy": 14})
    p = ctx.new_page(); errs = []
    p.on("pageerror", lambda e: errs.append(str(e)))
    p.goto(B + "/dept"); p.wait_for_selector("#email")
    p.click("#send"); p.wait_for_timeout(100)
    ok(p.get_attribute("#email", "aria-invalid") == "true" and p.inner_text("#email-err") != "", "dept sign-in: email checked")
    p.fill("#email", "ae.zone3@nic.in"); p.click("#send"); p.wait_for_selector("#code")
    ok(p.evaluate("document.activeElement.id") == "code", "code box focused")
    p.fill("#code", "000000"); p.click("#verify"); p.wait_for_timeout(300)
    ok("isn't right" in p.inner_text("#code-err") or "expired" in p.inner_text("#code-err"), "wrong code message", p.inner_text("#code-err"))
    p.fill("#code", last_code("ae.zone3@nic.in")); p.click("#verify"); p.wait_for_selector(".tiles")
    ok("Ravi Kumar" in p.inner_text("#who") and "Jal Kal Vibhag Lucknow" in p.inner_text("#who"), "signed in: name and office shown")
    ok(p.locator('[data-open="g21"]').count() == 1, "new case listed")
    ok("Your office's performance" in p.inner_text("main") and "First reply" in p.inner_text("main"), "officer: own office's performance shown")
    p.click('[data-open="g21"]'); p.wait_for_selector("#case-h")
    body = p.inner_text("main")
    ok("Second leak" in body and "9876500000" not in body and "citizen" in body.lower(), "case: complaint shown, no phone", body[:300])
    p.click("#r-send"); p.wait_for_timeout(100)
    ok(p.inner_text("#kind-err") != "", "reply: choose the position")
    p.check('input[name=kind][value=SCHEDULED]'); p.wait_for_selector("#r-date")
    p.click("#r-send"); p.wait_for_timeout(100)
    ok(p.get_attribute("#r-date", "aria-invalid") == "true", "reply: date needed")
    import datetime as _dt
    p.fill("#r-date", (_dt.date.today() + _dt.timedelta(days=3)).isoformat()); p.fill("#r-note", "Team visiting Thursday")
    with p.expect_response(lambda r: "/api/dept/reply" in r.url): p.click("#r-send")
    p.wait_for_selector("#msg")
    ok("work is scheduled" in p.inner_text("#msg") and "scheduled the work" in p.inner_text("main"), "reply sent and shown in the history", p.inner_text("#msg"))
    # work done with a photo
    p.check('input[name=kind][value=DONE_CLAIMED]'); p.wait_for_selector("#r-photo")
    png = bytes([0x89, 0x50, 0x4e, 0x47, 0x0d, 0x0a, 0x1a, 0x0a]) + bytes(range(40))
    with p.expect_response(lambda r: "/api/dept/photo" in r.url): p.set_input_files("#r-photo", files=[{"name": "work.png", "mimeType": "image/png", "buffer": png}])
    p.wait_for_selector("[data-rm-photo]")
    ok(p.locator("[data-rm-photo]").count() == 1, "work photo added")
    chk = p.inner_text(".photo-item .ev-list")
    ok("You added this photo about" in chk and "km" in chk and "No date in photo" in chk, "officer sees the checks under the photo (grieviq-34)", chk)
    ok("Your location: shared" in p.inner_text("#r-loc"), "officer's location line", p.inner_text("#r-loc"))
    with p.expect_response(lambda r: "/api/dept/reply" in r.url): p.click("#r-send")
    p.wait_for_selector("#msg:has-text('check it on site')")
    ok("check it on site" in p.inner_text("#msg"), "work done sent", p.inner_text("#msg"))
    p.click("#back"); p.wait_for_selector(".tiles")
    ok("Waiting for check" in p.inner_text('li:has([data-open="g21"])'), "list shows waiting for the check")
    p.click("#sign-out"); p.wait_for_selector("#email")
    ok("signed out" in p.inner_text("#msg"), "signed out")
    ok(not errs, "no script errors (dept sign-in flow)", errs)
    ctx.close()
    # 2. Every theme, language and width (session made directly, to stay under the code limit)
    def dept_signin(ctx):
        # A code made directly in the database (the 3-codes-per-15-minutes limit
        # applies to asking for codes, not to signing in), then the real sign-in.
        code = "135790"
        d = _sq.connect(WORK + "/test.db", timeout=10)
        d.execute("INSERT INTO dept_codes (id, email, code_hash, expires_at, created_at) VALUES (?, 'ae.zone3@nic.in', ?, ?, ?)",
                  ("ui-" + str(_dt.datetime.utcnow().timestamp()), _hl.sha256(("ae.zone3@nic.in:" + code).encode()).hexdigest(),
                   (_dt.datetime.utcnow() + _dt.timedelta(minutes=10)).isoformat() + "Z", (_dt.datetime.utcnow() + _dt.timedelta(seconds=5)).isoformat() + "Z"))
        d.commit(); d.close()
        r = ctx.request.post(B + "/api/dept/verify", data={"email": "ae.zone3@nic.in", "code": code})
        return r.status
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx = br.new_context(viewport={"width": width, "height": 900}, color_scheme=theme)
                ok(dept_signin(ctx) == 200, "signed in " + tag)
                ctx.add_init_script("try{localStorage.setItem('%s','%s')}catch(e){}" % (LANGKEY, lang))
                p = ctx.new_page(); errs = []
                p.on("pageerror", lambda e: errs.append(str(e)))
                p.goto(B + "/dept"); p.wait_for_selector(".tiles")
                txt = p.inner_text("main")
                if lang == "hi":
                    ok("आपके कार्यालय को भेजी गई शिकायतें" in txt and not re.findall(r"dd\.\w+", txt), "dept list in Hindi " + tag, txt[:200])
                ok(no_hscroll(p), "dept list: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS); ok(not bad, "dept list contrast AA " + tag, bad[:5])
                small = p.evaluate(TARGET_JS); ok(not small, "dept list targets >= 24px " + tag, small[:5])
                shot(p, "27-dept-list-" + tag, full=(width == 375))
                p.click('[data-open="g20"]'); p.wait_for_selector("#case-h")
                ok(no_hscroll(p), "dept case: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS); ok(not bad, "dept case contrast AA " + tag, bad[:5])
                p.goto(B + "/dept#case=g21"); p.wait_for_selector("#case-h")
                if p.locator("input[name=kind]").count():
                    p.check('input[name=kind][value=CANT_DO]'); p.wait_for_selector("#r-note")
                    p.click("#r-send"); p.wait_for_timeout(100)
                    bad = p.evaluate(CONTRAST_JS); ok(not bad, "dept reply form contrast AA " + tag, bad[:5])
                    small = p.evaluate(TARGET_JS); ok(not small, "dept reply targets >= 24px " + tag, small[:5])
                shot(p, "28-dept-case-" + tag, full=True)
                ok(not errs, "no script errors (dept) " + tag, errs)
                ctx.close()
    # 3. Admin: officers panel
    for theme in (None, "light"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme or "dark", lang, width)
                ctx, p, errs = P("super@test.in", width=width, theme=theme, lang=lang)
                p.goto(B + "/admin-departments.html?area=lucknow"); p.wait_for_selector(".office")
                btn = p.locator("article:has-text('Jal Kal Vibhag Lucknow') [data-officers]")
                btn.click(); p.wait_for_selector(".op h5")
                ok(btn.get_attribute("aria-expanded") == "true", "officers panel opens " + tag)
                ptxt = p.inner_text(".op")
                if lang == "en":
                    ok("Signed on" in ptxt and "Ravi Kumar" in ptxt and "Asha Verma" in ptxt and "Removed" in ptxt, "panel: agreement and officers " + tag, ptxt[:400])
                else:
                    ok("GrievIQ के साथ समझौता" in ptxt and not re.findall(r"adm\.do_\w+", ptxt), "panel in Hindi " + tag)
                ok(no_hscroll(p), "officers panel: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS.replace("createTreeWalker(document.body", "createTreeWalker(document.querySelector('.op')", 1)); ok(not bad, "officers panel contrast AA " + tag, bad[:5])
                p.locator(".op").screenshot(path=os.path.join(SHOTS, "29-admin-officers-" + tag + ".png"))
                ok(not errs, "no script errors (officers panel) " + tag, errs)
                ctx.close()
    ctx, p, errs = P("super@test.in")
    p.goto(B + "/admin-departments.html?area=lucknow"); p.wait_for_selector(".office")
    p.locator("article:has-text('Lucknow Nagar Nigam Control Room') [data-officers]").first.click(); p.wait_for_selector("#op-agr-form")
    p.click("#op-agr-save"); p.wait_for_timeout(300)
    ok(p.get_attribute("#op-signed-on", "aria-invalid") == "true" and p.evaluate("document.activeElement.id") == "op-signed-on", "agreement form: errors next to fields, focus on the first")
    p.fill("#op-signed-on", "2026-10-05"); p.fill("#op-signed-by", "S. Mishra, Additional Municipal Commissioner"); p.fill("#op-doc", "LMC/IT/2026/88")
    with p.expect_response(lambda r: "dept-officers" in r.url and r.request.method == "POST"): p.click("#op-agr-save")
    p.wait_for_selector("#op-add-form")
    p.fill("#op-name", "Neha Gupta"); p.fill("#op-desig", "Zonal Officer, Zone 1"); p.fill("#op-email", "zo1.lmc@nic.in")
    with p.expect_response(lambda r: "dept-officers" in r.url and r.request.method == "POST"): p.click("#op-add-go")
    p.wait_for_selector("#op-msg")
    ok("Neha Gupta added" in p.inner_text("#op-msg") and "zo1.lmc@nic.in" in p.inner_text(".op"), "officer added from the page")
    p.locator("[data-op-leave]").first.click(); p.wait_for_selector("#op-lv-form")
    ok(p.evaluate("document.activeElement.id") == "op-lv-TRANSFER", "leaving: focus on the first choice")
    p.check("#op-lv-LEFT"); p.wait_for_selector("#op-lv-reason")
    ok(p.locator(".op-last").count() == 1, "leaving: the only officer: asked who replies after")
    with p.expect_response(lambda r: "dept-officers" in r.url and r.request.method == "POST"): p.click("#op-lv-go")
    p.wait_for_selector("#op-lv-reason[aria-invalid=true]")
    ok(p.evaluate("document.activeElement.id") == "op-lv-reason" and p.locator("#op-lv-succ-err").count() == 1, "removal needs a reason and a choice; focus on the first")
    p.fill("#op-lv-reason", "Left the zonal office"); p.check("#op-lv-EMPTY")
    ok(p.input_value("#op-lv-reason") == "Left the zonal office", "the reason is kept when choosing")
    with p.expect_response(lambda r: "dept-officers" in r.url and r.request.method == "POST"): p.click("#op-lv-go")
    p.wait_for_selector("#op-msg")
    ok("removed and signed out" in p.inner_text("#op-msg") and "nobody to reply" in p.inner_text("#op-msg"), "officer removed from the page, office left with nobody on purpose")
    ok("Nobody can reply" in p.inner_text(".op") and p.locator("article:has(h3:text-is('Lucknow Nagar Nigam Control Room')) .pill.nobody").count() == 1, "the office says it has nobody to reply")
    ok(not errs, "no script errors (officers add/remove)", errs)
    ctx.close()
    ctx, p, errs = P("audit@test.in")
    p.goto(B + "/admin-departments.html?area=lucknow"); p.wait_for_selector(".office")
    p.locator("article:has-text('Jal Kal Vibhag Lucknow') [data-officers]").click(); p.wait_for_selector(".op h5")
    ok(p.locator("#op-add-form, [data-op-leave], #op-agr-open").count() == 0, "auditor: officers panel read only")
    ctx.close()
    ctx, p, errs = P("deo@test.in")
    p.goto(B + "/admin-departments.html?area=lucknow"); p.wait_for_selector(".office")
    ok(p.locator("[data-officers]").count() == 0, "data entry operator: no officers button")
    ctx.close()
    # 4. Rep console: who replied from the department
    ctx, p, errs = rep_page("en", "light", 1280)
    show(p, steps_case(["FORWARDED", "SCHEDULED"], "WITH_DEPT", extra={"expectedDate": "2026-10-12"}))
    p.evaluate("""() => { const c = currentCases[0]; c.deptSteps[1].byDept = true; c.deptSteps[1].byDeptName = 'Ravi Kumar'; renderContent(); }""")
    ok("Replied by the department (Ravi Kumar)" in p.inner_text("#ds-c9"), "rep console: department reply labelled with the officer")
    ctx.close()

    # ---------- Department performance (grieviq-33) ----------
    for theme in (None, "light"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme or "dark", lang, width)
                ctx, p, errs = P("super@test.in", width=width, theme=theme, lang=lang)
                p.goto(B + "/admin-dept-performance.html"); p.wait_for_selector("table.pf")
                txt = p.inner_text("main")
                if lang == "en":
                    ok("Department performance" in txt and "Jal Kal Vibhag Lucknow" in txt and "Too few cases" in txt and "no single score or ranking" in txt, "performance page " + tag, txt[:400])
                else:
                    ok("विभागों का प्रदर्शन" in txt and not re.findall(r"(adm\.pf_|perf\.)\w+", txt), "performance page in Hindi " + tag, txt[:300])
                ok(no_hscroll(p), "performance page: no sideways page scroll " + tag)
                bad = p.evaluate(CONTRAST_JS); ok(not bad, "performance page contrast AA " + tag, bad[:5])
                small = p.evaluate(TARGET_JS); ok(not small, "performance page targets >= 24px " + tag, small[:5])
                shot(p, "30-dept-perf-" + tag, full=(width == 375))
                ok(not errs, "no script errors (performance) " + tag, errs)
                ctx.close()
    ctx, p, errs = P("super@test.in")
    p.goto(B + "/admin-dept-performance.html"); p.wait_for_selector("table.pf")
    with p.expect_response(lambda r: "detail=" in r.url): p.locator("[data-detail]").first.click()
    p.wait_for_selector("#pf-d-h")
    ok(p.evaluate("document.activeElement.id") == "pf-d-h" and p.locator("#pf-trend").count() == 1, "details: heading focused, 12-month chart shown")
    p.click("#pf-close"); p.wait_for_timeout(100)
    ok(p.locator("#pf-d-h").count() == 0 and p.evaluate("document.activeElement.dataset.detail") is not None, "details closed, focus back on the button")
    with p.expect_response(lambda r: "by=type" in r.url): p.click('[data-by="type"]')
    p.wait_for_timeout(300)
    ok(p.get_attribute('[data-by="type"]', "aria-pressed") == "true" and "department type" in p.inner_text("table.pf thead").lower(), "switch to department types")
    p.select_option("#pf-period", "custom"); p.wait_for_selector("#pf-from")
    p.fill("#pf-from", "2026-10-01"); p.fill("#pf-to", "2026-09-01")
    with p.expect_response(lambda r: "dept-performance" in r.url): p.click("#pf-apply")
    p.wait_for_timeout(200)
    ok("check the dates" in p.inner_text("main"), "wrong dates: message shown")
    href = p.get_attribute("#pf-csv", "href")
    ok("format=csv" in href, "CSV link")
    ok(not errs, "no script errors (performance actions)", errs)
    ctx.close()
    ctx, p, errs = P("deo@test.in")
    p.goto(B + "/admin-dashboard.html"); p.wait_for_timeout(800)
    ok(p.locator('a[href*="admin-dept-performance"]:visible').count() == 0, "data entry operator: no performance link")
    ctx.close()
    # Rep overview: departments on your cases
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            tag = "%s-%s" % (theme, lang)
            ctx, p, errs = rep_page(lang, theme, 375)
            p.evaluate("""() => { const dp = { min: 5, rows: [
                { key: 'id:o1', name: 'Jal Kal Zone 3', department: 'Water Supply', m: { forwarded: 12, withNow: 3, overdueNow: 2, firstReply: { n: 10, median: 1.5 }, onTime: { n: 6, d: 9, pct: 66.7 }, toFixed: { n: 7, median: 6 }, fixedFirst: { n: 5, d: 7, pct: 71.4 }, reopened: { n: 0, d: 4, pct: null }, rating: { count: 6, average: 3.8, word: 'SATISFIED' }, ownReplies: { n: 8, d: 12, pct: 66.7 }, notOurs: 1, cantDo: 0 } },
                { key: 'name:x', name: 'MVVNL sub-station', department: 'Electricity', m: { forwarded: 2, withNow: 1, overdueNow: 0, firstReply: { n: 2, median: null }, onTime: { n: 1, d: 2, pct: null }, toFixed: { n: 0, median: null }, fixedFirst: { n: 0, d: 0, pct: null }, reopened: { n: 0, d: 0, pct: null }, rating: { count: 0 }, ownReplies: { n: 0, d: 2, pct: null }, notOurs: 0, cantDo: 0 } } ] };
              document.getElementById('content').innerHTML = '<div id="dp-test" style="padding:12px">' + renderOvDepts(dp) + '</div>'; }""")
            t2 = p.inner_text("#dp-test")
            if lang == "en":
                ok("Departments on your cases" in t2 and "66.7%" in t2 and "6 of 9" in t2 and "Too few cases (2)" in t2 and "1.5 days" in t2, "rep: department table " + tag, t2[:400])
            else:
                ok("आपके मामलों पर विभाग" in t2 and not re.findall(r"(perf|rep\.ov_dp)\.?\w*_\w+", t2), "rep: department table in Hindi " + tag, t2[:300])
            bad = p.evaluate(CONTRAST_JS.replace("createTreeWalker(document.body", "createTreeWalker(document.getElementById('dp-test')", 1)); ok(not bad, "rep department table contrast AA " + tag, bad[:5])
            ok(no_hscroll(p), "rep department table: no page scroll " + tag)
            shot(p, "31-rep-depts-" + tag, full=False)
            ok(not errs, "no script errors (rep departments) " + tag, errs)
            ctx.close()

    # ---------- themes, languages, widths: contrast + reflow ----------
    for theme in (None, "light"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                ctx, p, errs = P("super@test.in", width=width, theme=theme, lang=lang)
                p.goto(B + "/admin-departments.html?area=lucknow"); p.wait_for_selector(".office")
                tag = "%s-%s-%d" % (theme or "dark", lang, width)
                ok(no_hscroll(p), "no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS)
                ok(not bad, "contrast AA " + tag, bad[:6])
                small = p.evaluate(TARGET_JS)
                ok(not small, "targets >= 24px " + tag, small[:6])
                if lang == "hi":
                    body = p.inner_text("main")
                    ok("विभाग" in body and "Departments" not in p.inner_text("h1"), "Hindi page " + tag)
                    leftover = re.findall(r"adm\.dp_\w+", body)
                    ok(not leftover, "no raw keys " + tag, leftover)
                if width == 375:
                    shot(p, "05-" + tag)
                else:
                    shot(p, "06-" + tag, full=False)
                # form too
                p.click("#add-office"); p.click("#o-save"); p.wait_for_selector("#o-summary")
                ok(no_hscroll(p), "form no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS); ok(not bad, "form contrast " + tag, bad[:6])
                if width == 375 or lang == "hi": shot(p, "07-form-" + tag)
                ok(not errs, "no script errors " + tag, errs)
                ctx.close()

    # ---------- data entry operator ----------
    ctx, p, errs = P("deo@test.in")
    p.goto(B + "/admin-departments.html?area=lucknow"); p.wait_for_selector(".office")
    ok("sent as requests" in p.inner_text(".note"), "operator told changes are requests")
    p.click("#add-office")
    ok(p.inner_text("#o-save") == "Send for approval" and p.is_visible("#o-reason"), "operator: Send for approval + reason")
    p.fill("#o-name-en", "KGMU Hospital Helpdesk"); p.check('[data-dept="Health"]'); p.fill("#o-helpline", "0522 2257540"); p.fill("#o-source", "https://www.kgmu.org")
    p.click("#o-save"); p.wait_for_timeout(150)
    ok("reason" in p.inner_text("#o-summary").lower(), "reason required")
    p.fill("#o-reason", "Found on the KGMU website"); save(p, "#o-save")
    ok("sent for approval" in p.inner_text("#top-msg"), "request sent")
    ok("KGMU Hospital Helpdesk" in p.inner_text("main") and p.locator('.office h3:has-text("KGMU")').count() == 0, "waiting note, not yet listed")
    p.click("#add-office"); p.wait_for_selector("#office-form")
    ok(p.inner_text("#nt-open") == "+ Ask for a new type", "operator: Ask for a new type")
    p.click("#nt-open"); p.fill("#nt-en", "Parks"); p.fill("#nt-hi", "उद्यान विभाग")
    p.click("#nt-save"); p.wait_for_timeout(150); ok(p.is_visible("#nt-reason-err"), "operator: reason needed")
    p.fill("#nt-reason", "Many stray cattle complaints"); save(p, "#nt-save")
    ok("sent for approval" in p.inner_text("#nt-msg"), "operator: request sent from the form")
    p.click("#o-cancel")
    c = p.locator('.office:has(h3:has-text("CMO Lucknow"))'); c.locator("[data-retire]").click()
    p.fill("#r-reason", "Number no longer answered, checked twice"); save(p, "#r-go")
    c = p.locator('.office:has(h3:has-text("CMO Lucknow"))')
    ok("Retirement waiting for approval" in c.inner_text() and c.locator("[data-edit]").count() == 0, "pill + no actions while waiting")
    shot(p, "08-operator")
    ok(not errs, "no script errors (operator)", errs)
    ctx.close()

    # ---------- approver sees requests in Change requests (EN + HI) ----------
    for lang in ("hi", "en"):
        ctx, p, errs = P("ops@test.in", lang=lang)
        p.goto(B + "/admin-change-requests.html"); p.wait_for_selector(".req")
        txt = p.inner_text("main")
        ok(("New department office: KGMU Hospital Helpdesk" if lang == "en" else "नया विभाग कार्यालय: KGMU Hospital Helpdesk") in txt, "change request label " + lang, txt[:300])
        ok(("Retire department office: CMO Lucknow" if lang == "en" else "विभाग कार्यालय सेवानिवृत्त करें: CMO Lucknow") in txt, "retire label " + lang)
        shot(p, "09-change-requests-" + lang)
        if lang == "en":
            p.check('.req:has-text("KGMU") .pick'); p.check('.req:has-text("CMO Lucknow") .pick')
            save(p, "button:has-text('Approve selected')")
            p.wait_for_timeout(800)
        ok(not errs, "no script errors (change requests %s)" % lang, errs)
        ctx.close()
    ctx, p, errs = P("ops@test.in")
    p.goto(B + "/admin-departments.html?area=lucknow"); p.wait_for_selector(".office")
    ok(p.locator('.office h3:has-text("KGMU Hospital Helpdesk")').count() == 1, "approved office now listed")
    ok(p.locator('.office h3:has-text("CMO Lucknow")').count() == 0, "approved retirement applied")
    ctx.close()

    # ---------- auditor, moderator ----------
    ctx, p, errs = P("audit@test.in")
    p.goto(B + "/admin-departments.html?area=lucknow"); p.wait_for_selector(".office")
    ok(p.locator("#add-office, [data-edit], [data-retire], #open-upload").count() == 0, "auditor: no change buttons")
    ok("view this page but not change it" in p.inner_text(".note"), "auditor told read-only")
    ctx.close()
    ctx, p, errs = P("mod@test.in")
    p.goto(B + "/admin-departments.html"); p.wait_for_selector("#adm-denied")
    ok("isn't part of your role" in p.inner_text("#adm-denied"), "moderator: denied page")
    ok(p.locator('a[href="/admin-departments.html"]:visible').count() == 0, "moderator: no menu link")
    ctx.close()

    # ---------- other admin page still gets the link ----------
    ctx, p, errs = P("super@test.in")
    p.goto(B + "/admin-areas.html"); p.wait_for_selector("#adm-menu")
    p.click('button[data-key="adm.nav_g_data"]')
    ok(p.is_visible('#adm-drop-data a[href="/admin-departments.html"]'), "link on other admin pages")
    ok(not errs, "no script errors (areas)", errs)
    ctx.close()
    # ---------- Forward through the dashboard; checks on department photos (grieviq-34) ----------
    DIR34 = _json.loads(_json.dumps(DIR)); DIR34["offices"][1]["hasDashboard"] = True; DIR34["offices"][0]["hasDashboard"] = False
    def show34(p, case):
        p.evaluate("""([dir, c]) => { takeDepartments({ departments: ['Water Supply','Electricity','Health','Other'], deptDirectory: dir }); currentCases = [c];
          const box = document.getElementById('content'); box.innerHTML = '<div id="ds-test" style="max-width:640px;padding:12px">' + renderFollowupControl(c) + '<p id="resolve-status" role="status"></p></div>';
          renderContent = () => { const cc = currentCases[0]; document.getElementById('ds-test').innerHTML = renderFollowupControl(cc) + '<p id="resolve-status" role="status"></p>'; attachFollowupHandlers(); };
          attachFollowupHandlers(); window.scrollTo(0, 0); }""", [DIR34, case])
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx, p, errs = rep_page(lang, theme, width)
                fresh = steps_case([], "NONE"); fresh["deptSteps"] = []; fresh["deptState"] = None
                show34(p, fresh)
                ok(p.inner_text(".followup-btn") == ("Forward" if lang == "en" else "अग्रेषित करें"), "button says Forward " + tag, p.inner_text(".followup-btn"))
                p.select_option("#followup-dept-c9", "Water Supply"); p.wait_for_timeout(50)
                p.select_option("#followup-office-c9", "o1"); p.wait_for_timeout(50)
                ok(p.locator('input.followup-ch[value="DASHBOARD"]').count() == 0, "no dashboard choice for an office without officers " + tag)
                p.select_option("#followup-office-c9", "o2"); p.wait_for_timeout(50)
                dash = p.locator('input.followup-ch[value="DASHBOARD"]')
                ok(dash.count() == 1 and dash.is_checked(), "dashboard choice first and chosen " + tag)
                hint = p.inner_text("#fwd-dash-hint-c9")
                ok(("Jal Kal Zone 3's officers see it" in hint) if lang == "en" else ("तुरंत ईमेल" in hint), "dashboard hint " + tag, hint)
                ok(no_hscroll(p), "forward form: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS.replace("document.body", "document.getElementById('ds-test')", 1)); ok(not bad, "forward form contrast AA " + tag, bad[:5])
                small = p.evaluate(TARGET_JS); ok(not small, "forward form targets >= 24px " + tag, small[:5])
                shot(p, "34-forward-" + tag, full=False)
                p.click(".followup-btn"); p.wait_for_timeout(300)
                ok(posts[-1][0] == "grievances/c9/add-followup" and posts[-1][1]["channel"] == "DASHBOARD" and posts[-1][1]["officeId"] == "o2", "sent through the dashboard " + tag, posts[-1:])
                # forwarded before: "Forward again"; the department's photo with its checks
                case = steps_case(["FORWARDED", "DONE_CLAIMED"], "NEEDS_CHECK")
                case["deptSteps"][1].update({"byDept": True, "byDeptName": "Ravi Kumar", "photos": [{"url": "/logo-mark.svg", "thumbUrl": "/logo-mark.svg",
                    "warnings": [{"code": "DEVICE_FAR_FROM_PIN", "level": "warn", "metres": 357000, "accuracy": 14}, {"code": "NO_DATE", "level": "info"}]}]})
                show34(p, case)
                ok(p.inner_text(".followup-btn") == ("Forward again" if lang == "en" else "फिर से अग्रेषित करें"), "button says Forward again " + tag)
                ev = p.inner_text("#ds-c9 .ev-list")
                ok(("The department added this photo about 357" in ev) if lang == "en" else ("विभाग ने यह फ़ोटो" in ev), "rep sees checks on the department's photo " + tag, ev)
                ok(not re.findall(r"rep\.evd?_\w+", p.inner_text("#ds-test")), "no raw keys " + tag)
                bad = p.evaluate(CONTRAST_JS.replace("document.body", "document.getElementById('ds-test')", 1)); ok(not bad, "department photo checks contrast AA " + tag, bad[:5])
                shot(p, "35-dept-photo-checks-" + tag, full=False)
                ok(not errs, "no script errors (grieviq-34 rep) " + tag, errs)
                ctx.close()
    # Officer: the checks on the reply form, every theme and language
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            tag = "%s-%s" % (theme, lang)
            ctx = br.new_context(viewport={"width": 375, "height": 900}, color_scheme=theme, permissions=["geolocation"], geolocation={"latitude": 26.85, "longitude": 80.94, "accuracy": 10})
            ok(dept_signin(ctx) == 200, "signed in " + tag)
            ctx.add_init_script("try{localStorage.setItem('%s','%s')}catch(e){}" % (LANGKEY, lang))
            p = ctx.new_page(); errs = []
            p.on("pageerror", lambda e: errs.append(str(e)))
            p.goto(B + "/dept#case=g21"); p.wait_for_selector("#case-h")
            hist = p.locator(".steps .ev-list")
            ok(hist.count() >= 1 and (("The department added this photo" in hist.first.inner_text() or "No location in the photo" in hist.first.inner_text()) if lang == "en" else "फ़ोटो" in hist.first.inner_text()), "history shows the checks " + tag, hist.first.inner_text() if hist.count() else "")
            bad = p.evaluate(CONTRAST_JS); ok(not bad, "officer case with checks contrast AA " + tag, bad[:5])
            shot(p, "36-dept-checks-" + tag, full=True)
            ok(not errs, "no script errors (grieviq-34 officer) " + tag, errs)
            ctx.close()

    # ---------- The bell for department officers (grieviq-35) ----------
    FAKE_PUSH = """(() => {
      let sub = null;
      const b64u = (a) => btoa(String.fromCharCode.apply(null, new Uint8Array(a))).replace(/\\+/g, '-').replace(/\\//g, '_').replace(/=+$/, '');
      let perm = 'default';
      if (window.Notification) {
        Object.defineProperty(Notification, 'permission', { configurable: true, get: () => perm });
        Notification.requestPermission = async () => { perm = 'granted'; return perm; };
      }
      if (window.PushManager) {
        PushManager.prototype.getSubscription = async function () { return sub; };
        PushManager.prototype.subscribe = async function () {
          const kp = await crypto.subtle.generateKey({ name: 'ECDH', namedCurve: 'P-256' }, true, ['deriveBits']);
          const raw = await crypto.subtle.exportKey('raw', kp.publicKey);
          const keys = { p256dh: b64u(raw), auth: b64u(crypto.getRandomValues(new Uint8Array(16))) };
          sub = { endpoint: 'https://fcm.googleapis.com/fcm/send/ui-test-device', toJSON: () => ({ endpoint: 'https://fcm.googleapis.com/fcm/send/ui-test-device', keys }), unsubscribe: async () => { sub = null; return true; } };
          return sub;
        };
      }
    })();"""
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx = br.new_context(viewport={"width": width, "height": 900}, color_scheme=theme, permissions=["notifications"])
                ok(dept_signin(ctx) == 200, "signed in " + tag)
                ctx.add_init_script("try{localStorage.setItem('%s','%s')}catch(e){}" % (LANGKEY, lang))
                ctx.add_init_script(FAKE_PUSH)
                p = ctx.new_page(); errs = []
                p.on("pageerror", lambda e: errs.append(str(e)))
                p.goto(B + "/dept"); p.wait_for_selector(".tiles"); p.wait_for_selector("#rn-bell")
                p.wait_for_timeout(400)
                badge = p.locator(".rn-badge")
                ok(badge.count() == 1 and int(badge.inner_text()) >= 1, "bell shows unread alerts " + tag)
                ok(p.get_attribute("#rn-bell", "aria-expanded") == "false", "bell is a closed disclosure " + tag)
                p.click("#rn-bell"); p.wait_for_selector("#rn-panel .rn-item")
                ok(p.evaluate("document.activeElement.id") == "rn-h", "focus moves to the list heading " + tag)
                p.wait_for_selector("#rn-on"); p.wait_for_timeout(300)
                txt = p.inner_text("#rn-panel")
                if lang == "en":
                    ok("Sent back after the check on site: GRV-DEPT23" in txt and "Past the target time: GRV-DEPT24" in txt and "Turn on alerts on this device" in txt, "officer's alerts listed " + tag, txt[:400])
                else:
                    ok("मौके की जाँच के बाद वापस भेजी गई: GRV-DEPT23" in txt and "इस डिवाइस पर सूचनाएँ चालू करें" in txt and not re.findall(r"(rn|dd)\.\w+_\w+", txt), "officer's alerts in Hindi " + tag, txt[:400])
                ok("Sewer" not in txt and "9876522222" not in txt, "no complaint text in the bell " + tag)
                ok(no_hscroll(p), "bell: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS.replace("document.body", "document.getElementById('rn-panel')", 1)); ok(not bad, "bell contrast AA " + tag, bad[:5])
                small = p.evaluate(TARGET_JS); ok(not small, "bell targets >= 24px " + tag, small[:5])
                shot(p, "37-dept-bell-" + tag, full=False)
                if theme == "light" and width == 1280:
                    # turn on alerts, send a test, then open an alert
                    with p.expect_response(lambda r: "/api/dept/notifications" in r.url and r.request.method == "POST"): p.click("#rn-on")
                    p.wait_for_selector("#rn-test")
                    ok(("Turn off on this device" if lang == "en" else "इस डिवाइस पर बंद करें") in p.inner_text("#rn-panel"), "alerts on for this device " + tag)
                    before = len(open(WORK + "/push.log").read().strip().split("\n"))
                    p.click("#rn-test"); p.wait_for_function("document.getElementById('rn-pstat') && document.getElementById('rn-pstat').textContent.length > 12")
                    for _w in range(50):   # the status text can show before the push is written
                        after = open(WORK + "/push.log").read().strip().split("\n")
                        if len(after) > before: break
                        p.wait_for_timeout(100)
                    ok(len(after) > before and "ui-test-device" in after[-1], "test alert went to the device " + tag, after[-1:])
                    reg = p.evaluate("navigator.serviceWorker.getRegistration('/dept').then(r => r ? r.scope : null)")
                    ok(reg and reg.endswith("/dept"), "alerts belong to the department dashboard " + tag, reg)
                    n0 = int(p.inner_text(".rn-badge"))
                    item = p.locator("#rn-panel .rn-item.unread [data-n]").filter(has_text="GRV-DEPT23").first
                    item.click()
                    p.wait_for_selector("#case-h"); p.wait_for_timeout(400)
                    ok("GRV-DEPT23" in p.inner_text("#case-h"), "alert opens the case " + tag)
                    nb = p.locator(".rn-badge")
                    ok((nb.count() == 0 and n0 == 1) or (nb.count() == 1 and int(nb.inner_text()) == n0 - 1), "opened alert marked read " + tag)
                    p.click("#rn-bell"); p.wait_for_selector("#rn-off"); p.click("#rn-off"); p.wait_for_selector("#rn-on")
                    ok(True, "alerts turned off " + tag)
                ok(not errs, "no script errors (dept bell) " + tag, errs)
                ctx.close()
    # A link from an alert (#case=..&n=..) opens the case and marks the alert read
    ctx = br.new_context(viewport={"width": 390, "height": 900})
    ok(dept_signin(ctx) == 200, "signed in for the alert link")
    d = _sq.connect(WORK + "/test.db", timeout=10)
    nid = d.execute("SELECT id FROM notifications WHERE kind = 'DEPT_O_LATE' AND grievance_id = 'g24' AND read_at IS NULL").fetchone()
    d.close()
    p = ctx.new_page(); errs = []
    p.on("pageerror", lambda e: errs.append(str(e)))
    if nid:
        p.goto(B + "/dept#case=g24&n=" + nid[0]); p.wait_for_selector("#case-h"); p.wait_for_timeout(800)
        d = _sq.connect(WORK + "/test.db", timeout=10)
        ok(d.execute("SELECT read_at FROM notifications WHERE id = ?", (nid[0],)).fetchone()[0] is not None and "GRV-DEPT24" in p.inner_text("#case-h"), "alert link opens the case and marks it read")
        d.close()
    else:
        ok(False, "an unread overdue alert to open")
    ok(not errs, "no script errors (alert link)", errs)
    ctx.close()

    # ---------- Daily limit on filing complaints; blocked-alerts wording (grieviq-36) ----------
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx = br.new_context(viewport={"width": width, "height": 900}, color_scheme=theme)
                ctx.add_init_script("try{localStorage.setItem('%s','%s')}catch(e){}" % (LANGKEY, lang))
                p = ctx.new_page(); errs = []
                p.on("pageerror", lambda e: errs.append(str(e)))
                p.route("**/api/grievances/submit", lambda rt: rt.fulfill(status=429, headers={"content-type": "application/json", "retry-after": "3600"},
                        body='{"error":"limit","code":"DAILY_LIMIT","which":"PHONE","max":5}'))
                p.goto(B + "/submit?ward=lu-hazratganj"); p.wait_for_timeout(900)
                p.evaluate("""() => { document.getElementById('category').value = 'water-sanitation'; document.getElementById('localUnitId').value = 'lu-hazratganj'; }""")
                p.fill("#description", "Water pipe leaking near the park gate since Monday")
                p.fill("#phone", "9811100001"); p.check("#consent")
                p.wait_for_timeout(3100)
                p.click("#submitBtn"); p.wait_for_timeout(600)
                lim = p.locator("#limitError")
                txt = lim.inner_text() if lim.is_visible() else ""
                if lang == "en":
                    ok("You've filed 5 complaints today with this phone number" in txt and "limit is 5 a day" in txt, "limit message next to the button " + tag, txt or p.inner_text("#errorMessage"))
                else:
                    ok("आप आज इस फ़ोन नंबर से 5 शिकायतें" in txt, "limit message in Hindi " + tag, txt)
                ok(p.input_value("#description").startswith("Water pipe leaking") and p.input_value("#phone") == "9811100001" and p.is_enabled("#submitBtn"), "form keeps what was typed " + tag)
                ok(p.get_attribute("#limitError", "role") == "alert", "limit message is announced " + tag)
                ok(no_hscroll(p), "limit message: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS.replace("document.body", "document.getElementById('limitError')", 1)); ok(not bad, "limit message contrast AA " + tag, bad[:5])
                if width == 375: shot(p, "38-limit-" + tag, full=False) if False else p.locator("#limitError").screenshot(path=os.path.join(SHOTS, "38-limit-" + tag + ".png"))
                ok(not errs, "no script errors (limit) " + tag, errs)
                ctx.close()
    # blocked alerts: the bell says what to do (Incognito or site settings)
    ctx = br.new_context(viewport={"width": 1280, "height": 900})
    ok(dept_signin(ctx) == 200, "signed in for the blocked-alerts check")
    ctx.add_init_script("try{Object.defineProperty(Notification,'permission',{configurable:true,get:()=>'denied'})}catch(e){}")
    p = ctx.new_page(); errs = []
    p.on("pageerror", lambda e: errs.append(str(e)))
    p.goto(B + "/dept"); p.wait_for_selector("#rn-bell"); p.click("#rn-bell"); p.wait_for_selector("#rn-panel .rn-push")
    ok("Incognito" in p.inner_text("#rn-panel .rn-push") and "site settings" in p.inner_text("#rn-panel .rn-push"), "blocked alerts: Incognito and site-settings advice", p.inner_text("#rn-panel .rn-push"))
    ok(not errs, "no script errors (blocked alerts)", errs)
    ctx.close()

    # ---------- The agreement inside GrievIQ; rules for officers (grieviq-37) ----------
    def reset_rules():
        d = _sq.connect(WORK + "/test.db", timeout=10)
        d.execute("UPDATE dept_officers SET rules_version = NULL, rules_accepted_at = NULL WHERE email = 'ae.zone3@nic.in'"); d.commit(); d.close()
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme, lang, width)
                reset_rules()
                ctx = br.new_context(viewport={"width": width, "height": 900}, color_scheme=theme)
                ok(dept_signin(ctx) == 200, "signed in " + tag)
                ctx.add_init_script("try{localStorage.setItem('%s','%s')}catch(e){}" % (LANGKEY, lang))
                p = ctx.new_page(); errs = []
                p.on("pageerror", lambda e: errs.append(str(e)))
                p.goto(B + "/dept"); p.wait_for_selector("#rules-h")
                ok(p.locator("ol.rules li").count() == 7 and p.locator(".tiles").count() == 0, "rules shown first, no cases " + tag)
                ok(p.evaluate("document.activeElement.id") == "rules-h", "focus on the rules heading " + tag)
                txt = p.inner_text("main")
                if lang == "en":
                    ok("Rules for officers" in txt and "Never share your sign-in code" in txt and "security@grieviq.in" in txt, "rules wording " + tag, txt[:300])
                else:
                    ok("अधिकारियों के लिए नियम" in txt and not re.findall(r"dd\.\w+", txt), "rules in Hindi " + tag, txt[:300])
                p.click("#rules-go"); p.wait_for_timeout(150)
                ok(p.inner_text("#rules-err") != "" and p.get_attribute("#rules-agree", "aria-invalid") == "true" and p.evaluate("document.activeElement.id") == "rules-agree", "must tick the box; focus moves to it " + tag)
                ok(no_hscroll(p), "rules: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS); ok(not bad, "rules contrast AA " + tag, bad[:5])
                small = p.evaluate(TARGET_JS); ok(not small, "rules targets >= 24px " + tag, small[:5])
                shot(p, "39-rules-" + tag, full=(width == 375))
                p.check("#rules-agree")
                with p.expect_response(lambda r: "/api/dept/rules" in r.url and r.request.method == "POST"): p.click("#rules-go")
                p.wait_for_selector(".tiles")
                ok(p.locator("#rn-bell").count() == 1, "after accepting: cases and bell " + tag)
                wel = p.inner_text("#msg") if p.locator("#msg").count() else ""
                ok(("Welcome to GrievIQ, Ravi Kumar" in wel and "Lucknow" in wel) if lang == "en" else ("GrievIQ में आपका स्वागत है, Ravi Kumar" in wel), "officer welcomed on first acceptance " + tag, wel)
                p.click("#rules-link"); p.wait_for_selector("#rules-h")
                ok(("You accepted these rules (version 1) on" in p.inner_text("main")) if lang == "en" else ("आपने ये नियम (संस्करण 1)" in p.inner_text("main")), "footer link: rules with the date accepted " + tag)
                p.click("#rules-back"); p.wait_for_selector(".tiles")
                ok(not errs, "no script errors (rules) " + tag, errs)
                ctx.close()
    # Admin: the agreement to print, the signed copy, who accepted the rules
    d = _sq.connect(WORK + "/test.db", timeout=10)
    jal_id = d.execute("SELECT id FROM dept_offices WHERE name_en = 'Jal Kal Vibhag Lucknow'").fetchone()[0]
    d.close()
    for theme in (None, "light"):
        for lang in ("en", "hi"):
            tag = "%s-%s" % (theme or "dark", lang)
            ctx, p, errs = P("super@test.in", theme=theme, lang=lang)
            p.goto(B + "/admin-departments.html?area=lucknow"); p.wait_for_selector(".office")
            p.locator("article:has-text('Jal Kal Vibhag Lucknow') [data-officers]").click(); p.wait_for_selector(".op h5")
            ptxt = p.inner_text(".op")
            if lang == "en":
                ok("Download agreement to sign" in ptxt and "Signed copy on file:" in ptxt and "Download signed copy" in ptxt and "Replace the signed copy" in ptxt, "agreement panel: print, signed copy, replace " + tag, ptxt[:500])
                ok("Accepted the rules (v1)" in ptxt, "officer row: rules accepted " + tag)
            else:
                ok("हस्ताक्षर हेतु समझौता डाउनलोड करें" in ptxt and not re.findall(r"adm\.do_\w+", ptxt), "agreement panel in Hindi " + tag)
            bad = p.evaluate(CONTRAST_JS.replace("createTreeWalker(document.body", "createTreeWalker(document.querySelector('.op')", 1)); ok(not bad, "agreement panel contrast AA " + tag, bad[:5])
            if lang == "en" and theme is None:
                p.click("#op-copy-go"); p.wait_for_timeout(150)
                ok(p.get_attribute("#op-copy-file", "aria-invalid") == "true" and "Choose the file" in p.inner_text("#op-copy-err"), "upload: file needed " + tag)
                with p.expect_response(lambda r: "dept-agreement-copy" in r.url and r.request.method == "POST"):
                    p.set_input_files("#op-copy-file", files=[{"name": "signed.pdf", "mimeType": "application/pdf", "buffer": b"%PDF-1.4\n%test\n%%EOF\n"}]); p.click("#op-copy-go")
                p.wait_for_selector("#op-msg")
                ok("Signed copy uploaded." in p.inner_text("#op-msg") and "Signed copy on file: PDF" in p.inner_text(".op"), "signed PDF uploaded from the page " + tag)
                href = p.get_attribute("#op-print", "href")
                ok(href == "/admin-dept-agreement.html?office=" + jal_id and p.get_attribute("#op-print", "target") == "_blank", "print link opens the agreement in a new tab " + tag)
            p.locator(".op").screenshot(path=os.path.join(SHOTS, "40-agreement-panel-" + tag + ".png"))
            ok(not errs, "no script errors (agreement panel) " + tag, errs)
            ctx.close()
    ctx, p, errs = P("audit@test.in")
    p.goto(B + "/admin-departments.html?area=lucknow"); p.wait_for_selector(".office")
    p.locator("article:has-text('Jal Kal Vibhag Lucknow') [data-officers]").click(); p.wait_for_selector(".op h5")
    ok(p.locator("#op-copy-open").count() == 1 and p.locator("#op-copy-form").count() == 0, "auditor: can download the signed copy, can't upload")
    ctx.close()
    # The agreement to print
    for width in (1280, 375):
        ctx, p, errs = P("super@test.in", width=width)
        p.goto(B + "/admin-dept-agreement.html?office=" + jal_id); p.wait_for_selector("article.sheet h1")
        en = p.inner_text("article[lang=en]"); hi_t = p.inner_text("article[lang=hi]")
        ok("GrievIQ Department Participation Agreement" in en and "Jal Kal Vibhag Lucknow" in en and "Lucknow, Uttar Pradesh" in en and "Ravi Kumar" in en, "agreement filled in for the office, officers listed (%d)" % width, en[:300])
        ok("GrievIQ विभाग सहभागिता समझौता" in hi_t and "Ravi Kumar" in hi_t, "Hindi agreement on its own page (%d)" % width)
        ok("Asha Verma" not in en, "removed officers not listed (%d)" % width)
        ok(no_hscroll(p), "agreement page: no sideways scroll (%d)" % width)
        bad = p.evaluate(CONTRAST_JS); ok(not bad, "agreement page contrast AA (%d)" % width, bad[:5])
        ok(p.is_visible("#print"), "print button (%d)" % width)
        p.emulate_media(media="print"); p.wait_for_timeout(100)
        ok(not p.is_visible("#print") and not p.is_visible("#note"), "print: buttons and note hidden (%d)" % width)
        p.emulate_media(media="screen")
        shot(p, "41-agreement-%d" % width, full=False)
        if width == 1280: p.pdf(path=os.path.join(SHOTS, "41-agreement.pdf"), format="A4") if hasattr(p, "pdf") else None
        ok(not errs, "no script errors (agreement page %d)" % width, errs)
        ctx.close()
    ctx, p, errs = P("deo@test.in")
    p.goto(B + "/admin-dept-agreement.html?office=" + jal_id); p.wait_for_selector(".err")
    ok("can't open department agreements" in p.inner_text(".err"), "data entry operator: agreement page refused")
    ctx.close()

    # ---------- grieviq-38: page not found, thanks, welcome, admin buttons ----------
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx = br.new_context(viewport={"width": width, "height": 900}, color_scheme=theme)
                ctx.add_init_script("try{localStorage.setItem('%s','%s')}catch(e){}" % (LANGKEY, lang))
                p = ctx.new_page(); errs = []
                p.on("pageerror", lambda e: errs.append(str(e)))
                resp = p.goto(B + "/this-page-is-not-here"); p.wait_for_selector("h1")
                ok(resp.status == 404, "unknown address: 404 " + tag, resp.status)
                txt = p.inner_text("main")
                if lang == "en":
                    ok("Page not found" in txt and "error 404" in txt.lower(), "not-found wording " + tag, txt[:200])
                else:
                    ok("पृष्ठ नहीं मिला" in txt and not re.findall(r"nf\.\w+", txt), "not-found in Hindi " + tag, txt[:200])
                hrefs = p.eval_on_selector_all("main a", "a => a.map(x => x.getAttribute('href'))")
                ok(all(h in hrefs for h in ("/", "/status", "/submit", "/help")), "ways out: home, track, file, help " + tag, hrefs)
                ok(p.locator(".footer-thanks").count() == 1 and ("Thank you for using GrievIQ" in p.inner_text(".footer-thanks") if lang == "en" else "धन्यवाद" in p.inner_text(".footer-thanks")), "footer thank-you " + tag)
                ok(no_hscroll(p), "not-found: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS); ok(not bad, "not-found contrast AA " + tag, bad[:5])
                small = p.evaluate(TARGET_JS.replace("!e.closest('dd,", "!e.closest('.footer-links, dd,", 1)); ok(not small, "not-found targets >= 24px " + tag, small[:5])
                shot(p, "42-not-found-" + tag, full=(width == 375))
                ok(not errs, "no script errors (not found %s)" % tag, errs)
                ctx.close()
    ctx = br.new_context(); p = ctx.new_page()
    p.goto(B + "/submit"); p.wait_for_selector("form")
    ok(p.locator("p.done-thanks[data-i18n='submit.done_thanks']").count() == 1 and "Thank you for speaking up" in p.locator("p.done-thanks").text_content(), "thank-you line on the complaint-filed screen")
    ctx.close()

    # Rep welcome card: once per person per device
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx = br.new_context(viewport={"width": width, "height": 900}, color_scheme=theme)
                ctx.add_cookies([{"name": "test_email", "value": "r1@x.in", "url": B}])
                ctx.add_init_script("try{localStorage.setItem('%s','%s')}catch(e){}" % (LANGKEY, lang))
                p = ctx.new_page(); errs = []
                p.on("pageerror", lambda e: errs.append(str(e)))
                p.goto(B + "/rep"); p.wait_for_selector(".welcome-card", timeout=15000)
                txt = p.inner_text(".welcome-card")
                if lang == "en":
                    ok("Congratulations, and welcome to GrievIQ." in txt and "a better India" in txt, "rep welcome wording " + tag, txt[:300])
                else:
                    ok("बधाई हो" in txt and not re.findall(r"rep\.wel_\w+", txt), "rep welcome in Hindi " + tag, txt[:300])
                ok(no_hscroll(p), "welcome: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS.replace("createTreeWalker(document.body", "createTreeWalker(document.querySelector('.welcome-card')", 1)); ok(not bad, "welcome card contrast AA " + tag, bad[:5])
                shot(p, "43-rep-welcome-" + tag, full=False)
                p.click("#welcome-go"); p.wait_for_timeout(200)
                ok(p.locator(".welcome-card").count() == 0, "Get started closes the card " + tag)
                ok(p.evaluate("document.activeElement.dataset.tab") == "cases", "focus moves to the Cases tab " + tag)
                p.reload(); p.wait_for_selector(".tab-btn", timeout=15000); p.wait_for_timeout(500)
                ok(p.locator(".welcome-card").count() == 0, "not shown again on this device " + tag)
                ok(not errs, "no script errors (welcome %s)" % tag, errs)
                ctx.close()

    # Admin primary buttons: readable in light and dark
    for theme in ("light", "dark"):
        for pg in ("/admin-departments.html", "/admin-staff.html", "/admin-cases.html", "/admin-dashboard.html"):
            ctx, p, errs = P("super@test.in", theme=theme)
            p.goto(B + pg); p.wait_for_load_state("networkidle"); p.wait_for_timeout(300)
            res = p.evaluate(r"""() => {
              function rgb(s){const m=s.match(/rgba?\(([^)]+)\)/); const p=m[1].split(',').map(Number); return p;}
              function lum(c){const f=v=>{v/=255;return v<=0.03928?v/12.92:Math.pow((v+0.055)/1.055,2.4)};return 0.2126*f(c[0])+0.7152*f(c[1])+0.0722*f(c[2]);}
              const out=[];
              document.querySelectorAll('button.primary, .admin-nav a[aria-current=page], [aria-pressed=true], [aria-selected=true]').forEach(b=>{
                const cs=getComputedStyle(b); const r=b.getBoundingClientRect(); if(r.width<2) return;
                const bgc=rgb(cs.backgroundColor); if(bgc.length>3 && bgc[3]===0) return;
                const L1=lum(rgb(cs.color)),L2=lum(bgc); const ratio=(Math.max(L1,L2)+0.05)/(Math.min(L1,L2)+0.05);
                out.push([b.textContent.trim().slice(0,30), ratio]);
              });
              return out;
            }""")
            low = [x for x in res if x[1] < 4.5]
            ok(not low, "filled admin buttons >= 4.5:1 %s %s" % (theme, pg), low[:5])
            ctx.close()

    # ---------- grieviq-39: For offices page, Prices & site details, Pilot requests ----------
    import urllib.request as _ur
    def api(method, path, body=None, email="super@test.in"):
        req = _ur.Request(B + path, method=method, data=json.dumps(body).encode() if body is not None else None, headers={"content-type": "application/json", "test-email": email})
        try:
            with _ur.urlopen(req) as r: return r.status, json.loads(r.read() or b"{}")
        except _ur.HTTPError as e: return e.code, json.loads(e.read() or b"{}")
    def live(): return api("GET", "/api/admin/site-settings")[1]["live"]["data"]
    def publish(data, note="ui test"):
        s1, b1 = api("POST", "/api/admin/site-settings", {"action": "save_draft", "data": data}); s2, b2 = api("POST", "/api/admin/site-settings", {"action": "publish", "note": note})
        return s1, s2
    START = live()
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx = br.new_context(viewport={"width": width, "height": 900}, color_scheme=theme)
                ctx.add_init_script("try{localStorage.setItem('%s','%s')}catch(e){}" % (LANGKEY, lang))
                p = ctx.new_page(); errs = []
                p.on("pageerror", lambda e: errs.append(str(e)))
                resp = p.goto(B + "/for-offices"); p.wait_for_selector(".price-table tbody tr")
                ok(resp.status == 200, "for offices page opens " + tag)
                txt = p.inner_text("main"); allt = p.evaluate("() => { const c = document.body.cloneNode(true); c.querySelectorAll('script,style').forEach(x => x.remove()); return c.textContent; }")
                if lang == "en":
                    ok("tracked to resolution" in p.inner_text("h1") and "Gram Pradhan office" in txt and "₹999" in txt and "₹1,179 including GST" in txt and "₹42,000" in txt and "Request a free pilot" in txt and "90-day" not in txt, "headline, Pradhan price with GST, department price " + tag, txt[:300])
                    ok("Amit Singh" in txt and "Sector-150, Noida" in txt and "within 2 working days" in txt and "support@grieviq.in" in txt, "Grievance Officer, address, callback, contact " + tag)
                else:
                    ok("प्रमाण सहित" in p.inner_text("h1") and "ग्राम प्रधान कार्यालय" in txt and "GST सहित ₹1,179" in txt and "निःशुल्क पायलट माँगें" in txt and "90 दिन का निःशुल्क" not in txt and not re.findall(r"lp\.\w+", allt), "Hindi page, no missing text " + tag, txt[:300])
                ok("GeM" not in allt and "data centres in India" not in allt and "भारत स्थित डेटा" not in allt and "security audit by" not in allt, "claims held back " + tag)
                ok(p.locator(".price-table tbody tr").count() == 6, "six plans listed " + tag)
                ok(no_hscroll(p), "for offices: no sideways scroll " + tag)
                for d in p.locator("details").all(): d.evaluate("e => e.open = true")
                bad = p.evaluate(CONTRAST_JS); ok(not bad, "for offices contrast AA " + tag, bad[:6])
                small = p.evaluate(TARGET_JS.replace("!e.closest('dd,", "!e.closest('.footer-links, .hp, label, dd,", 1)); ok(not small, "for offices targets >= 24px " + tag, small[:6])
                shot(p, "44-for-offices-" + tag, full=True)
                ok(not errs, "no script errors (for offices %s)" % tag, errs)
                ctx.close()

    # The pilot request form: errors, then a sent request
    for lang in ("en", "hi"):
        ctx = br.new_context(viewport={"width": 1280, "height": 900})
        ctx.add_init_script("try{localStorage.setItem('%s','%s')}catch(e){}" % (LANGKEY, lang))
        p = ctx.new_page(); errs = []
        p.on("pageerror", lambda e: errs.append(str(e)))
        p.goto(B + "/for-offices#pilot"); p.wait_for_selector(".price-table")
        p.wait_for_timeout(2100)
        p.click("#pf-send"); p.wait_for_timeout(150)
        ok(p.evaluate("document.activeElement.id") == "err-box" and p.locator("#err-box li").count() == 5, "empty form: 5 errors listed, focus on the list " + lang)
        ok(p.get_attribute("#pf-name", "aria-invalid") == "true" and p.get_attribute("#pf-mobile", "aria-invalid") == "true" and "pe-mobile" in (p.get_attribute("#pf-mobile", "aria-describedby") or ""), "fields marked and described " + lang)
        if lang == "hi": ok("अपना नाम लिखें।" in p.inner_text("#err-box"), "errors in Hindi")
        bad = p.evaluate(CONTRAST_JS.replace("createTreeWalker(document.body", "createTreeWalker(document.getElementById('pilot')", 1)); ok(not bad, "form errors contrast AA " + lang, bad[:5])
        p.fill("#pf-name", "Sita Devi"); p.check("input[name=office][value=PRADHAN]"); p.fill("#pf-place", "Gram Panchayat Malihabad, Lucknow")
        p.fill("#pf-mobile", "98765 1234")
        p.click("#pf-send"); p.wait_for_timeout(150)
        ok(p.locator("#err-box li").count() == 2 and p.get_attribute("#pf-name", "aria-invalid") is None, "fixed fields cleared; mobile format and consent left " + lang)
        p.fill("#pf-mobile", "98765 1234" + ("5" if lang == "en" else "6")); p.check("#pf-consent")
        with p.expect_response(lambda r: "/api/pilot-request" in r.url): p.click("#pf-send")
        p.wait_for_selector("#done-h")
        ok(p.evaluate("document.activeElement.id") == "done-h" and ("within 2 working days" in p.inner_text(".done") if lang == "en" else "2 कार्य-दिवसों के भीतर" in p.inner_text(".done")), "request sent; thank-you with callback time " + lang)
        shot(p, "45-pilot-sent-" + lang, full=False)
        ok(not errs, "no script errors (pilot form %s)" % lang, errs)
        ctx.close()

    # Home page link and card; footer link on citizen pages
    ctx, p, errs = P("x@x.in")
    p.goto(B + "/"); p.wait_for_selector(".home-more .feedback-card"); p.wait_for_timeout(300)
    ok(p.locator("#offices-card").count() == 0 and p.locator("main a[href='/for-offices'], .home-grid a[href='/for-offices'], .home-more a[href='/for-offices'], .site-nav a[href='/for-offices']").count() == 0, "citizen home: nothing for offices except the footer (owner decision)")
    ok(p.locator(".footer-links a[href='/for-offices']").count() == 1, "footer link to For offices")
    p.goto(B + "/privacy"); p.wait_for_timeout(600)
    ok("Amit Singh" in p.inner_text("body") and "Sector-150, Noida" in p.inner_text("body"), "privacy page: Grievance Officer and address")
    p.goto(B + "/help"); p.wait_for_timeout(600)
    ok("Amit Singh" in p.inner_text("body"), "help page: Grievance Officer")
    ctx.close()
    # Changing the details in settings changes those pages
    d = json.loads(json.dumps(START)); d["company"]["goName"] = "Neha Verma"; d["company"]["goEmail"] = "dpo@grieviq.in"; d["company"]["address"] = "12 Test Road, Lucknow, 226001"
    ok(publish(d) == (200, 200), "details changed and published")
    for lang in ("en", "hi"):
        ctx, p, errs = P("x@x.in", lang=lang)
        p.goto(B + "/privacy"); p.wait_for_function("document.body.innerText.indexOf('Neha Verma') >= 0", timeout=5000)
        t = p.inner_text("body")
        ok("Neha Verma" in t and "12 Test Road" in t and "dpo@grieviq.in" in t and p.locator("a[href='mailto:dpo@grieviq.in']").count() >= 1, "privacy page follows the settings " + lang)
        if lang == "hi":
            p.click("[data-lang-toggle]"); p.wait_for_timeout(300)
            p.wait_for_timeout(200); ok("Neha Verma" in p.inner_text("body"), "still filled in after switching language")
        ctx.close()
    ctx, p, errs = P("super@test.in")
    p.goto(B + "/admin-dept-agreement.html?office=" + jal_id); p.wait_for_selector("article.sheet h1")
    ok("12 Test Road, Lucknow, 226001" in p.inner_text("article[lang=en]"), "agreement shows the registered address from settings")
    ctx.close()
    ok(publish(START) == (200, 200), "starting details published again")

    # Admin: Prices & site details
    for theme in ("dark", "light"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx, p, errs = P("super@test.in", width=width, theme=theme, lang=lang)
                p.goto(B + "/admin-site-settings.html"); p.wait_for_selector("[role=tab]")
                ok(p.locator("[role=tab]").count() == 9 and p.get_attribute("#tab-plans", "aria-selected") == "true", "nine tabs, Plans open " + tag)
                for t in ("plans", "discounts", "company", "claims", "history"):
                    p.click("#tab-" + t); p.wait_for_timeout(100)
                    ok(no_hscroll(p), "settings %s: no sideways scroll %s" % (t, tag))
                    bad = p.evaluate(CONTRAST_JS); ok(not bad, "settings %s contrast AA %s" % (t, tag), bad[:5])
                    small = p.evaluate(TARGET_JS); ok(not small, "settings %s targets %s" % (t, tag), small[:5])
                    if t in ("plans", "discounts"): shot(p, "46-settings-%s-%s" % (t, tag), full=False)
                if lang == "hi": ok(not re.findall(r"adm\.\w+", p.inner_text("main")), "settings in Hindi, nothing missing " + tag)
                ok(not errs, "no script errors (settings %s)" % tag, errs)
                ctx.close()
    # The flow: change, check, save, preview, publish, bring back
    ctx, p, errs = P("super@test.in")
    p.goto(B + "/admin-site-settings.html"); p.wait_for_selector("[role=tab]")
    p.focus("#tab-plans"); p.keyboard.press("ArrowRight")
    ok(p.evaluate("document.activeElement.id") == "tab-notices" and p.get_attribute("#tab-notices", "aria-selected") == "true", "arrow keys move between tabs")
    p.click("#tab-plans")
    p.fill("#f-plans-0-monthly", "1199"); p.wait_for_timeout(100)
    ok("₹1,199 a month (₹1,415 with GST)" in p.inner_text(".card .today >> nth=0") and p.locator(".unsaved").count() == 1, "price worked out as you type; marked not saved")
    p.fill("#f-plans-1-nameEn", "")
    p.click("#ss-save"); p.wait_for_selector("#ss-errors:not([hidden])")
    ok(p.evaluate("document.activeElement.id") == "ss-errors" and p.get_attribute("#f-plans-1-nameEn", "aria-invalid") == "true" and "(1)" in p.inner_text("#tab-plans"), "error list focused, field marked, tab shows the count")
    p.click("#ss-errors a"); ok(p.evaluate("document.activeElement.id") == "f-plans-1-nameEn", "error link goes to the field")
    p.fill("#f-plans-1-nameEn", "Ward Councillor office")
    with p.expect_response(lambda r: "/api/admin/site-settings" in r.url and r.request.method == "POST"): p.click("#ss-save")
    p.wait_for_selector("#ss-publish")
    ok("Draft saved: 1 change" in p.inner_text("#ss-msg") and p.evaluate("document.activeElement.id") == "ss-msg", "draft saved message, focused")
    ok(api("GET", "/api/site-info")[1]["plans"][0]["monthly"] == 999, "draft not public yet")
    with ctx.expect_page() as pv: p.click("#ss-preview")
    pp = pv.value; pp.wait_for_selector(".price-table")
    ok(pp.is_visible("#preview-bar") and "₹1,199" in pp.inner_text(".price-table"), "preview shows the draft with a notice")
    pp.close()
    p.fill("#ss-note", "Pradhan price test")
    with p.expect_response(lambda r: "/api/admin/site-settings" in r.url and r.request.method == "POST"): p.click("#ss-publish")
    p.wait_for_timeout(300)
    ok("Published as version" in p.inner_text("#ss-msg") and api("GET", "/api/site-info")[1]["plans"][0]["monthly"] == 1199, "published; public price changed")
    p.click("#tab-history"); p.wait_for_selector("table")
    ok("Pradhan price test" in p.inner_text("table") and p.locator("[data-act=restore]").count() >= 1, "history lists the version with its note")
    p.locator("[data-act=view]").nth(1).click(); p.wait_for_selector("#ss-back")
    ok(p.locator("#f-plans-0-monthly").is_disabled() and "read only" in p.inner_text("#ss-status"), "an earlier version opens read only")
    with p.expect_response(lambda r: "/api/admin/site-settings" in r.url and r.request.method == "POST"): p.click("#ss-restore-this")
    p.wait_for_timeout(300)
    ok("is live again" in p.inner_text("#ss-msg") and api("GET", "/api/site-info")[1]["plans"][0]["monthly"] == 999, "earlier version brought back")
    ok(not errs, "no script errors (settings flow)", errs)
    ctx.close()
    # Preview signed out: told, public page shown
    ctx = br.new_context(); p = ctx.new_page(); p.goto(B + "/for-offices?preview=draft"); p.wait_for_selector(".price-table")
    ok(p.is_visible("#preview-bar") and "isn't available" in p.inner_text("#preview-bar"), "preview needs the super admin")
    ctx.close()
    # Read only for others
    ctx, p, errs = P("audit@test.in")
    p.goto(B + "/admin-site-settings.html"); p.wait_for_selector("[role=tab]")
    ok(p.is_visible("#readonly-tag") and p.locator("#ss-save").count() == 0 and p.locator("#f-plans-0-monthly").is_disabled() and p.locator("[data-act=addplan]").count() == 0, "auditor: view only")
    bad = p.evaluate(CONTRAST_JS); ok(not bad, "view-only settings contrast AA", bad[:5])
    ctx.close()
    # Pilot requests
    for who in ("ops@test.in", "audit@test.in"):
        for theme in ("dark", "light"):
            ctx, p, errs = P(who, theme=theme)
            p.goto(B + "/admin-pilot-requests.html"); p.wait_for_selector(".pr-card")
            t = p.inner_text("main")
            if who.startswith("ops"):
                ok("Sita Devi" in t and p.locator("a[href^='tel:+91']").count() >= 1 and p.locator("select").count() >= 1, "operations admin: full details and status " + theme)
            else:
                ok("XXXXXX" in t and p.locator("select").count() == 0 and p.is_visible("#readonly-tag"), "auditor: numbers partly hidden, view only " + theme)
            bad = p.evaluate(CONTRAST_JS); ok(not bad, "pilot requests contrast AA %s %s" % (who, theme), bad[:5])
            ok(no_hscroll(p), "pilot requests: no sideways scroll")
            shot(p, "47-pilot-requests-%s-%s" % (who[:3], theme), full=False)
            ok(not errs, "no script errors (pilot requests)", errs)
            ctx.close()
    ctx, p, errs = P("ops@test.in")
    p.goto(B + "/admin-pilot-requests.html"); p.wait_for_selector(".pr-card")
    first = p.locator(".pr-card").first
    first.locator("select").select_option("CONTACTED"); first.locator("input[type=text]").fill("Demo booked")
    with p.expect_response(lambda r: "/api/admin/pilot-requests" in r.url and r.request.method == "POST"): first.locator("[data-save]").click()
    p.wait_for_timeout(300)
    ok("Saved." in p.inner_text("main") and "Demo booked" in p.inner_text("main"), "status and note saved")
    ctx, p, errs = P("deo@test.in")
    p.goto(B + "/admin-pilot-requests.html"); p.wait_for_timeout(800)
    ok("does not have access" in p.inner_text("main"), "data entry operator: no access to pilot requests")
    ctx.close()

    # ---------- grieviq-40: guides (SOPs) inside the apps, positioning line ----------
    DLG_CONTRAST = CONTRAST_JS.replace("createTreeWalker(document.body", "createTreeWalker(document.querySelector('.gq-drawer')", 1)
    DLG_TARGETS = r"""() => Array.from(document.querySelectorAll('.gq-drawer button, .gq-drawer a')).filter(e => { const r = e.getBoundingClientRect(); return r.width > 0 && (r.height < 24 || r.width < 24); }).map(e => e.textContent.trim().slice(0, 30))"""
    def sop_ids(p): return p.eval_on_selector_all(".gq-drawer .gq-sop", "els => els.map(e => e.id)")
    # Citizen guide page
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx = br.new_context(viewport={"width": width, "height": 900}, color_scheme=theme)
                ctx.add_init_script("try{localStorage.setItem('%s','%s')}catch(e){}" % (LANGKEY, lang))
                p = ctx.new_page(); errs = []
                p.on("pageerror", lambda e: errs.append(str(e)))
                resp = p.goto(B + "/guide"); p.wait_for_selector("#gq-page .gq-sop")
                ok(resp.status == 200, "guide page opens " + tag)
                txt = p.inner_text("main")
                if lang == "en":
                    ok("How to use GrievIQ" in p.inner_text("h1") and "File a complaint" in txt and "How a complaint moves" in txt and "112" in txt, "citizen guide in English " + tag, txt[:200])
                else:
                    ok("GrievIQ का उपयोग कैसे करें" in p.inner_text("h1") and "शिकायत दर्ज करें" in txt and not re.findall(r"(guide|common)\.\w+", txt), "citizen guide in Hindi, nothing missing " + tag, txt[:200])
                ok(p.locator("#gq-page .gq-sop").count() == 6 and p.locator("#gq-page ol.gq-flow li").count() == 5, "six guides and the five-step journey " + tag)
                ok(p.locator("#gq-page strong").count() > 5 and p.locator("#gq-page script").count() == 0, "bold words drawn, nothing else " + tag)
                ok(no_hscroll(p), "guide page: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS); ok(not bad, "guide page contrast AA " + tag, bad[:6])
                small = p.evaluate(TARGET_JS.replace("!e.closest('dd,", "!e.closest('.footer-links, dd,", 1)); ok(not small, "guide page targets >= 24px " + tag, small[:6])
                if width == 375 or lang == "en": shot(p, "48-guide-page-" + tag, full=(width == 375))
                ok(not errs, "no script errors (guide page %s)" % tag, errs)
                ctx.close()
    ctx, p, errs = P("x@x.in")
    p.goto(B + "/guide#gq-track"); p.wait_for_selector("#gq-track"); p.wait_for_timeout(200)
    ok(p.evaluate("document.activeElement.id") == "gq-track", "link to a part of the guide focuses it")
    p.click("a[data-gq-jump=rate]"); ok(p.evaluate("document.activeElement.id") == "gq-rate" and p.evaluate("location.hash") == "#gq-rate", "contents link moves focus")
    p.evaluate("() => { window.__printed = 0; window.print = () => { window.__printed++; }; }"); p.click("#gq-print")
    ok(p.evaluate("window.__printed") == 1, "print / save as PDF")
    p.click("[data-lang-toggle]"); p.wait_for_timeout(200)
    ok("शिकायत दर्ज करें" in p.inner_text("#gq-page") and "File a complaint" not in p.inner_text("#gq-page"), "switching language redraws the guide")
    p.click("[data-lang-toggle]"); p.wait_for_timeout(200)
    p.goto(B + "/"); p.wait_for_timeout(500)
    ok(p.locator(".footer-links a[href='/guide']").count() == 1, "footer link to the guide")
    p.goto(B + "/submit"); p.wait_for_timeout(500)
    ok(p.get_attribute(".guide-link a", "href") == "/guide#gq-file", "file page: how filing works")
    p.goto(B + "/status"); p.wait_for_timeout(500)
    ok(p.get_attribute(".guide-link a", "href") == "/guide#gq-track", "track page: how tracking works")
    for pg in ("/submit", "/status"):
        for th in ("light", "dark"):
            c2 = br.new_context(viewport={"width": 375, "height": 900}, color_scheme=th); p2 = c2.new_page(); p2.goto(B + pg); p2.wait_for_timeout(600)
            bad = p2.evaluate(CONTRAST_JS.replace("createTreeWalker(document.body", "createTreeWalker(document.querySelector('.guide-link')", 1)); ok(not bad, "how-it-works link contrast %s %s" % (pg, th), bad)
            c2.close()
    p.goto(B + "/for-offices"); p.wait_for_selector(".eyebrow")
    ok(p.inner_text(".eyebrow").lower() in ("constituency service & grievance redressal platform", "क्षेत्र सेवा एवं शिकायत निवारण मंच") and ("Constituency Service" in p.title() or "क्षेत्र सेवा" in p.title()), "For offices: the positioning line", p.inner_text(".eyebrow"))
    ctx.close()

    # Representative console: Help panel
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx = br.new_context(viewport={"width": width, "height": 900}, color_scheme=theme)
                ctx.add_cookies([{"name": "test_email", "value": "r1@x.in", "url": B}])
                ctx.add_init_script("try{localStorage.setItem('%s','%s')}catch(e){}" % (LANGKEY, lang))
                p = ctx.new_page(); errs = []
                p.on("pageerror", lambda e: errs.append(str(e)))
                p.goto(B + "/rep"); p.wait_for_selector("#help-btn", timeout=15000)
                sub = p.inner_text(".console-subtitle")
                ok(sub == ("Constituency Service & Grievance Redressal Platform" if lang == "en" else "क्षेत्र सेवा एवं शिकायत निवारण मंच"), "console subtitle: positioning line " + tag, sub)
                ok(p.locator("#welcome-guide").count() == 1, "welcome card links to the guide " + tag)
                ok(no_hscroll(p), "rep header with Help: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS.replace("createTreeWalker(document.body", "createTreeWalker(document.querySelector('.console-header')", 1)); ok(not bad, "rep header contrast " + tag, bad)
                bad = p.evaluate(CONTRAST_JS.replace("createTreeWalker(document.body", "createTreeWalker(document.querySelector('.welcome-card')", 1)); ok(not bad, "welcome card with guide link contrast " + tag, bad)
                p.click("#help-btn"); p.wait_for_selector(".gq-drawer")
                ok(p.get_attribute(".gq-drawer", "role") == "dialog" and p.get_attribute(".gq-drawer", "aria-modal") == "true" and p.evaluate("document.activeElement.id") == "gq-dtitle", "Help opens as a dialog, focus on its title " + tag)
                ok(sop_ids(p) == ["gqd-signin", "gqd-daily", "gqd-forward", "gqd-fieldcheck", "gqd-resolve", "gqd-team", "gqd-reports", "gqd-messages", "gqd-billing", "gqd-notices", "gqd-share", "gqd-takeover"], "representative sees all twelve guides " + tag, sop_ids(p))
                ok(p.evaluate("document.getElementById('app').inert") is True, "the console behind is inert " + tag)
                if lang == "hi": ok(not re.findall(r"guide\.\w+", p.inner_text(".gq-drawer")) and "सहायता" in p.inner_text("#help-btn"), "Help in Hindi " + tag)
                ok(p.evaluate("document.documentElement.scrollWidth <= window.innerWidth + 1") and p.evaluate("(() => { const d = document.querySelector('.gq-dbody'); return d.scrollWidth <= d.clientWidth + 1; })()"), "Help panel: no sideways scroll " + tag)
                bad = p.evaluate(DLG_CONTRAST); ok(not bad, "Help panel contrast AA " + tag, bad[:6])
                ok(not p.evaluate(DLG_TARGETS), "Help panel targets >= 24px " + tag, p.evaluate(DLG_TARGETS))
                shot(p, "49-rep-help-" + tag, full=False)
                p.keyboard.press("Shift+Tab"); ok(p.evaluate("document.activeElement.closest('.gq-drawer') !== null"), "Shift+Tab stays in the panel " + tag)
                p.keyboard.press("Escape"); p.wait_for_timeout(100)
                ok(p.locator(".gq-drawer").count() == 0 and p.evaluate("document.activeElement.id") == "help-btn" and p.evaluate("document.getElementById('app').inert") is False, "Esc closes; focus back on Help " + tag)
                ok(not errs, "no script errors (rep help %s)" % tag, errs)
                ctx.close()
    # Contextual links and role filtering
    ctx = br.new_context(viewport={"width": 1280, "height": 900}); ctx.add_cookies([{"name": "test_email", "value": "r1@x.in", "url": B}])
    p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
    p.goto(B + "/rep"); p.wait_for_selector("#help-btn", timeout=15000)
    p.click("#welcome-guide"); p.wait_for_selector(".gq-drawer")
    ok(p.evaluate("document.activeElement.id") == "gq-dtitle", "welcome guide link opens the guide"); p.keyboard.press("Escape")
    ok(p.evaluate("document.activeElement.id") == "welcome-guide", "focus back on the welcome link")
    p.evaluate("GIQGuide.open('rep', { focus: 'forward' })"); p.wait_for_selector(".gq-drawer")
    ok(p.evaluate("document.activeElement.id") == "gqd-forward", "a 'how this works' link opens the guide at that part")
    p.click(".gq-drawer [data-gq-jump=team]"); ok(p.evaluate("document.activeElement.id") == "gqd-team", "contents inside the panel move focus")
    p.evaluate("() => { window.__cls = ''; window.print = () => { window.__cls = document.documentElement.className; }; }"); p.click("[data-gq-print]")
    ok("gq-print" in p.evaluate("window.__cls"), "Print prints only the guide")
    p.evaluate("document.documentElement.classList.add('gq-print')"); p.emulate_media(media="print")
    ok(p.evaluate("getComputedStyle(document.getElementById('app')).display") == "none" and p.evaluate("getComputedStyle(document.querySelector('.gq-dbody')).overflowY") == "visible", "printout holds only the guide")
    p.emulate_media(media="screen"); p.evaluate("document.documentElement.classList.remove('gq-print')")
    p.click("[data-gq-close]"); p.wait_for_timeout(100)
    ok(p.locator(".gq-drawer").count() == 0, "Close button closes")
    ctxs = p.locator("[data-guide-focus]").count()
    if ctxs:
        b = p.locator("[data-guide-focus]").first; f = b.get_attribute("data-guide-focus"); b.click(); p.wait_for_selector(".gq-drawer")
        ok(p.evaluate("document.activeElement.id") == "gqd-" + f, "the link next to the action opens its guide (%s)" % f)
        p.keyboard.press("Escape")
    ok(not errs, "no script errors (rep contextual)", errs)
    ctx.close()
    ctx = br.new_context(viewport={"width": 375, "height": 900}); ctx.add_cookies([{"name": "test_email", "value": "fw@x.in", "url": B}])
    p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
    p.goto(B + "/rep"); p.wait_for_selector("#help-btn", timeout=15000)
    p.click("#help-btn"); p.wait_for_selector(".gq-drawer")
    ok(sop_ids(p) == ["gqd-signin", "gqd-fieldcheck"], "field worker sees only sign-in and field check", sop_ids(p))
    shot(p, "50-fw-help", full=False)
    ok(not errs, "no script errors (field worker help)", errs)
    ctx.close()

    # Department dashboard
    for theme in ("light", "dark"):
        ctx = br.new_context(viewport={"width": 375, "height": 900}, color_scheme=theme)
        ok(dept_signin(ctx) == 200, "dept signed in " + theme)
        p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
        p.goto(B + "/dept"); p.wait_for_selector(".tiles")
        p.click("#guide-link"); p.wait_for_selector(".gq-drawer")
        ok(sop_ids(p) == ["gqd-signin", "gqd-dashboard", "gqd-reply", "gqd-sentback", "gqd-officers", "gqd-performance", "gqd-billing", "gqd-notices"] and p.locator(".gq-drawer ol.gq-flow li").count() == 5, "department guide with its journey " + theme, sop_ids(p))
        bad = p.evaluate(DLG_CONTRAST); ok(not bad, "dept Help contrast AA " + theme, bad[:6])
        shot(p, "51-dept-help-" + theme, full=False)
        p.keyboard.press("Escape"); ok(p.evaluate("document.activeElement.id") == "guide-link", "focus back on the guide link " + theme)
        p.goto(B + "/dept#case=g24"); p.wait_for_selector("#case-h"); p.wait_for_timeout(300)
        if p.locator("#guide-reply").count():
            bad = p.evaluate(CONTRAST_JS.replace("createTreeWalker(document.body", "createTreeWalker(document.querySelector('.guide-row')", 1)); ok(not bad, "how to reply link contrast " + theme, bad)
            p.click("#guide-reply"); p.wait_for_selector(".gq-drawer")
            ok(p.evaluate("document.activeElement.id") == "gqd-reply", "How to reply opens that part " + theme)
        ok(not errs, "no script errors (dept help %s)" % theme, errs)
        ctx.close()

    # Admin console: Help shows the parts for the person's role
    for who, want in (("super@test.in", 21), ("deo@test.in", ["gqd-data", "gqd-addrep", "gqd-seats", "gqd-handover"]), ("audit@test.in", ["gqd-bills", "gqd-subs", "gqd-notices", "gqd-gov", "gqd-staff", "gqd-audit", "gqd-retention"])):
        for theme in ("dark", "light"):
            ctx, p, errs = P(who, theme=theme, width=375 if who.startswith("deo") else 1280)
            p.goto(B + "/admin-dashboard.html"); p.wait_for_selector("#adm-help", state="attached")
            if not p.is_visible("#adm-help"): p.click("#adm-menu-btn")
            p.click("#adm-help"); p.wait_for_selector(".gq-drawer")
            ids = sop_ids(p)
            ok((len(ids) == want) if isinstance(want, int) else (ids == want), "admin Help for %s %s" % (who, theme), ids)
            bad = p.evaluate(DLG_CONTRAST); ok(not bad, "admin Help contrast AA %s %s" % (who, theme), bad[:6])
            if who.startswith("super"): ok("underline" in p.evaluate("getComputedStyle(document.querySelector('.gq-toc a')).textDecorationLine"), "admin Help: contents look like links " + theme)
            if who.startswith("super"): shot(p, "52-admin-help-" + theme, full=False)
            p.keyboard.press("Escape"); ok(p.evaluate("document.activeElement.id") == "adm-help", "admin: focus back on Help " + theme)
            ok(not errs, "no script errors (admin help %s)" % who, errs)
            ctx.close()

    # ---------- grieviq-60: the Rep console never shows a blank page ----------
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx = br.new_context(viewport={"width": width, "height": 900}, color_scheme=theme)
                ctx.add_cookies([{"name": "test_email", "value": "r1@x.in", "url": B}])
                ctx.add_init_script("try{localStorage.setItem('%s','%s')}catch(e){}" % (LANGKEY, lang))
                p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
                fails = {"n": 1}
                def _route(route):
                    if fails["n"] > 0:
                        fails["n"] -= 1; route.fulfill(status=500, body="Internal error")
                    else: route.continue_()
                p.route("**/api/grievances?*", _route); p.route("**/api/grievances", _route)
                p.goto(B + "/rep"); p.wait_for_selector("#load-retry", timeout=15000)
                txt = p.inner_text("#app")
                ok(("couldn't load" in txt and "support@grieviq.in" in txt) if lang == "en" else ("लोड नहीं हो सका" in txt and not re.findall(r"rep\.\w+", txt)), "rep: a failed load says so, with Try again " + tag, txt[:200])
                ok(p.get_attribute(".notice", "role") == "alert" and p.evaluate("document.activeElement.id") == "load-retry", "rep: failure is announced and Try again has focus " + tag)
                bad = p.evaluate(CONTRAST_JS.replace("createTreeWalker(document.body", "createTreeWalker(document.querySelector('.notice')", 1)); ok(not bad, "rep: failure message contrast AA " + tag, bad)
                ok(p.evaluate("[...document.querySelectorAll('.notice-actions button')].every(b => b.getBoundingClientRect().height >= 44)"), "rep: failure buttons 44px " + tag)
                ok(no_hscroll(p), "rep: failure message, no sideways scroll " + tag)
                if lang == "en" or width == 375: shot(p, "60-rep-load-failed-" + tag, full=False)
                p.click("#load-retry"); p.wait_for_selector("#help-btn", timeout=15000)
                ok(p.locator("#load-retry").count() == 0, "rep: Try again loads the console " + tag)
                ok(not errs, "no script errors (rep load failure %s)" % tag, errs)
                ctx.close()
    # A representative holding 6 ward offices sees the console (live bug of 23 Oct: compound SELECT limit)
    import sqlite3 as _sq60
    _c60 = _sq60.connect(WORK + "/test.db", timeout=10)
    for _i in range(1, 7):
        _c60.execute("INSERT OR IGNORE INTO local_units (id,name,unit_type,mla_constituency_id,rep_email,area_id) VALUES (?,?,?,?,?,?)", ("lu-six-ui-%d" % _i, "Six UI %d" % _i, "URBAN", "mla-c", "six-ui@x.in", "lucknow"))
    _c60.commit()
    ctx = br.new_context(viewport={"width": 1280, "height": 900}); ctx.add_cookies([{"name": "test_email", "value": "six-ui@x.in", "url": B}])
    p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
    p.goto(B + "/rep"); p.wait_for_selector("#help-btn", timeout=15000)
    ok(p.locator("#load-retry").count() == 0 and "Six UI" in p.inner_text("body"), "rep with 6 ward offices: console opens", p.inner_text("body")[:300])
    ok(not errs, "no script errors (6-ward console)", errs)
    ctx.close()
    for _i in range(1, 7): _c60.execute("DELETE FROM local_units WHERE id = ?", ("lu-six-ui-%d" % _i,))
    _c60.commit(); _c60.close()

    # ---------- grieviq-59: Help on the government dashboards (officials' guide) ----------
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx = br.new_context(viewport={"width": width, "height": 900}, color_scheme=theme)
                ctx.add_init_script("try{localStorage.setItem('%s','%s')}catch(e){}" % (LANGKEY, lang))
                p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
                p.goto(B + ("/state" if width == 375 else "/national") + "?demo=1"); p.wait_for_selector(".tiles")
                ok(p.inner_text(".help-btn").strip() == ("Help" if lang == "en" else "सहायता"), "gov Help button labelled " + tag)
                ok(p.evaluate("(() => { const r = document.querySelector('.help-btn').getBoundingClientRect(); return r.height >= 44 && r.width >= 44; })()"), "gov Help button 44px " + tag)
                ok(no_hscroll(p), "gov header with Help: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS.replace("createTreeWalker(document.body", "createTreeWalker(document.querySelector('header.top')", 1)); ok(not bad, "gov header contrast " + tag, bad)
                p.click(".help-btn"); p.wait_for_selector(".gq-drawer")
                ok(sop_ids(p) == ["gqd-signin", "gqd-read", "gqd-reports", "gqd-followup", "gqd-access", "gqd-demo"], "officials' guide: six parts " + tag, sop_ids(p))
                txt = p.inner_text(".gq-drawer")
                ok(("fewer than 10" in txt and "never named or ranked" in txt) if lang == "en" else ("10 से कम" in txt and not re.findall(r"guide\.\w+", txt)), "officials' guide rules in " + lang + " " + tag)
                bad = p.evaluate(DLG_CONTRAST); ok(not bad, "gov Help contrast AA " + tag, bad[:6])
                ok(not p.evaluate(DLG_TARGETS), "gov Help targets >= 24px " + tag, p.evaluate(DLG_TARGETS))
                ok(p.evaluate("(() => { const d = document.querySelector('.gq-dbody'); return d.scrollWidth <= d.clientWidth + 1; })()"), "gov Help: no sideways scroll " + tag)
                if lang == "en" or width == 375: shot(p, "59-gov-help-" + tag, full=False)
                p.keyboard.press("Escape"); p.wait_for_timeout(100)
                ok(p.locator(".gq-drawer").count() == 0 and "help-btn" in (p.evaluate("document.activeElement.className") or ""), "gov: Esc closes, focus back on Help " + tag)
                ok(not errs, "no script errors (gov help %s)" % tag, errs)
                ctx.close()
    ctx = br.new_context(viewport={"width": 1280, "height": 900}); p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
    p.goto(B + "/state"); p.wait_for_selector("#f-email")
    ok(p.locator(".help-btn").count() == 1, "Help is there before signing in")
    p.click(".help-btn"); p.wait_for_selector(".gq-drawer")
    p.click(".gq-drawer [data-gq-jump=signin]"); ok(p.evaluate("document.activeElement.id") == "gqd-signin", "officials' guide opens at sign-in")
    p.keyboard.press("Escape"); ok(not errs, "no script errors (gov sign-in help)", errs); ctx.close()
    # The department reply guide uses the dashboard's own words
    ctx = br.new_context(viewport={"width": 1280, "height": 900}); p = ctx.new_page()
    p.goto(B + "/guide"); p.wait_for_selector("#gq-page")
    p.evaluate("document.getElementById('gq-page').innerHTML = GIQGuide.render('dept', { prefix: 'x-' })")
    t = p.inner_text("#gq-page")
    ok(all(w in t for w in ("The work is scheduled for a date", "The work has started", "The work is done", "This isn't our office's work", "It can't be done", "What is the position?")), "department guide uses the reply words on the dashboard")
    ok("Scheduled" not in t and "Work started" not in t, "old reply words are gone")
    ctx.close()

    # ---------- grieviq-61: search box in Help ----------
    def _first_hit(p): return p.eval_on_selector_all("[data-gq-results] .gq-stitle", "els => els.map(e => e.textContent)")
    for theme in ("light", "dark"):
        for width in (1280, 375):
            tag = "%s-%d" % (theme, width)
            ctx, p, errs = P("super@test.in", theme=theme, width=width)
            p.goto(B + "/admin-dashboard.html"); p.wait_for_selector("#adm-help", state="attached")
            if not p.is_visible("#adm-help"): p.click("#adm-menu-btn")
            p.click("#adm-help"); p.wait_for_selector(".gq-drawer [data-gq-q]")
            ok(p.locator(".gq-drawer label[for=gqd-q]").count() == 1 and p.locator(".gq-drawer .gq-lens").count() == 1, "admin Help: labelled search box with a lens " + tag)
            ok(p.evaluate("document.getElementById('gqd-q').getBoundingClientRect().height >= 44"), "search box 44px " + tag)
            p.fill("#gqd-q", "how do I add a new staff member"); p.wait_for_timeout(700)
            hits = _first_hit(p)
            ok(hits and hits[0].startswith("Staff accounts"), "'add a new staff member' finds Staff accounts first " + tag, hits)
            ok(p.locator("[data-gq-results] mark").count() >= 2, "matching words are marked " + tag)
            ok(p.inner_text(".gq-drawer [data-gq-status]").strip().endswith("matching parts of the guide") or "1 matching" in p.inner_text(".gq-drawer [data-gq-status]"), "result count announced " + tag, p.inner_text(".gq-drawer [data-gq-status]"))
            bad = p.evaluate(DLG_CONTRAST); ok(not bad, "search results contrast AA " + tag, bad[:6])
            ok(not p.evaluate(DLG_TARGETS), "search results targets >= 24px " + tag, p.evaluate(DLG_TARGETS))
            ok(p.evaluate("(() => { const d = document.querySelector('.gq-dbody'); return d.scrollWidth <= d.clientWidth + 1; })()"), "search: no sideways scroll " + tag)
            shot(p, "61-admin-search-" + tag, full=False)
            p.focus("#gqd-q"); p.keyboard.press("Enter"); p.wait_for_timeout(150)
            ok(p.evaluate("document.activeElement.id") == "gqd-staff", "Enter opens the first result " + tag, p.evaluate("document.activeElement.id"))
            p.focus("#gqd-q"); p.keyboard.press("Escape"); p.wait_for_timeout(100)
            ok(p.locator(".gq-drawer").count() == 1 and p.input_value("#gqd-q") == "" and p.locator("[data-gq-results] li").count() == 0, "Esc first clears the search " + tag)
            p.fill("#gqd-q", "zzqx flibber"); p.wait_for_timeout(100)
            ok("Nothing matched" in p.inner_text("[data-gq-results]"), "nothing found says so " + tag)
            p.fill("#gqd-q", ""); p.keyboard.press("Escape"); p.wait_for_timeout(100)
            ok(p.locator(".gq-drawer").count() == 0, "Esc on an empty search closes Help " + tag)
            ok(not errs, "no script errors (admin search %s)" % tag, errs)
            ctx.close()
    # Hindi question, and the language switch keeps the question
    ctx, p, errs = P("super@test.in", theme="light", width=1280)
    p.goto(B + "/admin-dashboard.html"); p.wait_for_selector("#adm-help", state="attached")
    if not p.is_visible("#adm-help"): p.click("#adm-menu-btn")
    p.click("#adm-help"); p.wait_for_selector("#gqd-q")
    p.fill("#gqd-q", "payment not matching bank statement"); p.wait_for_timeout(100)
    ok(_first_hit(p)[:1] == ["Bills and payments"], "'payment not matching' finds Bills and payments", _first_hit(p))
    p.evaluate("GIQ.setLang ? GIQ.setLang('hi') : document.querySelector('[data-lang-toggle]').click()"); p.wait_for_timeout(300)
    ok(p.input_value("#gqd-q") == "payment not matching bank statement" and len(_first_hit(p)) >= 1 and _first_hit(p)[0] == "बिल और भुगतान", "language switch keeps the question; titles in Hindi", _first_hit(p))
    p.fill("#gqd-q", "कर्मचारी अवकाश"); p.wait_for_timeout(100)
    ok(_first_hit(p)[:1] and _first_hit(p)[0].startswith("स्टाफ़ खाते"), "Hindi question finds staff leave", _first_hit(p))
    # grieviq-62: the owner's question "how to add a representative" (it found nothing useful before)
    p.fill("#gqd-q", "जनप्रतिनिधि कैसे जोड़ें"); p.wait_for_timeout(100)
    ok(_first_hit(p)[:1] and _first_hit(p)[0].startswith("जनप्रतिनिधि जोड़ें"), "Hindi: add a representative", _first_hit(p))
    p.evaluate("GIQ.setLang ? GIQ.setLang('en') : document.querySelector('[data-lang-toggle]').click()"); p.wait_for_timeout(300)
    for q62 in ("how to add a representative", "add a councillor", "add mla", "add pradhan"):
        p.fill("#gqd-q", q62); p.wait_for_timeout(100)
        ok(_first_hit(p)[:1] and _first_hit(p)[0].startswith("Add a representative"), "'%s' finds Add a representative first" % q62, _first_hit(p))
    p.fill("#gqd-q", "complaint stuck"); p.wait_for_timeout(100)
    ok(_first_hit(p)[:1] == ["Exceptions and nudges"], "'complaint stuck' finds Exceptions and nudges", _first_hit(p))
    p.fill("#gqd-q", "change time limit"); p.wait_for_timeout(100)
    ok(_first_hit(p)[:1] == ["Issue types and department types"], "'change time limit' finds Issue types", _first_hit(p))
    p.fill("#gqd-q", "no code received"); p.wait_for_timeout(100)
    ok(_first_hit(p)[:1] == ["No sign-in code arrives"], "'no code received' finds the sign-in problem", _first_hit(p))
    # grieviq-63: a result opens the step that answers the question, marked
    p.fill("#gqd-q", "data entry operator request"); p.wait_for_timeout(100)
    p.click("[data-gq-results] a.gq-stitle >> nth=0"); p.wait_for_timeout(200)
    ok(p.locator("#gqd-data li.gq-hit").count() == 1 and "Data entry operator" in p.inner_text("#gqd-data li.gq-hit"), "the result opens and marks the step that answers it")
    bad = p.evaluate(DLG_CONTRAST); ok(not bad, "marked step contrast AA", bad[:6])
    p.fill("#gqd-q", "staff kaise add kare"); p.wait_for_timeout(100)
    ok(_first_hit(p)[:1] and _first_hit(p)[0].startswith("Staff accounts"), "Hinglish question works", _first_hit(p))
    p.fill("#gqd-q", "backup"); p.wait_for_timeout(100)
    ok("Nothing matched" in p.inner_text("[data-gq-results]"), "a question the guide doesn't cover says so instead of guessing", _first_hit(p))
    # grieviq-63: the question bank. Every question people ask (dev-tests/guide-questions.json)
    # must bring the right part of the guide first, in the real browser code.
    QB = json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "guide-questions.json"), encoding="utf-8"))
    QROLES = {"admin": ["super_admin"], "rep": ["REPRESENTATIVE", "MULTI_SEAT"]}
    for qlang in ("en", "hi"):
        p.evaluate("GIQ.setLang ? GIQ.setLang('%s') : 0" % qlang); p.wait_for_timeout(200)
        misses = []
        for kind in ("admin", "rep", "dept", "citizen", "gov"):
            got = p.evaluate("(a) => a.qs.map(q => (GIQGuide.find(a.kind, a.roles, q[0])[0] || 'NOTHING'))", {"kind": kind, "roles": QROLES.get(kind), "qs": QB[kind]})
            for (q, want), g in zip(QB[kind], got):
                if g not in want.split("|"): misses.append("%s: %s -> %s (want %s)" % (kind, q, g, want))
        nq = sum(len(QB[k]) for k in ("admin", "rep", "dept", "citizen", "gov"))
        ok(not misses, "question bank (%d questions, Help in %s): right guide part first for every one" % (nq, qlang), misses[:10])
    p.evaluate("GIQ.setLang ? GIQ.setLang('en') : 0"); p.wait_for_timeout(200)
    ok(not errs, "no script errors (Hindi search)", errs); ctx.close()
    # Citizen guide page, rep, dept and gov panels
    ctx = br.new_context(viewport={"width": 375, "height": 900}); p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
    p.goto(B + "/guide"); p.wait_for_selector("#gq-q")
    p.fill("#gq-q", "how to reopen complaint"); p.wait_for_timeout(100)
    ok(_first_hit(p)[:1] == ["Confirm the fix, or reopen"], "citizen guide page: reopen", _first_hit(p))
    p.click("[data-gq-results] a.gq-stitle >> nth=0"); p.wait_for_timeout(150)
    ok(p.evaluate("document.activeElement.id") == "gq-confirm", "citizen guide page: result opens that part")
    ok(no_hscroll(p), "citizen guide page search: no sideways scroll")
    bad = p.evaluate(CONTRAST_JS.replace("createTreeWalker(document.body", "createTreeWalker(document.querySelector('.gq-search')", 1)); ok(not bad, "citizen search contrast AA", bad)
    shot(p, "61-guide-search-375", full=False)
    ok(not errs, "no script errors (citizen search)", errs); ctx.close()
    ctx = br.new_context(viewport={"width": 1280, "height": 900}, color_scheme="dark"); ctx.add_cookies([{"name": "test_email", "value": "r1@x.in", "url": B}])
    p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
    p.goto(B + "/rep"); p.wait_for_selector("#help-btn", timeout=15000); p.click("#help-btn"); p.wait_for_selector("#gqd-q")
    p.fill("#gqd-q", "how do I add a field worker"); p.wait_for_timeout(100)
    ok(_first_hit(p)[:1] == ["Manage the team"], "rep Help: add a field worker", _first_hit(p))
    bad = p.evaluate(DLG_CONTRAST); ok(not bad, "rep search contrast AA dark", bad[:6])
    ok(not errs, "no script errors (rep search)", errs); ctx.close()
    ctx = br.new_context(viewport={"width": 375, "height": 900}); p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
    p.goto(B + "/state?demo=1"); p.wait_for_selector(".tiles"); p.click(".help-btn"); p.wait_for_selector("#gqd-q")
    p.fill("#gqd-q", "download excel report"); p.wait_for_timeout(100)
    ok(_first_hit(p)[:1] == ["The six MIS reports"], "officials' Help: excel report", _first_hit(p))
    ok(not errs, "no script errors (gov search)", errs); ctx.close()
    ctx = br.new_context(viewport={"width": 1280, "height": 900})
    ok(dept_signin(ctx) == 200, "dept signed in (search)")
    p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
    p.goto(B + "/dept"); p.wait_for_selector(".tiles"); p.click("#guide-link"); p.wait_for_selector("#gqd-q")
    p.fill("#gqd-q", "case is not ours"); p.wait_for_timeout(100)
    ok(_first_hit(p)[:1] == ["Reply on a forwarded case"], "department Help: not ours", _first_hit(p))
    ok(not errs, "no script errors (dept search)", errs); ctx.close()

    # ---------- grieviq-41: payment details, bills, billing for offices ----------
    import numpy as _np, cv2 as _cv2
    def qr_text(p, sel):
        png = p.locator(sel).first.screenshot()
        img = _cv2.imdecode(_np.frombuffer(png, _np.uint8), _cv2.IMREAD_GRAYSCALE)
        img = _cv2.copyMakeBorder(img, 20, 20, 20, 20, _cv2.BORDER_CONSTANT, value=255)
        return _cv2.QRCodeDetector().detectAndDecode(img)[0]
    def api41(method, path, body=None, email="super@test.in"):
        req = _ur.Request(B + path, method=method, data=json.dumps(body).encode() if body is not None else None, headers={"content-type": "application/json", "test-email": email})
        try:
            with _ur.urlopen(req) as r: return r.status, json.loads(r.read() or b"{}")
        except _ur.HTTPError as e: return e.code, json.loads(e.read() or b"{}")
    st, nb = api41("POST", "/api/admin/bills", {"action": "create", "officeType": "REP", "officeTier": "LOCAL", "officeId": "lu-hazratganj", "planId": "councillor", "period": "MONTH", "from": "2026-12-01"})
    ok(st == 200, "bill for the console tests", nb)
    REF = nb.get("ref")
    # Admin: Payment details tab
    for theme in ("dark", "light"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx, p, errs = P("super@test.in", width=width, theme=theme, lang=lang)
                p.goto(B + "/admin-site-settings.html#payments"); p.wait_for_selector("#panel-payments")
                ok(p.get_attribute("#tab-payments", "aria-selected") == "true" and p.locator("#panel-payments fieldset.card").count() == 2, "link opens Payment details with UPI and bank " + tag)
                ok(p.locator("#panel-payments .giq-qr").count() == 1, "QR preview drawn from the UPI ID " + tag)
                if lang == "hi": ok(not re.findall(r"adm\.\w+", p.inner_text("main")), "payment details in Hindi, nothing missing " + tag)
                ok(no_hscroll(p), "payment details: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS); ok(not bad, "payment details contrast AA " + tag, bad[:5])
                small = p.evaluate(TARGET_JS); ok(not small, "payment details targets " + tag, small[:5])
                if width == 1280 or lang == "en": shot(p, "53-pay-details-" + tag, full=True)
                ok(not errs, "no script errors (payment details %s)" % tag, errs)
                ctx.close()
    ctx, p, errs = P("super@test.in")
    p.goto(B + "/admin-site-settings.html#payments"); p.wait_for_selector("#panel-payments")
    ok(qr_text(p, "#panel-payments .giq-qr").startswith("upi://pay?pa=iotc%40hdfcbank&pn="), "the preview QR code scans to the UPI ID")
    p.fill("#f-payments-methods-1-ifsc", "HDFC123"); p.click("#ss-save"); p.wait_for_selector("#ss-errors:not([hidden])")
    ok("IFSC" in p.inner_text("#ss-errors") and "(1)" in p.inner_text("#tab-payments"), "wrong IFSC listed, tab shows the count")
    p.click("#ss-errors a"); ok(p.evaluate("document.activeElement.id") == "f-payments-methods-1-ifsc", "error link goes to the IFSC field")
    p.fill("#f-payments-methods-1-ifsc", "HDFC0009999")
    with p.expect_response(lambda r: "/api/admin/site-settings" in r.url and r.request.method == "POST"): p.click("#ss-save")
    p.wait_for_selector("#ss-publish")
    with p.expect_response(lambda r: "/api/admin/site-settings" in r.url and r.request.method == "POST"): p.click("#ss-publish")
    p.wait_for_timeout(400)
    ok("go live on" in p.inner_text("#ss-msg"), "publishing says when the new details go live", p.inner_text("#ss-msg"))
    p.click("#tab-payments"); p.wait_for_selector("#pay-wait")
    ok("HDFC0001234" in p.inner_text(".pay-now") and p.is_visible("#pay-cancel"), "waiting change shown, with what offices see now")
    bad = p.evaluate(CONTRAST_JS); ok(not bad, "waiting box contrast AA", bad[:5])
    shot(p, "54-pay-waiting", full=False)
    with p.expect_response(lambda r: "/api/admin/site-settings" in r.url and r.request.method == "POST"): p.click("#pay-cancel")
    p.wait_for_timeout(300)
    ok("cancelled" in p.inner_text("#ss-msg") and p.locator("#pay-wait").count() == 0, "change cancelled from the page")
    ok(not errs, "no script errors (payment flow)", errs)
    ctx.close()
    ctx, p, errs = P("audit@test.in")
    p.goto(B + "/admin-site-settings.html#payments"); p.wait_for_selector("#panel-payments")
    ok("XXXXXXXXXX5678" in p.inner_text("#panel-payments") or p.input_value("#f-payments-methods-1-accountNumber") == "XXXXXXXXXX5678", "auditor: account number partly hidden")
    ok(p.locator("[data-act=addupi]").count() == 0 and p.locator("#pay-cancel").count() == 0, "auditor: view only")
    ctx.close()

    # Admin: Bills page
    for theme in ("dark", "light"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx, p, errs = P("super@test.in", width=width, theme=theme, lang=lang)
                p.goto(B + "/admin-bills.html"); p.wait_for_selector(".bl-card")
                p.click("#blf-all"); p.wait_for_timeout(100)
                if lang == "hi": ok(not re.findall(r"adm\.\w+", p.inner_text("main")), "bills in Hindi, nothing missing " + tag)
                ok(no_hscroll(p), "bills: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS); ok(not bad, "bills contrast AA " + tag, bad[:5])
                small = p.evaluate(TARGET_JS); ok(not small, "bills targets " + tag, small[:5])
                if width == 1280 or lang == "en": shot(p, "55-bills-" + tag, full=(width == 375))
                ok(not errs, "no script errors (bills %s)" % tag, errs)
                ctx.close()
    ctx, p, errs = P("ops@test.in")
    p.goto(B + "/admin-bills.html"); p.wait_for_selector("#bl-make")
    ok(p.locator('a[href="/admin-bills.html"]').count() >= 1, "Bills in the admin menu")
    p.click("#bl-make"); p.wait_for_timeout(150)
    ok(p.evaluate("document.activeElement.id") == "bl-q" and p.get_attribute("#bl-q", "aria-invalid") == "true", "no office chosen: told, focus on the search")
    p.fill("#bl-q", "Tilak"); p.press("#bl-q", "Enter"); p.wait_for_selector("#bl-o-0")
    p.check("#bl-o-0"); p.wait_for_function("document.getElementById('bl-quote').textContent.indexOf('2,359') >= 0")
    ok(p.input_value("#bl-plan") == "councillor", "plan suggested from the office")
    p.check("#bl-p-YEAR"); p.wait_for_function("document.getElementById('bl-quote').textContent.indexOf('23,588') >= 0")
    ok("₹19,990 + GST 18% ₹3,598 = ₹23,588" in p.inner_text("#bl-quote"), "yearly amount worked out from published prices")
    with p.expect_response(lambda r: "/api/admin/bills" in r.url and r.request.method == "POST"): p.click("#bl-make")
    p.wait_for_timeout(400)
    ok("GIQ-PAY-" in p.inner_text("#bl-msg") and p.evaluate("document.activeElement.id") == "bl-msg", "bill made; message focused")
    # the department's payment waiting for a check (from the API tests)
    p.click("#blf-REPORTED"); p.wait_for_selector("[data-confirm]")
    card = p.locator(".bl-card").first
    card.locator("[data-return]").click(); p.wait_for_timeout(200)
    ok(p.evaluate("document.activeElement.id").startswith("rr-") and "at least 5" in card.inner_text(), "turning back without a reason: told, focus on the reason")
    with p.expect_response(lambda r: "/api/admin/bills" in r.url and r.request.method == "POST"): card.locator("[data-confirm]").click()
    p.wait_for_timeout(400)
    ok("bill is paid" in p.inner_text("#bl-msg"), "payment confirmed from the page")
    ok(not errs, "no script errors (bills flow)", errs)
    ctx.close()

    # Representative console: Billing
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            for width in (375, 1280):
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx = br.new_context(viewport={"width": width, "height": 900}, color_scheme=theme)
                ctx.add_cookies([{"name": "test_email", "value": "om@x.in", "url": B}])
                ctx.add_init_script("try{localStorage.setItem('%s','%s');localStorage.setItem('giq-welcome-om@x.in','1')}catch(e){}" % (LANGKEY, lang))
                p = ctx.new_page(); errs = []
                p.on("pageerror", lambda e: errs.append(str(e)))
                p.goto(B + "/rep"); p.wait_for_selector('[data-tab="billing"]', timeout=15000)
                p.click('[data-tab="billing"]'); p.wait_for_selector(".gb-bill")
                ok(p.locator(".gb-bill").count() >= 2 and p.locator(".gb-st-DUE").count() == 1, "office manager sees the office's bills, one to pay " + tag)
                txt = p.inner_text("#billing-box")
                if lang == "en": ok("Pay only to the details shown here." in txt and "IOTC" in txt and REF in txt and "₹2,359" in txt, "billing wording " + tag, txt[:300])
                else: ok("भुगतान केवल यहाँ दिखाए गए विवरण पर करें।" in txt and not re.findall(r"bill\.\w+", txt), "billing in Hindi, nothing missing " + tag, txt[:300])
                ok(no_hscroll(p), "billing: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS.replace("createTreeWalker(document.body", "createTreeWalker(document.getElementById('billing-box')", 1)); ok(not bad, "billing contrast AA " + tag, bad[:5])
                small = p.evaluate("() => Array.from(document.querySelectorAll('#billing-box button, #billing-box a, #billing-box input')).filter(e => { const r = e.getBoundingClientRect(); return r.width > 0 && (e.type !== 'radio') && (r.height < 24 || r.width < 24); }).map(e => e.textContent || e.id)")
                ok(not small, "billing targets >= 24px " + tag, small)
                shot(p, "56-rep-billing-" + tag, full=True)
                ok(not errs, "no script errors (rep billing %s)" % tag, errs)
                ctx.close()
    ctx = br.new_context(viewport={"width": 390, "height": 900}, permissions=["clipboard-read", "clipboard-write"])
    ctx.add_cookies([{"name": "test_email", "value": "r1@x.in", "url": B}])
    p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
    p.goto(B + "/rep"); p.wait_for_selector('[data-tab="billing"]', timeout=15000); p.click('[data-tab="billing"]'); p.wait_for_selector(".gb-st-DUE")
    due = p.locator(".gb-st-DUE")
    q = qr_text(p, ".gb-st-DUE .gb-qr")
    ok(q.startswith("upi://pay?pa=iotc%40hdfcbank&pn=") and "am=2359.00" in q and "cu=INR" in q and "tn=" + REF in q, "bill QR scans to the UPI ID with the amount and bill number", q)
    ok(due.locator("a.gb-applink").get_attribute("href") == q, "UPI app link matches the QR")
    due.locator("[data-copy]").first.click(); p.wait_for_timeout(200)
    ok("copied" in p.inner_text("#gb-copied") and p.evaluate("navigator.clipboard.readText()") == "iotc@hdfcbank", "copy button copies the UPI ID and says so")
    due.locator("button[type=submit]").click(); p.wait_for_timeout(200)
    ok(p.evaluate("document.activeElement.className") == "gb-errbox" and p.get_attribute(".gb-st-DUE input[name=utr]", "aria-invalid") == "true", "empty reference: error list focused, field marked")
    due.locator(".gb-errbox a").first.click(); ok(p.evaluate("document.activeElement.name") == "utr", "error link goes to the field")
    p.fill(".gb-st-DUE input[name=utr]", "61234567"); due.locator("button[type=submit]").click(); p.wait_for_timeout(400)
    ok("10 to 22" in p.inner_text(".gb-st-DUE .gb-errbox"), "short reference refused by the server, in words")
    p.fill(".gb-st-DUE input[name=utr]", "777766665555")
    with p.expect_response(lambda r: "/api/billing" in r.url and r.request.method == "POST"): p.locator(".gb-st-DUE button[type=submit]").click()
    p.wait_for_selector("#gb-msg")
    ok(p.evaluate("document.activeElement.id") == "gb-msg" and "Thank you" in p.inner_text("#gb-msg") and p.locator(".gb-st-REPORTED").count() == 1, "payment recorded; thank-you focused; bill shows it's being checked")
    shot(p, "57-rep-billing-sent", full=False)
    ok(not errs, "no script errors (rep billing flow)", errs)
    ctx.close()
    ctx = br.new_context(); ctx.add_cookies([{"name": "test_email", "value": "fw@x.in", "url": B}])
    p = ctx.new_page(); p.goto(B + "/rep"); p.wait_for_selector(".tab-btn", timeout=15000)
    ok(p.locator('[data-tab="billing"]').count() == 0, "field worker: no Billing tab")
    ctx.close()

    # Department dashboard: Billing
    for theme in ("light", "dark"):
        ctx = br.new_context(viewport={"width": 375, "height": 900}, color_scheme=theme)
        ok(dept_signin(ctx) == 200, "dept signed in for billing " + theme)
        p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
        p.goto(B + "/dept"); p.wait_for_selector("#go-billing")
        p.click("#go-billing"); p.wait_for_selector(".gb-bill")
        ok(p.evaluate("document.activeElement.id") == "gb-title" and p.evaluate("location.hash") == "#billing", "billing opens, heading focused " + theme)
        ok(p.locator(".gb-st-PAID").count() == 1, "department sees its paid bill " + theme)
        bad = p.evaluate(CONTRAST_JS.replace("createTreeWalker(document.body", "createTreeWalker(document.getElementById('billing-box')", 1)); ok(not bad, "dept billing contrast AA " + theme, bad[:5])
        ok(no_hscroll(p), "dept billing: no sideways scroll " + theme)
        shot(p, "58-dept-billing-" + theme, full=True)
        p.click("#back"); p.wait_for_timeout(200)
        ok(p.evaluate("document.activeElement.id") == "go-billing", "back to the list, focus on Open billing " + theme)
        ok(not errs, "no script errors (dept billing %s)" % theme, errs)
        ctx.close()

    # ---------- grieviq-42: home page layout (owner's design, 11 Oct 2026) ----------
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx = br.new_context(viewport={"width": width, "height": 900}, color_scheme=theme)
                ctx.add_init_script("try{localStorage.setItem('%s','%s')}catch(e){}" % (LANGKEY, lang))
                p = ctx.new_page(); errs = []
                p.on("pageerror", lambda e: errs.append(str(e)))
                p.goto(B + "/"); p.wait_for_selector(".home-more .feedback-card"); p.wait_for_timeout(400)
                ok(p.locator(".stats, #stat-wards, #stat-device").count() == 0, "number boxes removed " + tag)
                ok(p.evaluate("document.querySelector('.home-side').firstElementChild.id") == "loc-card", "Where is the problem? comes first on the right " + tag)
                box = lambda sel: p.evaluate("(s) => { const r = document.querySelector(s).getBoundingClientRect(); return [r.top + scrollY, r.left, r.width, r.height]; }", sel)
                em, fb, tr = box(".home-more .emergency"), box(".home-more .feedback-card"), box('.home-side a[href="/status"]')
                if width == 1280:
                    ok(abs(em[0] - fb[0]) < 2 and abs(em[3] - fb[3]) < 2 and em[1] < fb[1], "Emergency left, Feedback right, level and the same height " + tag, (em, fb))
                    ok(p.locator(".site-nav a").count() == 3 and p.locator(".site-nav a[href='/for-offices']").count() == 0, "top bar: no For offices link " + tag)
                else:
                    ok(tr[0] < fb[0] < em[0], "phone order: track, feedback, emergency last " + tag, (tr, fb, em))
                ok(p.locator("#offices-card").count() == 0 and p.locator(".footer-links a[href='/for-offices']").count() == 1, "offices: only the small footer link " + tag)
                ok(no_hscroll(p), "home: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS); ok(not bad, "home contrast AA " + tag, bad[:5])
                small = p.evaluate("() => Array.from(document.querySelectorAll('.home-more a, .home-side a, .home-side button')).filter(e => { const r = e.getBoundingClientRect(); return r.width > 0 && r.height > 0 && (r.height < 24 || r.width < 24); }).map(e => e.textContent.trim().slice(0, 30))")
                ok(not small, "home targets >= 24px " + tag, small)
                shot(p, "59-home-" + tag, full=True)
                ok(not errs, "no script errors (home %s)" % tag, errs)
                ctx.close()
    # Reports saved on this device: one line with "Forget", only when there are some
    ctx = br.new_context(viewport={"width": 1280, "height": 900})
    ctx.add_init_script("try{if(!sessionStorage.getItem('x')){localStorage.setItem('griq-my-reports', JSON.stringify(['GRV-TEST01','GRV-TEST02']));sessionStorage.setItem('x','1')}}catch(e){}")
    p = ctx.new_page(); p.goto(B + "/"); p.wait_for_selector(".home-more .feedback-card"); p.wait_for_timeout(300)
    ok(p.is_visible("#device-note") and "2" in p.inner_text('.home-side a[href="/status"]'), "with saved reports: count on Track and the Forget line")
    p.click("#forget-device"); p.wait_for_timeout(100)
    ok(not p.is_visible("#device-note"), "Forget clears them and hides the line")
    ctx.close()
    ctx = br.new_context(); p = ctx.new_page(); p.goto(B + "/"); p.wait_for_selector(".home-more .feedback-card"); p.wait_for_timeout(300)
    ok(not p.is_visible("#device-note"), "first visit: no device line")
    ctx.close()

    # grieviq-43: offices find plans on the sign-in pages
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            ctx = br.new_context(viewport={"width": 375, "height": 900}, color_scheme=theme)
            ctx.add_init_script("try{localStorage.setItem('%s','%s')}catch(e){}" % (LANGKEY, lang))
            p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
            for path, sel in (("/rep", ".signin-new a"), ("/dept", ".new-office a")):
                p.goto(B + path); p.wait_for_selector(sel, timeout=15000)
                txt = p.inner_text(sel)
                ok(p.get_attribute(sel, "href") == "/for-offices" and (("New to GrievIQ?" in txt) if lang == "en" else ("GrievIQ में नए हैं?" in txt)), "sign-in page links offices to plans %s %s %s" % (path, theme, lang), txt)
                bad = p.evaluate(CONTRAST_JS); ok(not bad, "sign-in page contrast %s %s %s" % (path, theme, lang), bad[:5])
                small = p.evaluate("(s) => { const r = document.querySelector(s).getBoundingClientRect(); return r.height < 24; }", sel); ok(not small, "sign-in link target %s %s" % (path, theme))
                if lang == "en": shot(p, "60-signin%s-%s" % (path.replace("/", "-"), theme), full=False)
            ok(not errs, "no script errors (sign-in pages %s %s)" % (theme, lang), errs)
            ctx.close()

    # ---------- grieviq-45: a copy of the pages reports where it runs ----------
    ctx = br.new_context(); p = ctx.new_page(); seen = []
    p.route("https://grieviq.in/api/b**", lambda route: (seen.append(route.request.url), route.fulfill(status=204)))
    p.goto(B.replace("localhost", "127.0.0.1") + "/"); p.wait_for_timeout(800)
    ok(len(seen) >= 1 and "h=127.0.0.1" in seen[0] and "k=c" in seen[0], "page on another address reports it", seen)
    n = len(seen); p.reload(); p.wait_for_timeout(600)
    ok(len(seen) == n, "only once a day per browser")
    p.goto(B + "/"); p.wait_for_timeout(600)
    ok(len(seen) == n, "GrievIQ's own address (here localhost) reports nothing")
    ctx.close()
    for theme in ("dark", "light"):
        for lang in ("en", "hi"):
            ctx, p, errs = P("super@test.in", theme=theme, lang=lang, width=375 if lang == "hi" else 1280)
            p.goto(B + "/admin-copies.html"); p.wait_for_selector("#cp-count, #cp-none")
            ok("copy.example.com" in p.inner_text("main"), "Copies page lists the copy, address as written " + theme + " " + lang, p.inner_text("main")[-300:])
            if lang == "hi": ok(not re.findall(r"adm\.\w+", p.inner_text("main")), "Copies page in Hindi, nothing missing")
            bad = p.evaluate(CONTRAST_JS); ok(not bad, "Copies page contrast AA %s %s" % (theme, lang), bad[:5])
            ok(no_hscroll(p) or p.evaluate("document.querySelector('.table-wrap').scrollWidth > 0"), "Copies page: no page-wide sideways scroll")
            shot(p, "61-copies-%s-%s" % (theme, lang), full=False)
            ok(not errs, "no script errors (copies %s %s)" % (theme, lang), errs)
            ctx.close()
    ctx, p, errs = P("ops@test.in"); p.goto(B + "/admin-copies.html"); p.wait_for_timeout(800)
    ok("does not have access" in p.inner_text("main"), "operations admin: no access to Copies")
    ctx.close()

    # ---------- grieviq-46: office subscriptions ----------
    _t46 = (_dt.datetime.utcnow() + _dt.timedelta(hours=5, minutes=30)).date()
    D46 = lambda n: (_t46 + _dt.timedelta(days=n)).isoformat()
    def subrow(sql, args=()):
        d = _sq.connect(WORK + "/test.db", timeout=10); d.execute(sql, args); d.commit(); d.close()
    r2m = api41("GET", "/api/me", email="r2@x.in")[1]["mandates"][0]
    R2 = (r2m["tier"], r2m["id"])
    jal = _sq.connect(WORK + "/test.db").execute("SELECT office_id FROM dept_officers WHERE email = 'ae.zone3@nic.in' AND status = 'ACTIVE'").fetchone()[0]
    WHY = "Checked with the office by phone"
    # r2's office in grace, r1's office paused by the super admin, the department due soon
    subrow("UPDATE office_subscriptions SET paid_through = ?, last_notice = NULL WHERE office_tier = ? AND office_id = ?", (D46(-3), R2[0], R2[1]))
    st, _ = api41("POST", "/api/admin/subscriptions", {"action": "pause", "type": "REP", "tier": "LOCAL", "id": "lu-hazratganj", "reason": "Office asked for a break"}); ok(st == 200, "r1's office paused for the console tests")
    subrow("UPDATE office_subscriptions SET paid_through = ? WHERE office_type = 'DEPT' AND office_id = ?", (D46(3), jal))
    for theme in ("dark", "light"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx, p, errs = P("super@test.in", width=width, theme=theme, lang=lang)
                p.goto(B + "/admin-subscriptions.html"); p.wait_for_selector(".sb-card, #sb-none")
                ok(p.locator(".sb-card").count() >= 3 and p.locator(".sb-card.s-PAUSED").count() >= 1 and p.locator(".sb-card.s-GRACE").count() >= 1 and p.locator(".sb-card.s-DUE_SOON").count() >= 1, "Needs attention lists paused, grace and due-soon offices " + tag)
                if lang == "hi": ok(not re.findall(r"adm\.\w+", p.inner_text("main")), "Subscriptions page in Hindi, nothing missing " + tag, re.findall(r"adm\.\w+", p.inner_text("main"))[:5])
                p.locator(".sb-card.s-GRACE details summary").first.click(); p.wait_for_timeout(100)
                if lang == "hi": ok(not re.findall(r"adm\.\w+", p.inner_text("main")), "change form in Hindi, nothing missing " + tag, re.findall(r"adm\.\w+", p.inner_text("main"))[:5])
                ok(no_hscroll(p), "Subscriptions: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS); ok(not bad, "Subscriptions contrast AA " + tag, bad[:5])
                small = p.evaluate(TARGET_JS); ok(not small, "Subscriptions targets " + tag, small[:5])
                shot(p, "62-subs-" + tag, full=True)
                ok(not errs, "no script errors (subscriptions %s)" % tag, errs)
                ctx.close()
    # the super admin changes pilot dates from the page; a missing reason is pointed out
    ctx, p, errs = P("super@test.in")
    p.goto(B + "/admin-subscriptions.html"); p.wait_for_selector(".sb-card")
    p.check("#sbf-all"); p.fill("#sb-q", r2m["name"][:6]); p.dispatch_event("#sb-q", "input"); p.wait_for_timeout(150)
    card = p.locator(".sb-card.s-GRACE").first
    card.locator("summary").click()
    pre = "#sb-REP-" + R2[0] + "-" + re.sub(r"[^a-zA-Z0-9]", "-", R2[1])
    p.fill(pre + "-ps", D46(-60)); p.fill(pre + "-pe", D46(20))
    with p.expect_response(lambda r: "/api/admin/subscriptions" in r.url and r.request.method == "POST"): card.locator("[data-act=set_pilot]").click()
    p.wait_for_timeout(150)
    ok(p.evaluate("document.activeElement.id") == pre[1:] + "-r" and "reason" in p.inner_text(".sb-card .row-msg.error").lower(), "no reason: message shown and the reason box focused", p.evaluate("document.activeElement.id"))
    p.fill(pre + "-r", WHY)
    with p.expect_response(lambda r: "/api/admin/subscriptions" in r.url and r.request.method == "POST"): p.locator(pre.replace("#", "#") + "-r").locator("xpath=ancestor::article").locator("[data-act=set_pilot]").click()
    p.wait_for_selector("#sb-msg.success")
    ok(p.evaluate("document.activeElement.id") == "sb-msg" and "Saved" in p.inner_text("#sb-msg"), "pilot dates saved, message focused", p.inner_text("#sb-msg"))
    row = _sq.connect(WORK + "/test.db").execute("SELECT pilot_start, pilot_end FROM office_subscriptions WHERE office_tier = ? AND office_id = ?", R2).fetchone()
    ok(row == (D46(-60), D46(20)), "pilot start and end changed by the super admin", row)
    p.locator(".sb-card [data-hist]").first.click(); p.wait_for_selector(".sb-hist li")
    ok("Pilot dates changed" in p.inner_text(".sb-hist") and WHY in p.inner_text(".sb-hist"), "history shows the change with its reason")
    shot(p, "63-subs-changed", full=False)
    ok(not errs, "no script errors (subscriptions change)", errs)
    ctx.close()
    for who in ("ops@test.in", "audit@test.in"):
        ctx, p, errs = P(who); p.goto(B + "/admin-subscriptions.html"); p.wait_for_selector(".sb-card")
        ok(p.is_visible("#readonly-tag") and p.locator("[data-act]").count() == 0, "view only for " + who)
        ctx.close()
    ctx, p, errs = P("mod@test.in"); p.goto(B + "/admin-subscriptions.html"); p.wait_for_timeout(800)
    ok("does not have access" in p.inner_text("main"), "moderator: no access to Subscriptions")
    ctx.close()
    # dashboard tile
    ctx, p, errs = P("super@test.in"); p.goto(B + "/admin-dashboard.html"); p.wait_for_selector('a.card[href="/admin-subscriptions.html"]')
    ok("due soon" in p.inner_text('a.card[href="/admin-subscriptions.html"]'), "dashboard: subscriptions tile")
    ok(not errs, "no script errors (dashboard tile)", errs)
    ctx.close()

    # representative console: back in grace, then paused (unpaid), field worker of a paused office
    subrow("UPDATE office_subscriptions SET pilot_end = ?, paid_through = ? WHERE office_tier = ? AND office_id = ?", (D46(-30), D46(-3), R2[0], R2[1]))
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            for width in (375, 1280):
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx = br.new_context(viewport={"width": width, "height": 900}, color_scheme=theme)
                ctx.add_cookies([{"name": "test_email", "value": "r2@x.in", "url": B}])
                ctx.add_init_script("try{localStorage.setItem('%s','%s');localStorage.setItem('giq-welcome-r2@x.in','1')}catch(e){}" % (LANGKEY, lang))
                p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
                p.goto(B + "/rep"); p.wait_for_selector(".tab-btn", timeout=15000)
                ok(p.locator(".plan-note").count() == 1 and p.locator(".plan-note.paused").count() == 0 and p.locator("[data-plan-pay]").count() == 1, "grace note with Pay " + tag)
                txt = p.inner_text(".plan-note")
                if lang == "en": ok("ended on" in txt and "Everything keeps working until" in txt, "grace wording " + tag, txt)
                else: ok(not re.findall(r"rep\.\w+", txt) and "तक सब कुछ चलता रहेगा" in txt, "grace note in Hindi " + tag, txt)
                bad = p.evaluate(CONTRAST_JS); ok(not bad, "rep plan note contrast AA " + tag, bad[:5])
                ok(no_hscroll(p), "rep plan note: no sideways scroll " + tag)
                if width == 375: shot(p, "64-rep-grace-" + tag, full=False)
                ok(not errs, "no script errors (rep grace %s)" % tag, errs)
                ctx.close()
    subrow("UPDATE office_subscriptions SET paid_through = ? WHERE office_tier = ? AND office_id = ?", (D46(-20), R2[0], R2[1]))
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            tag = "%s-%s" % (theme, lang)
            ctx = br.new_context(viewport={"width": 375, "height": 900}, color_scheme=theme)
            ctx.add_cookies([{"name": "test_email", "value": "r2@x.in", "url": B}])
            ctx.add_init_script("try{localStorage.setItem('%s','%s');localStorage.setItem('giq-welcome-r2@x.in','1')}catch(e){}" % (LANGKEY, lang))
            p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
            p.goto(B + "/rep"); p.wait_for_selector(".plan-note.paused", timeout=15000)
            txt = p.inner_text(".plan-note")
            if lang == "en": ok("paused (view only)" in txt and "move up to the next level on time" in txt, "paused wording " + tag, txt)
            else: ok(not re.findall(r"rep\.\w+", txt) and "केवल देखना" in txt, "paused note in Hindi " + tag, txt)
            bad = p.evaluate(CONTRAST_JS); ok(not bad, "paused note contrast AA " + tag, bad[:5])
            shot(p, "65-rep-paused-" + tag, full=False)
            if lang == "en" and theme == "light":
                p.click("[data-plan-pay]"); p.wait_for_selector("#billing-box .gb-bill, #billing-box .gb-empty, #billing-box h2", timeout=10000)
                ok(p.evaluate("document.activeElement.getAttribute('data-tab')") == "billing" or p.locator('[data-tab="billing"].active').count() == 1, "Pay to continue opens Billing")
            ok(not errs, "no script errors (rep paused %s)" % tag, errs)
            ctx.close()
    ctx = br.new_context(viewport={"width": 375, "height": 900}); ctx.add_cookies([{"name": "test_email", "value": "fw@x.in", "url": B}])
    ctx.add_init_script("try{localStorage.setItem('giq-welcome-fw@x.in','1')}catch(e){}")
    p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
    p.goto(B + "/rep"); p.wait_for_selector(".plan-note.paused", timeout=15000)
    txt = p.inner_text(".plan-note")
    ok("ask your representative" in txt and p.locator("[data-plan-pay]").count() == 0 and not re.search(r"\d{4}", txt), "field worker: paused, ask your office (no dates, no Pay)", txt)
    ok(not errs, "no script errors (field worker paused)", errs)
    ctx.close()
    ctx = br.new_context(viewport={"width": 375, "height": 900}); ctx.add_cookies([{"name": "test_email", "value": "om@x.in", "url": B}])
    ctx.add_init_script("try{localStorage.setItem('giq-welcome-om@x.in','1')}catch(e){}")
    p = ctx.new_page(); p.goto(B + "/rep"); p.wait_for_selector(".plan-note.paused", timeout=15000)
    ok("support@grieviq.in" in p.inner_text(".plan-note") and p.locator("[data-plan-pay]").count() == 0, "paused by GrievIQ: write to support, no Pay")
    ctx.close()

    # department dashboard: due soon, then paused (no reply)
    for theme in ("light", "dark"):
        ctx = br.new_context(viewport={"width": 375, "height": 900}, color_scheme=theme)
        ok(dept_signin(ctx) == 200, "dept signed in for plan note " + theme)
        p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
        p.goto(B + "/dept"); p.wait_for_selector(".plan-note")
        ok("paid up to" in p.inner_text(".plan-note") and p.locator("#plan-pay").count() == 1, "department: due-soon note with Billing link " + theme)
        bad = p.evaluate(CONTRAST_JS); ok(not bad, "dept plan note contrast AA " + theme, bad[:5])
        shot(p, "66-dept-due-" + theme, full=False)
        ok(not errs, "no script errors (dept due %s)" % theme, errs)
        ctx.close()
    st, _ = api41("POST", "/api/admin/subscriptions", {"action": "pause", "type": "DEPT", "tier": "", "id": jal, "reason": "Testing the pause"}); ok(st == 200, "department office paused")
    ctx = br.new_context(viewport={"width": 375, "height": 900})
    ok(dept_signin(ctx) == 200, "dept signed in while paused")
    p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
    p.goto(B + "/dept"); p.wait_for_selector(".plan-note.paused")
    ok("view only" in p.inner_text(".plan-note"), "department sees the pause")
    if p.locator("[data-open]").count():
        p.locator("[data-open]").first.click(); p.wait_for_selector("#case-h")
        ok(p.locator("#f-reply").count() == 0 and "can't reply" in p.inner_text("main"), "paused: no reply form, reason shown")
        shot(p, "67-dept-paused-case", full=False)
    ok(not errs, "no script errors (dept paused)", errs)
    ctx.close()

    # Bills: an automatic bill waits until staff send it
    subrow("UPDATE office_subscriptions SET pilot_end = ?, paid_through = ?, last_notice = NULL WHERE office_tier = ? AND office_id = ?", (D46(-30), D46(4), R2[0], R2[1]))
    run = json.loads(_ur.urlopen(_ur.Request(B + "/__test/run-subscriptions", method="POST")).read())
    ok(run.get("bills") == 1, "hourly job made the next bill", run)
    for theme in ("dark", "light"):
        for lang in ("en", "hi"):
            tag = "%s-%s" % (theme, lang)
            ctx, p, errs = P("super@test.in", theme=theme, lang=lang, width=375 if lang == "hi" else 1280)
            p.goto(B + "/admin-bills.html"); p.wait_for_selector("#blf-WAITING")
            p.check("#blf-WAITING"); p.wait_for_selector(".bl-card.st-WAITING")
            ok(p.locator(".bl-card.st-WAITING [data-send]").count() == 1, "waiting bill with Send " + tag)
            if lang == "hi": ok(not re.findall(r"adm\.\w+", p.inner_text("main")), "Bills in Hindi, nothing missing " + tag)
            bad = p.evaluate(CONTRAST_JS); ok(not bad, "waiting bill contrast AA " + tag, bad[:5])
            shot(p, "68-bills-waiting-" + tag, full=False)
            ok(not errs, "no script errors (bills waiting %s)" % tag, errs)
            ctx.close()
    ok(not any(b.get("waiting") for b in api41("GET", "/api/billing", email="r2@x.in")[1]["bills"]) and len([b for b in api41("GET", "/api/billing", email="r2@x.in")[1]["bills"] if b.get("auto") and not b.get("sentAt")]) == 0, "office doesn't see it yet")
    ctx, p, errs = P("ops@test.in"); p.goto(B + "/admin-bills.html"); p.wait_for_selector("#blf-WAITING"); p.check("#blf-WAITING"); p.wait_for_selector("[data-send]")
    with p.expect_response(lambda r: "/api/admin/bills" in r.url and r.request.method == "POST"): p.click("[data-send]")
    p.wait_for_selector("#bl-msg.success")
    ok("Sent" in p.inner_text("#bl-msg") and p.evaluate("document.activeElement.id") == "bl-msg", "operations admin sends it; message focused")
    ok(not errs, "no script errors (bill send)", errs)
    ctx.close()
    ok(len([b for b in api41("GET", "/api/billing", email="r2@x.in")[1]["bills"] if b.get("auto")]) == 2, "office now sees the sent bill")
    # back to normal for anything after this
    api41("POST", "/api/admin/subscriptions", {"action": "resume", "type": "REP", "tier": "LOCAL", "id": "lu-hazratganj", "reason": "Break over"})
    api41("POST", "/api/admin/subscriptions", {"action": "resume", "type": "DEPT", "tier": "", "id": jal, "reason": "Testing over"})

    # ---------- grieviq-47: closed cases a page at a time; Speed page ----------
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            for width in (375, 1280):
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx = br.new_context(viewport={"width": width, "height": 900}, color_scheme=theme)
                ctx.add_cookies([{"name": "test_email", "value": "r1@x.in", "url": B}])
                ctx.add_init_script("try{localStorage.setItem('%s','%s');localStorage.setItem('giq-welcome-r1@x.in','1')}catch(e){}" % (LANGKEY, lang))
                p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
                p.goto(B + "/rep"); p.wait_for_selector("#more-closed", timeout=20000)
                ok(p.locator("#find-case").count() == 1, "find box shown when older cases aren't loaded " + tag)
                txt = p.inner_text(".more-closed")
                if lang == "en": ok("Showing 50 of" in txt and "closed cases" in txt, "show-more wording " + tag, txt)
                else: ok(not re.findall(r"rep\.\w+", p.inner_text("#content")) and "बंद मामले" in txt, "show-more in Hindi " + tag, txt)
                bad = p.evaluate(CONTRAST_JS.replace("createTreeWalker(document.body", "createTreeWalker(document.getElementById('content')", 1)); ok(not bad, "case list contrast AA " + tag, bad[:5])
                small = p.evaluate("() => Array.from(document.querySelectorAll('#find-case input, #find-case button, #more-closed')).filter(e => { const r = e.getBoundingClientRect(); return r.height < 24 || r.width < 24; }).map(e => e.id || e.textContent)")
                ok(not small, "find and show-more targets >= 24px " + tag, small)
                ok(no_hscroll(p), "case list: no sideways scroll " + tag)
                if width == 375:
                    p.locator("#find-case").scroll_into_view_if_needed(); shot(p, "69-find-" + tag, full=False)
                    p.locator(".more-closed").scroll_into_view_if_needed(); shot(p, "70-more-" + tag, full=False)
                ok(not errs, "no script errors (case paging %s)" % tag, errs)
                ctx.close()
    ctx = br.new_context(viewport={"width": 390, "height": 900})
    ctx.add_cookies([{"name": "test_email", "value": "r1@x.in", "url": B}])
    ctx.add_init_script("try{localStorage.setItem('giq-welcome-r1@x.in','1')}catch(e){}")
    p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
    p.goto(B + "/rep"); p.wait_for_selector("#more-closed", timeout=20000)
    n0 = p.locator(".case-row").count()
    with p.expect_response(lambda r: "/api/grievances?closed=100" in r.url): p.click("#more-closed")
    p.wait_for_timeout(300)
    ok(p.locator(".case-row").count() == n0 + 50, "Show 50 more adds 50 cases", (n0, p.locator(".case-row").count()))
    ok(p.evaluate("document.activeElement.className").startswith("case-row"), "focus moves to the first new case")
    p.fill("#find-ref", "abc"); p.click(".find-btn"); p.wait_for_timeout(150)
    ok(p.get_attribute("#find-ref", "aria-invalid") == "true" and "GRV-" in p.inner_text("#find-msg") and p.evaluate("document.activeElement.id") == "find-ref", "wrong format: explained, field marked and focused")
    p.fill("#find-ref", "GRV-NOPE77")
    with p.expect_response(lambda r: "ref=GRV-NOPE77" in r.url): p.click(".find-btn")
    p.wait_for_timeout(300)
    ok("GRV-NOPE77" in p.inner_text("#find-msg") and p.evaluate("document.activeElement.id") == "find-ref", "unknown number: told, focus back in the box")
    p.fill("#find-ref", "grv-old119")
    with p.expect_response(lambda r: "ref=GRV-OLD119" in r.url): p.click(".find-btn")
    p.wait_for_timeout(400)
    ok(p.evaluate("document.activeElement.id") == "case-gc119" and "GRV-OLD119" in p.inner_text("#case-gc119"), "old case found by number and focused")
    shot(p, "71-found-case", full=False)
    ok(not errs, "no script errors (find a case)", errs)
    ctx.close()
    # department dashboard: more than 50 closed in the last 30 days
    _d47 = _sq.connect(WORK + "/test.db", timeout=10)
    _jal = _d47.execute("SELECT office_id FROM dept_officers WHERE email = 'ae.zone3@nic.in' AND status = 'ACTIVE'").fetchone()[0]
    for i in range(40):
        _t = (_dt.datetime.now(_dt.timezone.utc) - _dt.timedelta(hours=i + 3)).strftime("%Y-%m-%dT%H:%M:%SZ")
        _d47.execute("INSERT OR IGNORE INTO grievances (id, tracking_ref, citizen_phone, local_unit_id, category_id, status, description, current_tier, created_at, resolved_at, closed_at) VALUES (?, ?, '9000000097', 'lu-hazratganj', 'water-sanitation', 'CLOSED', 'Recently closed', 'LOCAL', ?, ?, ?)", ("gd%d" % i, "GRV-DCL%03d" % i, _t, _t, _t))
        _d47.execute("INSERT OR IGNORE INTO case_dept_steps (id, grievance_id, kind, department, office_id, office_name, channel, actor, actor_role, created_at) VALUES (?, ?, 'FORWARDED', 'Water Supply / Sanitation', ?, 'Jal Kal', 'DASHBOARD', 'r1@x.in', 'OFFICE', ?)", ("sdu%d" % i, "gd%d" % i, _jal, _t))
    _d47.commit(); _d47.close()
    for theme in ("light", "dark"):
        ctx = br.new_context(viewport={"width": 375, "height": 900}, color_scheme=theme)
        ok(dept_signin(ctx) == 200, "dept signed in for paging " + theme)
        p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
        p.goto(B + "/dept"); p.wait_for_selector("#f-show"); p.select_option("#f-show", "CLOSED"); p.wait_for_selector("#more-closed")
        ok("last 30 days" in p.inner_text(".more-closed"), "department: show more closed " + theme)
        n0 = p.locator("[data-open]").count()
        with p.expect_response(lambda r: "/api/dept/cases?closed=100" in r.url): p.click("#more-closed")
        p.wait_for_timeout(300)
        ok(p.locator("[data-open]").count() > n0 and p.evaluate("document.activeElement.hasAttribute('data-open')"), "department: more cases, focus on the first new one " + theme)
        bad = p.evaluate(CONTRAST_JS); ok(not bad, "department list contrast AA " + theme, bad[:5])
        ok(not errs, "no script errors (dept paging %s)" % theme, errs)
        ctx.close()
    # Speed page. The slow figures are recorded again here, for the 2 hours
    # before now: the API tests' figures can fall outside "the last 2 hours"
    # when a test run crosses the top of an hour (grieviq-61).
    import time as _t61, urllib.request as _ur61
    _h0 = int(_t61.time() // 3600) * 3600
    _iso = lambda sec: _dt.datetime.fromtimestamp(sec, _dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.000Z")
    _samples = []
    for i in range(20):
        _samples += [{"area": "REP", "ms": 200 + i * 40, "at": _iso(_h0 - 7200 + (5 + i) * 60)}, {"area": "REP", "ms": 200 + i * 40, "at": _iso(_h0 - 3600 + (5 + i) * 60)}, {"area": "CITIZEN", "ms": 80 + i, "at": _iso(_h0 - 3600 + (10 + i) * 60)}]
    _ur61.urlopen(_ur61.Request(B + "/__test/speed-sample", data=json.dumps({"samples": _samples}).encode(), method="POST")).read()
    _ur61.urlopen(_ur61.Request(B + "/__test/run-speed", data=b"{}", method="POST")).read()
    for theme in ("dark", "light"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx, p, errs = P("super@test.in", theme=theme, lang=lang, width=width)
                p.goto(B + "/admin-speed.html"); p.wait_for_selector(".sp-status")
                ok(p.locator(".sp-card .sp-svg").count() == 4 and p.locator(".sp-meter").count() == 1, "Speed page: 4 parts with charts, database size " + tag)
                st = p.inner_text(".sp-status")
                if lang == "en": ok("Slow" in st and "Representative console" in st, "status in words names the slow part " + tag, st)
                else: ok(not re.findall(r"adm\.\w+", p.inner_text("main")) and "धीमा" in st, "Speed page in Hindi, nothing missing " + tag, re.findall(r"adm\.\w+", p.inner_text("main"))[:5])
                ok(no_hscroll(p), "Speed page: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS); ok(not bad, "Speed page contrast AA " + tag, bad[:5])
                small = p.evaluate(TARGET_JS); ok(not small, "Speed page targets " + tag, small[:5])
                shot(p, "72-speed-" + tag, full=True)
                ok(not errs, "no script errors (speed %s)" % tag, errs)
                ctx.close()
    ctx, p, errs = P("ops@test.in"); p.goto(B + "/admin-speed.html"); p.wait_for_selector(".sp-status")
    ok(True, "operations admin opens Speed"); ctx.close()
    ctx, p, errs = P("audit@test.in"); p.goto(B + "/admin-speed.html"); p.wait_for_timeout(800)
    ok("does not have access" in p.inner_text("main"), "auditor: no access to Speed"); ctx.close()
    ctx, p, errs = P("super@test.in"); p.goto(B + "/admin-dashboard.html"); p.wait_for_selector('a.card[href="/admin-speed.html"]')
    ok("Slow" in p.inner_text('a.card[href="/admin-speed.html"]'), "dashboard: Speed tile")
    ok("Ward and office data" in p.inner_text("header") and "Lucknow" not in p.inner_text("header"), "admin header: no city name")
    ok(not errs, "no script errors (dashboard speed tile)", errs)
    ctx.close()

    # ---------- grieviq-48: states and districts, themes, government access, sharing ----------
    for theme in ("dark", "light"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx, p, errs = P("super@test.in", theme=theme, lang=lang, width=width)
                p.goto(B + "/admin-areas.html"); p.wait_for_selector("#geo-state")
                ok(p.locator("#geo-state option").count() == 36 and re.search(r"\b7[5-9]\b", p.inner_text("details.geo-list summary")), "Areas: 36 states, districts of Uttar Pradesh listed " + tag, p.inner_text("details.geo-list summary"))
                p.click("details.geo-list summary"); p.fill("#geo-filter", "luck")
                ok(p.locator("details.geo-list tbody tr:not([hidden])").count() == 1 and "162" in p.inner_text("details.geo-list tbody tr:not([hidden])"), "district list filters as you type " + tag)
                if lang == "hi": ok(not re.findall(r"adm\.\w+", p.inner_text("main")), "Areas in Hindi, nothing missing " + tag, re.findall(r"adm\.\w+", p.inner_text("main"))[:5])
                bad = p.evaluate(CONTRAST_JS); ok(not bad, "Areas contrast AA " + tag, bad[:5])
                ok(no_hscroll(p), "Areas: no sideways scroll " + tag)
                if width == 375: p.locator("#geo").scroll_into_view_if_needed(); shot(p, "73-districts-" + tag, full=False)
                p.goto(B + "/admin-areas.html?id=lucknow"); p.wait_for_selector("#place-form")
                ok(p.input_value("#p-d") == "uttar-pradesh--lucknow" and p.input_value("#p-lgd") == "800123", "area shows its district and LGD code " + tag)
                bad = p.evaluate(CONTRAST_JS); ok(not bad, "area page contrast AA " + tag, bad[:5])
                p.goto(B + "/admin-gov.html"); p.wait_for_selector("#gv-form")
                ok(p.locator(".gv-card").count() >= 3, "Government access: people listed " + tag)
                if lang == "hi": ok(not re.findall(r"adm\.\w+", p.inner_text("main")), "Government access in Hindi, nothing missing " + tag, re.findall(r"adm\.\w+", p.inner_text("main"))[:5])
                bad = p.evaluate(CONTRAST_JS); ok(not bad, "Government access contrast AA " + tag, bad[:5])
                small = p.evaluate(TARGET_JS); ok(not small, "Government access targets " + tag, small[:5])
                ok(no_hscroll(p), "Government access: no sideways scroll " + tag)
                shot(p, "74-gov-" + tag, full=True)
                p.goto(B + "/admin-issue-types.html"); p.wait_for_selector(".theme-select")
                ok(p.locator(".theme-select").first.locator("option").count() == 20, "Issue types: 19 themes and 'not grouped' " + tag)
                bad = p.evaluate(CONTRAST_JS); ok(not bad, "Issue types contrast AA " + tag, bad[:5])
                ok(not errs, "no script errors (grieviq-48 admin %s)" % tag, errs)
                ctx.close()
    # give access: errors listed, linked and focused; then success
    ctx, p, errs = P("super@test.in", width=390)
    p.goto(B + "/admin-gov.html"); p.wait_for_selector("#gv-form")
    p.check("#gv-lv-d"); p.wait_for_timeout(100)
    ok(p.locator("#gv-districtId").count() == 1 and p.evaluate("document.activeElement.id") == "gv-lv-d", "choosing one district shows the district list, focus kept")
    p.fill("#gv-email", "someone-at-gmail"); p.click("#gv-add"); p.wait_for_selector("#gv-errs")
    ok(p.evaluate("document.activeElement.id") == "gv-errs" and "gov.in" in p.inner_text("#gv-errs") and p.get_attribute("#gv-email", "aria-invalid") == "true", "errors listed at the top and focused")
    p.click('#gv-errs a[href="#gv-email"]'); ok(p.evaluate("document.activeElement.id") == "gv-email", "error link moves to the field")
    p.fill("#gv-stateId >> xpath=preceding-sibling::div[1]//input", "uttar pr"); p.wait_for_timeout(200)
    ok(p.input_value("#gv-stateId") == "uttar-pradesh", "typing in the state search chooses the only match")
    p.fill("#gv-districtId >> xpath=preceding-sibling::div[1]//input", "kanpur")
    ok(p.locator("#gv-districtId option").count() == 3 and "2 found" in p.inner_text(".giq-find-n >> nth=1"), "district search narrows the list", p.locator("#gv-districtId option").count())
    p.select_option("#gv-districtId", "uttar-pradesh--kanpur-nagar")
    p.fill("#gv-name", "Rakesh Kumar"); p.fill("#gv-designation", "District Magistrate"); p.fill("#gv-email", "dm-kanpur@up.gov.in")
    p.fill("#gv-letterRef", "PG/2026/120"); p.fill("#gv-letterDate", "2026-10-12")
    with p.expect_response(lambda r: r.url.endswith("/api/admin/gov") and r.request.method == "POST"): p.click("#gv-add")
    p.wait_for_timeout(400)
    ok("Rakesh Kumar now has access" in p.inner_text("#gv-msg") and p.evaluate("document.activeElement.id") == "gv-msg", "access given: confirmed in words, focus on the message")
    ok("Kanpur Nagar district" in p.inner_text("#gv-app"), "district official shows their district")
    ok(not errs, "no script errors (give access)", errs)
    ctx.close()
    # representative chooses to share; office manager only sees it
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            tag = "%s-%s" % (theme, lang)
            ctx = br.new_context(viewport={"width": 375, "height": 900}, color_scheme=theme)
            ctx.add_cookies([{"name": "test_email", "value": "r1@x.in", "url": B}])
            ctx.add_init_script("try{localStorage.setItem('%s','%s');localStorage.setItem('giq-welcome-r1@x.in','1')}catch(e){}" % (LANGKEY, lang))
            p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
            p.goto(B + "/rep"); p.click('.tab-btn[data-tab="team"]'); p.wait_for_selector("#gov-share-t")
            if lang == "hi": ok(not re.findall(r"rep\.\w+", p.inner_text("#content")), "sharing box in Hindi " + tag)
            p.click("#gs-start"); ok(p.evaluate("document.activeElement.id") == "gs-yes", "asks first; focus on Yes " + tag)
            bad = p.evaluate(CONTRAST_JS); ok(not bad, "sharing box contrast AA " + tag, bad[:5])
            p.locator(".gov-share").scroll_into_view_if_needed(); shot(p, "75-share-ask-" + tag, full=False)
            with p.expect_response(lambda r: r.url.endswith("/api/team") and r.request.method == "POST"): p.click("#gs-yes")
            p.wait_for_selector("#gs-stop")
            ok(p.evaluate("document.activeElement.id") == "team-notice" and p.locator(".gov-share.on").count() == 1, "shared: told, box marked " + tag)
            bad = p.evaluate(CONTRAST_JS); ok(not bad, "shared box contrast AA " + tag, bad[:5])
            with p.expect_response(lambda r: r.url.endswith("/api/team") and r.request.method == "POST"): p.click("#gs-stop")
            p.wait_for_selector("#gs-start")
            ok(not errs, "no script errors (sharing %s)" % tag, errs)
            ctx.close()
    ctx = br.new_context(viewport={"width": 375, "height": 900})
    ctx.add_cookies([{"name": "test_email", "value": "om@x.in", "url": B}])
    ctx.add_init_script("try{localStorage.setItem('giq-welcome-om@x.in','1')}catch(e){}")
    p = ctx.new_page(); p.goto(B + "/rep"); p.click('.tab-btn[data-tab="team"]'); p.wait_for_selector("#gov-share-t")
    ok(p.locator("#gs-start").count() == 0 and "Only the representative" in p.inner_text(".gov-share"), "office manager sees the choice but can't change it")
    ctx.close()

    # ---------- grieviq-49: government dashboard ----------
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx = br.new_context(viewport={"width": width, "height": 900}, color_scheme=theme)
                ctx.add_init_script("try{localStorage.setItem('%s','%s')}catch(e){}" % (LANGKEY, lang))
                p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
                p.goto(B + ("/state" if width == 375 else "/national") + "?demo=1"); p.wait_for_selector(".tiles")
                ok(p.is_visible("#demo-bar") and (("Demo data" in p.inner_text("#demo-bar")) if lang == "en" else ("डेमो" in p.inner_text("#demo-bar"))), "demo clearly labelled " + tag)
                ok(p.locator(".tiles .tile").count() == 11, "11 key figure tiles " + tag)
                if lang == "hi": ok(not re.findall(r"(gv|theme|dd)\.\w+", p.inner_text("body")), "dashboard in Hindi, nothing missing " + tag, re.findall(r"(gv|theme|dd)\.\w+", p.inner_text("body"))[:5])
                bad = p.evaluate(CONTRAST_JS); ok(not bad, "dashboard contrast AA " + tag, bad[:5])
                small = p.evaluate(TARGET_JS); ok(not small, "dashboard targets " + tag, small[:5])
                ok(no_hscroll(p), "dashboard: no sideways scroll " + tag)
                shot(p, "76-gov-demo-" + tag, full=width == 375)
                ok(not errs, "no script errors (gov demo %s)" % tag, errs)
                ctx.close()
    ctx = br.new_context(viewport={"width": 1280, "height": 900})
    p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
    p.goto(B + "/national?demo=1"); p.wait_for_selector("a[data-state]")
    ok("Demo Country" in p.inner_text("h1") and p.locator("a[data-state]").count() == 6, "demo opens at the country: 6 demo states")
    p.click('th button[data-sort="received"]'); p.wait_for_timeout(100)
    ok(p.get_attribute('th:has(button[data-sort="received"])', "aria-sort") == "descending" and p.evaluate("document.activeElement.getAttribute('data-sort')") == "received", "sort by a column; focus stays on it")
    vals = p.eval_on_selector_all("section[aria-labelledby=rows-h] tbody tr td:nth-child(2)", "els => els.map(e => +e.textContent.replace(/,/g, ''))")
    ok(vals == sorted(vals, reverse=True), "rows sorted largest first", vals)
    p.fill("#rows-q", "state c"); p.wait_for_timeout(100)
    ok(p.locator("a[data-state]").count() == 1 and p.evaluate("document.activeElement.id") == "rows-q", "find a state as you type")
    p.click("a[data-state='demo-c']"); p.wait_for_selector("a[data-district]")
    ok(p.evaluate("document.activeElement.id") == "h1" and "Demo State C" in p.inner_text("h1") and "state=demo-c" in p.url and p.locator("nav.crumbs li").count() == 2, "state opened from the country; breadcrumb country › state")
    p.click("a[data-district='demo-c-3']"); p.wait_for_selector("nav.crumbs")
    ok("Demo District C3" in p.inner_text("h1") and "district=demo-c-3" in p.url and p.locator("nav.crumbs li").count() == 3, "district opened, address keeps it")
    p.click("a[data-dept]"); p.wait_for_selector("#dtrend-h")
    ok(p.locator("nav.crumbs li").count() == 4, "breadcrumb: country › state › district › department")
    p.go_back(); p.wait_for_timeout(500)
    ok("Demo District C3" in p.inner_text("h1") and "dept=" not in p.url, "back button returns to the district")
    p.click("nav.crumbs a[data-go='state']"); p.wait_for_timeout(500)
    ok("Demo State C" in p.inner_text("h1"), "breadcrumb back to the state")
    p.click("nav.crumbs a[data-go='country']"); p.wait_for_timeout(500)
    ok("Demo Country" in p.inner_text("h1"), "breadcrumb back to the country")
    p.select_option("#month", index=1); p.wait_for_timeout(500)
    ok("month=" in p.url and p.evaluate("document.activeElement.id") == "month", "another month; focus kept")
    shot(p, "81-gov-country", full=False)
    ok(not errs, "no script errors (gov drill-down)", errs)
    ctx.close()
    # sign in as a real state official (added in the API tests)
    _d49 = _sq.connect(WORK + "/test.db", timeout=10)
    for theme in ("light", "dark"):
        ctx = br.new_context(viewport={"width": 390, "height": 900}, color_scheme=theme)
        p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
        p.goto(B + "/state"); p.wait_for_selector("#f-email")
        bad = p.evaluate(CONTRAST_JS); ok(not bad, "gov sign-in contrast AA " + theme, bad[:5])
        ok(p.locator("a[href='/state?demo=1']").count() == 1 and p.locator("a[href='/national']").count() == 1, "sign-in page links to its demo and to the other dashboard " + theme)
        _d49.execute("DELETE FROM gov_codes WHERE email = 'ps.pg@up.gov.in'"); _d49.commit()   # tests sign in more often than the 3-per-15-minutes limit
        p.fill("#email", "ps.pg@up.gov.in"); p.click("#send"); p.wait_for_selector("#code")
        import hashlib as _hl
        _d49.execute("UPDATE gov_codes SET code_hash = ? WHERE email = 'ps.pg@up.gov.in' AND used_at IS NULL", (_hl.sha256(b"gov:ps.pg@up.gov.in:515151").hexdigest(),)); _d49.commit()
        p.fill("#code", "515151"); p.click("#verify"); p.wait_for_selector(".tiles")
        ok("Uttar Pradesh" in p.inner_text("h1") and not p.is_visible("#demo-bar") and "Priya Verma" in p.inner_text("#who"), "official signed in: their state, no demo label " + theme)
        txt = p.inner_text("main")
        ok("fewer than 10" in txt and "r1@x.in" not in txt and "GRV-" not in txt, "small numbers as 'fewer than 10'; no people or case numbers " + theme)
        bad = p.evaluate(CONTRAST_JS); ok(not bad, "official dashboard contrast AA " + theme, bad[:5])
        ok(no_hscroll(p), "official dashboard: no sideways scroll " + theme)
        shot(p, "77-gov-official-" + theme, full=True)
        p.click("#sign-out"); p.wait_for_selector("#f-email")
        ok("signed out" in p.inner_text("#msg").lower(), "signed out " + theme)
        ok(not errs, "no script errors (gov sign-in %s)" % theme, errs)
        ctx.close()
    _d49.close()
    # admin: another email shows the exception box
    ctx, p, errs = P("super@test.in", width=390)
    p.goto(B + "/admin-gov.html"); p.wait_for_selector("#gv-form")
    ok(not p.is_visible("#gv-other"), "exception box hidden for an empty email")
    p.fill("#gv-email", "someone@up.gov.in"); ok(not p.is_visible("#gv-other"), "hidden for a gov.in email")
    p.fill("#gv-email", "someone@gmail.com"); ok(p.is_visible("#gv-other") and "90" in p.inner_text("#gv-other"), "shown for another email, with the 90-day review")
    p.locator("#gv-other").scroll_into_view_if_needed(); shot(p, "78-gov-other-email", full=False)
    p.check("#gv-lv-n"); p.wait_for_timeout(100)
    ok(p.locator("#gv-stateId").count() == 0 and p.is_visible("#gv-national-h") and p.evaluate("document.activeElement.id") == "gv-lv-n", "whole country: no state to choose, the rule explained, focus kept")
    p.check("#gv-lv-s"); p.wait_for_timeout(100)
    ok(p.locator("#gv-stateId").count() == 1, "back to one state: the state list returns")
    bad = p.evaluate(CONTRAST_JS); ok(not bad, "exception box contrast AA", bad[:5])
    ok(not errs, "no script errors (other email)", errs)
    ctx.close()

    # ---------- grieviq-50: MIS reports ----------
    import openpyxl as _oxl
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme, lang, width)
                ctx = br.new_context(viewport={"width": width, "height": 900}, color_scheme=theme, accept_downloads=True)
                ctx.add_init_script("try{localStorage.setItem('%s','%s')}catch(e){}" % (LANGKEY, lang))
                p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
                p.goto(B + "/state?demo=1"); p.wait_for_selector("#reports")
                ok(p.locator("#reports ul.reports li").count() == 6, "six reports listed " + tag)
                bad = p.evaluate(CONTRAST_JS.replace("createTreeWalker(document.body", "createTreeWalker(document.getElementById('reports')", 1)); ok(not bad, "reports list contrast AA " + tag, bad[:5])
                p.goto(B + "/gov-report?r=summary&demo=1"); p.wait_for_selector("#h1")
                ok("DEMO" in p.inner_text(".demo") or "डेमो" in p.inner_text(".demo"), "report says demo " + tag)
                if lang == "hi": ok(not re.findall(r"(gv|theme|dd)\.\w+", p.inner_text("body")), "report in Hindi, nothing missing " + tag, re.findall(r"(gv|theme|dd)\.\w+", p.inner_text("body"))[:5])
                bad = p.evaluate(CONTRAST_JS); ok(not bad, "report page contrast AA " + tag, bad[:5])
                small = p.evaluate(TARGET_JS); ok(not small, "report page targets " + tag, small[:5])
                ok(no_hscroll(p), "report page: no sideways scroll " + tag)
                with p.expect_download() as d: p.click("#xlsx")
                path = WORK + "/dl-" + tag + ".xlsx"; d.value.save_as(path); wb = _oxl.load_workbook(path)
                ok(len(wb.sheetnames) == 2 and wb.worksheets[0]["A8"].value is not None and isinstance(wb.worksheets[0]["B8"].value, (int, float)), "Excel opens; numbers are numbers " + tag, (wb.sheetnames, wb.worksheets[0]["A8"].value, wb.worksheets[0]["B8"].value))
                if width == 375: shot(p, "79-report-" + tag, full=True)
                ok(not errs, "no script errors (reports %s)" % tag, errs)
                ctx.close()
    ctx = br.new_context(viewport={"width": 1280, "height": 900}, accept_downloads=True)
    p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
    for r in ("districts", "departments", "pending30", "themes", "reps"):
        p.goto(B + "/gov-report?r=%s&demo=1" % r); p.wait_for_selector("#h1")
        ok(p.locator("table tbody tr").count() >= 1 and not p.is_disabled("#pdf"), "demo report opens: " + r)
    with p.expect_download() as d: p.goto(B + "/gov-report?r=themes&demo=1&x=xlsx")
    d.value.save_as(WORK + "/dl-themes.xlsx"); wb = _oxl.load_workbook(WORK + "/dl-themes.xlsx"); ws = wb.worksheets[0]
    hdr = [c.value for c in ws[7]]
    ok(len(hdr) == 14 and hdr[0] == "Theme" and hdr[-1] == "Total", "Excel button on the dashboard downloads straight away; 12 month columns", hdr)
    ok(not errs, "no script errors (report kinds)", errs)
    ctx.close()
    # signed-in official: email choices saved
    _d50 = _sq.connect(WORK + "/test.db", timeout=10)
    ctx = br.new_context(viewport={"width": 390, "height": 900})
    p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
    _d50.execute("DELETE FROM gov_codes WHERE email = 'ps.pg@up.gov.in'"); _d50.commit()
    p.goto(B + "/state"); p.wait_for_selector("#f-email"); p.fill("#email", "ps.pg@up.gov.in"); p.click("#send"); p.wait_for_selector("#code")
    import hashlib as _hl2
    _d50.execute("UPDATE gov_codes SET code_hash = ? WHERE email = 'ps.pg@up.gov.in' AND used_at IS NULL", (_hl2.sha256(b"gov:ps.pg@up.gov.in:616161").hexdigest(),)); _d50.commit()
    p.fill("#code", "616161"); p.click("#verify"); p.wait_for_selector("#reports")
    ok(p.locator("#mail-monthly").count() == 1, "email choices shown to a signed-in official")
    was = p.is_checked("#mail-weekly")
    with p.expect_response(lambda r: "/api/gov/prefs" in r.url and r.request.method == "POST"): p.click("#mail-weekly")
    p.wait_for_timeout(200)
    ok("Saved" in p.inner_text("#mail-msg"), "choice saved, told in words")
    ok(_d50.execute("SELECT weekly FROM gov_report_prefs p JOIN gov_users u ON u.id = p.user_id WHERE u.email = 'ps.pg@up.gov.in' AND u.status = 'ACTIVE'").fetchone()[0] == (0 if was else 1), "choice stored")
    p.click("#reports a >> nth=0"); p.wait_for_selector("#h1")
    ok("Uttar Pradesh" in p.inner_text("main") and p.locator(".demo").count() == 0, "real report: the official's state, no demo label")
    shot(p, "80-report-official", full=True)
    ok(not errs, "no script errors (official reports)", errs)
    ctx.close(); _d50.close()

    # ---------- grieviq-52: State Dashboard and National Dashboard ----------
    for theme in ("light", "dark"):
        for lang in ("en", "hi"):
            tag = theme + "-" + lang
            ctx = br.new_context(viewport={"width": 375, "height": 900}, color_scheme=theme)
            ctx.add_init_script("try{localStorage.setItem('%s','%s')}catch(e){}" % (LANGKEY, lang))
            p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
            p.goto(B + "/gov"); p.wait_for_selector("ul.doors")
            ok(p.locator("li.door").count() == 2 and p.locator("li.door a[href='/state?demo=1']").count() == 1 and p.locator("li.door a[href='/national?demo=1']").count() == 1, "grieviq.in/gov: two doors, each with its demo " + tag)
            if lang == "hi": ok(not re.findall(r"(gv|dd)\.\w+", p.inner_text("body")), "doors in Hindi " + tag)
            bad = p.evaluate(CONTRAST_JS); ok(not bad, "doors contrast AA " + tag, bad[:5])
            small = p.evaluate(TARGET_JS); ok(not small, "doors targets " + tag, small[:5])
            shot(p, "82-doors-" + tag, full=True)
            for door, want in (("state", "State Dashboard"), ("national", "National Dashboard")):
                p.goto(B + "/" + door + "?demo=1"); p.wait_for_selector(".tiles")
                band = p.inner_text(".mode-band")
                ok((want.lower() in band.lower()) if lang == "en" else ("डैशबोर्ड" in band), "band says which dashboard: " + door + " " + tag, band)
                bad = p.evaluate(CONTRAST_JS); ok(not bad, door + " demo contrast AA " + tag, bad[:5])
                if door == "state":
                    ok(("Demo State A" in p.inner_text("h1") or "डेमो राज्य A" in p.inner_text("h1")) and p.locator("nav.crumbs").count() == 0, "State Dashboard demo opens at a demo state, no country above it " + tag)
                else:
                    ok(p.locator("a[data-state]").count() == 6, "National Dashboard demo opens at the demo country " + tag)
                shot(p, "83-%s-demo-%s" % (door, tag), full=False)
            ok(not errs, "no script errors (doors %s)" % tag, errs)
            ctx.close()
    # state demo: drill to a district and back; reports link back to the State Dashboard
    ctx = br.new_context(viewport={"width": 1280, "height": 900})
    p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
    p.goto(B + "/state?demo=1"); p.wait_for_selector("a[data-district]")
    p.click("a[data-district]"); p.wait_for_selector("nav.crumbs")
    ok(p.locator("nav.crumbs li").count() == 2 and "state=demo-a" in p.url, "state demo: state › district, no country")
    p.click("#reports a >> nth=0"); p.wait_for_selector("#h1")
    ok("/state?" in p.get_attribute("#back", "href") and "State Dashboard" in p.inner_text(".head"), "report says State Dashboard and goes back to it")
    p.goto(B + "/national?demo=1"); p.wait_for_selector("#reports"); p.click("#reports a >> nth=0"); p.wait_for_selector("#h1")
    ok("/national" in p.get_attribute("#back", "href") and "National Dashboard" in p.inner_text(".head"), "national report says National Dashboard")
    ok(not errs, "no script errors (state demo drill)", errs)
    ctx.close()
    # signed-in officials are sent to their own dashboard
    _d52 = _sq.connect(WORK + "/test.db", timeout=10)
    import hashlib as _hl3
    def gov_signin(p, door, email, code):
        _d52.execute("DELETE FROM gov_codes WHERE email = ?", (email,)); _d52.commit()
        p.goto(B + door); p.wait_for_selector("#f-email"); p.fill("#email", email); p.click("#send"); p.wait_for_selector("#code")
        _d52.execute("UPDATE gov_codes SET code_hash = ? WHERE email = ? AND used_at IS NULL", (_hl3.sha256(("gov:%s:%s" % (email, code)).encode()).hexdigest(), email)); _d52.commit()
        p.fill("#code", code); p.click("#verify")
    ctx = br.new_context(viewport={"width": 1280, "height": 900})
    p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
    gov_signin(p, "/national", "ps.pg@up.gov.in", "727272"); p.wait_for_selector(".tiles")
    ok(p.url.split("?")[0].endswith("/state") and "state or district access" in p.inner_text("main") and "Uttar Pradesh" in p.inner_text("h1"), "state official signing in on /national is sent to the State Dashboard, told why", p.url)
    p.goto(B + "/gov"); p.wait_for_selector(".tiles")
    ok(p.url.split("?")[0].endswith("/state"), "grieviq.in/gov sends a signed-in state official to the State Dashboard")
    p.goto(B + "/national"); p.wait_for_selector(".tiles")
    ok(p.url.split("?")[0].endswith("/state") and p.locator("a[data-state]").count() == 0, "a state official never sees the national view")
    ctx.close()
    ctx = br.new_context(viewport={"width": 1280, "height": 900})
    p = ctx.new_page()
    gov_signin(p, "/state", "js.pg@darpg.gov.in", "828282"); p.wait_for_selector(".tiles")
    ok(p.url.split("?")[0].endswith("/national") and "national access" in p.inner_text("main") and p.locator("a[data-state]").count() >= 1, "national official signing in on /state is sent to the National Dashboard", p.url)
    shot(p, "84-national-official", full=False)
    ok(not errs, "no script errors (doors sign-in)", errs)
    ctx.close(); _d52.close()

    # ---------- grieviq-53: change of representative ----------
    import urllib.request, urllib.error
    def ho_api(method, body=None, email="super@test.in", path="/api/admin/handovers"):
        req = urllib.request.Request(B + path, method=method, data=(_json.dumps(body).encode() if body is not None else None), headers={"content-type": "application/json", "test-email": email})
        try:
            with urllib.request.urlopen(req) as r: return r.status, _json.loads(r.read())
        except urllib.error.HTTPError as e: return e.code, _json.loads(e.read() or b"{}")
    _d53 = _sq.connect(WORK + "/test.db", timeout=10)
    for lang in ("en", "hi"):
        for theme in ("dark", "light"):
            for width in (1280, 375):
                tag = "%s %s %d" % (lang, theme, width)
                ctx, p, errs = P("super@test.in", width=width, theme=theme, lang=lang)
                p.goto(B + "/admin-handovers.html"); p.wait_for_selector("#ho-form")
                ok(p.locator("#ho-seat option").count() > 3 and p.locator("#ho-apply").count() == 1 and p.locator("#ho-ask").count() == 1, "handover page: super admin can apply or ask " + tag)
                p.select_option("#ho-seat", "lu-tilak"); p.wait_for_selector(".ho-cur")
                ok("r2@x.in" in p.inner_text(".ho-cur"), "handover page: shows who holds the seat now " + tag)
                if lang == "hi": ok(not re.findall(r"adm\.\w+", p.inner_text("main")), "handover page in Hindi " + tag, re.findall(r"adm\.\w+", p.inner_text("main"))[:5])
                ok(no_hscroll(p), "handover page: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS); ok(not bad, "handover page contrast AA " + tag, bad[:5])
                small = p.evaluate(TARGET_JS); ok(not small, "handover page targets " + tag, small[:5])
                if width == 1280: shot(p, "85-handover-admin-%s-%s" % (lang, theme), full=True)
                ok(not errs, "no script errors (handover page %s)" % tag, errs)
                ctx.close()
    # an operator asks with nothing filled in: errors listed and linked; then a real request
    ctx, p, errs = P("deo@test.in", width=375)
    p.goto(B + "/admin-handovers.html"); p.wait_for_selector("#ho-form")
    ok(p.locator("#ho-apply").count() == 0, "operator: only 'Send for approval'")
    p.click("#ho-ask"); p.wait_for_selector("#ho-errs")
    ok(p.evaluate("document.activeElement.id") == "ho-errs" and p.locator("#ho-errs li").count() >= 4 and p.get_attribute("#ho-seat", "aria-invalid") == "true", "operator: empty request lists what to fix, focus on the list")
    p.click("#ho-errs a[href='#ho-seat']"); ok(p.evaluate("document.activeElement.id") == "ho-seat", "error link moves to the field")
    p.select_option("#ho-seat", "lu-tilak"); p.wait_for_selector(".ho-cur")
    p.fill("#ho-newName", "Rina New"); p.fill("#ho-newEmail", "r3@x.in"); p.select_option("#ho-reason", "ELECTION")
    p.fill("#ho-evidence", "SEC/UP/2027/123"); p.fill("#ho-effectiveDate", p.get_attribute("#ho-effectiveDate", "max"))
    p.click("#ho-ask"); p.wait_for_selector("#ho-msg.success")
    ok(p.locator(".ho-sec .sb-card").count() == 1 and p.locator("[data-act=withdraw]").count() == 1 and p.locator("[data-act=approve]").count() == 0, "operator: request waits, can withdraw, cannot approve")
    shot(p, "86-handover-operator-375", full=True)
    ok(not errs, "no script errors (handover operator)", errs)
    ctx.close()
    # a second admin approves; the change shows in the record
    ctx, p, errs = P("super@test.in")
    p.goto(B + "/admin-handovers.html"); p.wait_for_selector("[data-act=approve]")
    p.click("[data-act=reject]"); p.wait_for_selector(".row-msg.error")
    ok(p.evaluate("document.activeElement.getAttribute('aria-invalid')") == "true", "reject without a reason: asked for one, focus on the box")
    p.click("[data-act=approve]"); p.wait_for_selector("#ho-msg.success")
    ok(p.locator(".ho-sec .sb-card").count() == 0 and "Rina New" in p.inner_text(".ho-log") and p.locator(".ho-log .st-APPLIED").count() >= 1, "approved: shown in recent changes")
    ok(_d53.execute("SELECT rep_email FROM local_units WHERE id = 'lu-tilak'").fetchone()[0] == "r3@x.in", "approved: the seat has the new representative")
    ok(not errs, "no script errors (handover approve)", errs)
    ctx.close()
    # the new representative's first sign-in: the summary card
    hid = _d53.execute("SELECT id FROM seat_handovers WHERE new_email = 'r3@x.in' AND status = 'APPLIED' ORDER BY requested_at DESC").fetchone()[0]
    for lang in ("en", "hi"):
        for theme in ("light", "dark"):
            tag = "%s %s" % (lang, theme)
            ctx = br.new_context(viewport={"width": 390 if theme == "dark" else 1280, "height": 900})
            ctx.add_cookies([{"name": "test_email", "value": "r3@x.in", "url": B}])
            ctx.add_init_script("try{localStorage.setItem('griq-theme','%s');localStorage.setItem('%s','%s')}catch(e){}" % (theme, LANGKEY, lang))
            p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
            p.goto(B + "/rep.html"); p.wait_for_selector(".ho-card")
            txt = p.inner_text(".ho-card")
            ok(("Tilak Nagar" in txt) and p.locator(".ho-card .ho-stats li").count() == 4 and p.locator("[data-ho-act=accept_plan]").count() == 1, "new representative sees the handover summary " + tag, txt[:200])
            if lang == "hi": ok(not re.findall(r"rep\.\w+", txt), "summary in Hindi " + tag)
            bad = p.evaluate(CONTRAST_JS.replace("createTreeWalker(document.body", "createTreeWalker(document.querySelector('.ho-card')", 1)); ok(not bad, "summary contrast AA " + tag, bad[:5])
            small = p.evaluate(TARGET_JS); ok(not small, "summary targets " + tag, small[:5])
            ok(no_hscroll(p), "summary: no sideways scroll " + tag)
            p.evaluate("document.querySelector('.ho-card').scrollIntoView()")
            shot(p, "87-handover-rep-%s-%s" % (lang, theme), full=False)
            ok(not errs, "no script errors (handover summary %s)" % tag, errs)
            ctx.close()
    ctx = br.new_context(viewport={"width": 1280, "height": 900})
    ctx.add_cookies([{"name": "test_email", "value": "r3@x.in", "url": B}])
    p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
    p.goto(B + "/rep.html"); p.wait_for_selector(".ho-card")
    p.click("[data-ho-act=accept_plan]"); p.wait_for_selector(".ho-card .ho-ok")
    ok(p.evaluate("document.activeElement.className") == "ho-ok" and _d53.execute("SELECT plan_accepted_at FROM seat_handovers WHERE id = ?", (hid,)).fetchone()[0], "continue the plan: recorded, focus on the confirmation")
    p.click("[data-ho-act=seen]"); p.wait_for_timeout(400)
    ok(p.locator(".ho-card").count() == 0 and p.evaluate("document.activeElement.getAttribute('data-tab')") == "cases", "Done: card goes, focus on Cases")
    p.reload(); p.wait_for_selector(".tab-bar")
    ok(p.locator(".ho-card").count() == 0, "summary not shown again")
    ok(not errs, "no script errors (handover summary actions)", errs)
    ctx.close()
    # back as it was
    st, r = ho_api("POST", {"action": "propose", "seatType": "ward", "seatId": "lu-tilak", "kind": "NEW_REP", "newName": "R Two", "newEmail": "r2@x.in", "reason": "OTHER", "reasonNote": "test reset", "evidence": "test reset", "effectiveDate": _d53.execute("SELECT date('now', '+330 minutes')").fetchone()[0], "applyNow": True})
    ok(st == 200 and r.get("status") == "APPLIED", "handover: seat put back", r)
    _d53.execute("UPDATE local_units SET rep_name = NULL WHERE id = 'lu-tilak'"); _d53.execute("UPDATE seat_handovers SET welcome_seen_at = 'x' WHERE welcome_seen_at IS NULL"); _d53.commit(); _d53.close()

    # ---------- grieviq-54: transfers ----------
    _d54 = _sq.connect(WORK + "/test.db", timeout=10)
    ctrl_id = _d54.execute("SELECT id FROM dept_offices WHERE name_en = 'Lucknow Nagar Nigam Control Room'").fetchone()[0]
    for lang in ("en", "hi"):
        for theme in (None, "light"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme or "dark", lang, width)
                ctx, p, errs = P("super@test.in", width=width, theme=theme, lang=lang)
                p.goto(B + "/admin-departments.html?area=lucknow"); p.wait_for_selector(".office")
                p.locator("article:has-text('Jal Kal Vibhag Lucknow') [data-officers]").click(); p.wait_for_selector("[data-op-leave]")
                p.locator("[data-op-leave]").first.click(); p.wait_for_selector("#op-lv-form")
                p.check("#op-lv-TRANSFER"); p.wait_for_selector("#op-lv-to")
                ok(p.locator("#op-lv-to option[value='%s']" % ctrl_id).count() == 1 and p.input_value("#op-lv-email") == "ae.zone3@nic.in", "transfer form: offices to choose, email filled in " + tag)
                p.check("#op-lv-ADD"); p.wait_for_selector("#op-lv-sname")
                if lang == "hi": ok(not re.findall(r"adm\.\w+", p.inner_text(".op")), "transfer form in Hindi " + tag, re.findall(r"adm\.\w+", p.inner_text(".op"))[:5])
                ok(no_hscroll(p), "transfer form: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS.replace("createTreeWalker(document.body", "createTreeWalker(document.querySelector('.op')", 1)); ok(not bad, "transfer form contrast AA " + tag, bad[:5])
                small = p.evaluate(TARGET_JS); ok(not small, "transfer form targets " + tag, small[:5])
                if width == 1280 or lang == "en": p.locator("#op-lv-form").screenshot(path=os.path.join(SHOTS, "88-officer-transfer-%s.png" % tag))
                p.click("#op-lv-cancel"); ok(p.evaluate("document.activeElement.getAttribute('data-op-leave')") is not None, "cancel: focus back on the button " + tag)
                ok(not errs, "no script errors (transfer form %s)" % tag, errs)
                ctx.close()
    # a real transfer, with the successor added in the same step
    ctx, p, errs = P("super@test.in")
    p.goto(B + "/admin-departments.html?area=lucknow"); p.wait_for_selector(".office")
    p.locator("article:has-text('Jal Kal Vibhag Lucknow') [data-officers]").click(); p.wait_for_selector("[data-op-leave]")
    p.locator("[data-op-leave]").first.click(); p.wait_for_selector("#op-lv-form")
    p.check("#op-lv-TRANSFER"); p.wait_for_selector("#op-lv-to")
    with p.expect_response(lambda r: "dept-officers" in r.url and r.request.method == "POST"): p.click("#op-lv-go")
    p.wait_for_selector("#op-lv-to[aria-invalid=true]")
    ok(p.evaluate("document.activeElement.id") == "op-lv-to" and p.locator("#op-lv-order[aria-invalid=true]").count() == 1 and p.locator("#op-lv-succ-err").count() == 1, "transfer: new office, order and successor asked for; focus on the first")
    p.select_option("#op-lv-to", ctrl_id); p.fill("#op-lv-order", "LMC/TR/2026/41"); p.fill("#op-lv-desig", "Assistant Engineer, Control Room")
    p.check("#op-lv-ADD"); p.wait_for_selector("#op-lv-sname")
    ok(p.input_value("#op-lv-order") == "LMC/TR/2026/41" and p.input_value("#op-lv-to") == ctrl_id, "choices kept while the form changes")
    p.fill("#op-lv-sname", "Pooja Mishra"); p.fill("#op-lv-semail", "ae.zone5@nic.in")
    with p.expect_response(lambda r: "dept-officers" in r.url and r.request.method == "POST"): p.click("#op-lv-go")
    p.wait_for_selector("#op-msg")
    msg = p.inner_text("#op-msg")
    ok("transferred to Lucknow Nagar Nigam Control Room" in msg and "Pooja Mishra added as the successor" in msg and p.evaluate("document.activeElement.id") == "op-msg", "transfer done: said plainly, focus on the message", msg)
    ok("Transferred to Lucknow Nagar Nigam Control Room" in p.inner_text(".op") and "LMC/TR/2026/41" in p.inner_text(".op"), "the old office shows where they went")
    ok(p.locator("article:has(h3:text-is('Lucknow Nagar Nigam Control Room')) .pill.nobody").count() == 0, "the new office no longer says nobody to reply")
    p.locator("article:has(h3:text-is('Lucknow Nagar Nigam Control Room')) [data-officers]").click(); p.wait_for_selector("text=Transferred here from")
    ok("Assistant Engineer, Control Room" in p.inner_text(".op"), "the new office shows the officer and where they came from")
    shot(p, "89-officer-transferred", full=False)
    ok(not errs, "no script errors (officer transfer)", errs)
    ctx.close()
    # government access: transfer, and the last official for a state
    ps_id = _d54.execute("SELECT id FROM gov_users WHERE email = 'ps.pg@up.gov.in' AND status = 'ACTIVE'").fetchone()[0]
    dm_id = _d54.execute("SELECT id FROM gov_users WHERE email = 'dm@lucknow.nic.in' AND status = 'ACTIVE'").fetchone()[0]
    for lang in ("en", "hi"):
        for theme in (None, "light"):
            for width in (1280, 375):
                tag = "%s-%s-%d" % (theme or "dark", lang, width)
                ctx, p, errs = P("super@test.in", width=width, theme=theme, lang=lang)
                p.goto(B + "/admin-gov.html"); p.wait_for_selector("details[data-id='%s']" % ps_id, state="attached")
                det = "details[data-id='%s']" % ps_id
                p.click(det + " > summary"); p.wait_for_selector(det + " .gv-alone")
                p.click(det + " [data-act=transfer]"); p.wait_for_selector("#gt-form")
                ok(p.evaluate("document.activeElement.id") == "gt-h" and p.is_checked("#gt-lv-s") and p.input_value("#gt-email") == "ps.pg@up.gov.in", "official transfer form: opens filled in, focus on its heading " + tag)
                if lang == "hi": ok(not re.findall(r"adm\.\w+", p.inner_text(det)), "official transfer in Hindi " + tag, re.findall(r"adm\.\w+", p.inner_text(det))[:5])
                ok(no_hscroll(p), "official transfer: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS.replace("createTreeWalker(document.body", "createTreeWalker(document.querySelector(\"%s\")" % det, 1)); ok(not bad, "official transfer contrast AA " + tag, bad[:5])
                small = p.evaluate(TARGET_JS); ok(not small, "official transfer targets " + tag, small[:5])
                if width == 1280 or lang == "en": p.locator(det).locator("xpath=..").screenshot(path=os.path.join(SHOTS, "90-official-transfer-%s.png" % tag))
                p.click("#gt-cancel"); ok(p.evaluate("document.activeElement.getAttribute('data-act')") == "transfer", "cancel: focus back on Transfer " + tag)
                ok(not errs, "no script errors (official transfer %s)" % tag, errs)
                ctx.close()
    ctx, p, errs = P("super@test.in")
    p.goto(B + "/admin-gov.html"); p.wait_for_selector("details[data-id='%s']" % dm_id, state="attached")
    det = "details[data-id='%s']" % dm_id
    p.click(det + " > summary"); p.click(det + " [data-act=transfer]"); p.wait_for_selector("#gt-form")
    ok(p.locator(det + " .gv-alone").count() == 0, "a district official covered by the state official: no warning")
    p.click("#gt-go"); p.wait_for_selector("#gt-letterRef[aria-invalid=true]")
    ok(p.locator("#gt-lv-d[aria-invalid], #gt-form .ferr").count() >= 3 and p.evaluate("document.activeElement.id") in ("gt-lv-d", "gt-letterRef"), "official transfer: what to fix shown; focus on the first", p.evaluate("document.activeElement.id"))
    p.select_option("#gt-districtId", "uttar-pradesh--agra"); p.fill("#gt-designation", "District Magistrate, Agra")
    p.fill("#gt-letterRef", "UP/APP/2026/311"); p.fill("#gt-letterDate", p.get_attribute("#gt-letterDate", "max")); p.fill("#gt-reason", "Transferred as DM Agra")
    with p.expect_response(lambda r: "/api/admin/gov" in r.url and r.request.method == "POST"): p.click("#gt-go")
    p.wait_for_selector("#gv-msg.success")
    ok("access moved" in p.inner_text("#gv-msg") and "District Magistrate, Agra" in p.inner_text("#gv-app") and p.evaluate("document.activeElement.id") == "gv-msg", "official moved: said plainly, new access listed")
    shot(p, "91-official-transferred", full=False)
    # the state's only official: removing needs the tick
    det = "details[data-id='%s']" % ps_id
    p.click(det + " > summary"); p.fill(det + " input[type=text]", "Retired from service")
    p.click(det + " [data-act=remove]"); p.wait_for_selector(det + " .gv-alone .ferr")
    ok(p.evaluate("document.activeElement.id").endswith("-alone") and p.input_value(det + " input[type=text]") == "Retired from service", "last official for the state: asked to tick first, reason kept, focus on the tick")
    ok(_d54.execute("SELECT status FROM gov_users WHERE id = ?", (ps_id,)).fetchone()[0] == "ACTIVE", "not removed without the tick")
    ok(not errs, "no script errors (official transfer flow)", errs)
    ctx.close(); _d54.close()

    # ---------- grieviq-55: paid public notices ----------
    _d55 = _sq.connect(WORK + "/test.db", timeout=10)
    import datetime as _dt55
    today55 = (_dt55.datetime.utcnow() + _dt55.timedelta(hours=5, minutes=30)).strftime("%Y-%m-%d")
    # Prices & site details: the notice plans
    for lang in ("en", "hi"):
        for theme in ("dark", "light"):
            tag = "%s %s" % (lang, theme)
            ctx, p, errs = P("super@test.in", theme=theme, lang=lang)
            p.goto(B + "/admin-site-settings.html"); p.wait_for_selector("#tab-notices")
            p.click("#tab-notices"); p.wait_for_selector(".ncard")
            ok(p.locator(".ncard").count() == 3 and p.locator("#f-notices-enabled").is_checked(), "Prices: three notice plans, switched on " + tag)
            if lang == "en":
                p.fill("#f-notices-plans-0-priceLocal", "349")
                ok("₹412" in p.inner_text(".ncard >> nth=0"), "Prices: price with GST worked out as you type")
            if lang == "hi": ok(not re.findall(r"adm\.\w+", p.inner_text("main")), "Prices notices tab in Hindi", re.findall(r"adm\.\w+", p.inner_text("main"))[:5])
            bad = p.evaluate(CONTRAST_JS); ok(not bad, "Prices notices tab contrast AA " + tag, bad[:5])
            shot(p, "92-prices-notices-%s-%s" % (lang, theme), full=False)
            ok(not errs, "no script errors (prices notices %s)" % tag, errs)
            ctx.close()
    # representative console: ask for a notice
    for lang in ("en", "hi"):
        for theme in ("light", "dark"):
            for width in (1280, 390):
                tag = "%s %s %d" % (lang, theme, width)
                ctx = br.new_context(viewport={"width": width, "height": 900})
                ctx.add_cookies([{"name": "test_email", "value": "r2@x.in", "url": B}])
                ctx.add_init_script("try{localStorage.setItem('griq-theme','%s');localStorage.setItem('%s','%s')}catch(e){}" % (theme, LANGKEY, lang))
                p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e))); p.on("dialog", lambda d: d.accept())
                p.goto(B + "/rep.html"); p.wait_for_selector('.tab-btn[data-tab="notices"]')
                p.click('.tab-btn[data-tab="notices"]'); p.wait_for_selector("#gn-form")
                ok(p.locator('input[name="gn-category"]').count() == 7 and p.locator('input[name="gn-planId"]').count() == 3, "rep notices: the form with types and plans " + tag)
                p.check("#gn-category-CAMP"); p.fill("#gn-titleEn", "Ration card help desk"); p.fill("#gn-bodyEn", "Bring your Aadhaar card. Help with new and corrected ration cards.")
                p.fill("#gn-titleHi", "राशन कार्ड सहायता डेस्क"); p.fill("#gn-bodyHi", "अपना आधार कार्ड साथ लाएँ। नए और सुधार वाले राशन कार्ड में सहायता।"); p.fill("#gn-place", "Ward office, Tilak Nagar")
                ok("Paid notice" in p.inner_text("#gn-preview") or "सशुल्क सूचना" in p.inner_text("#gn-preview"), "rep notices: live preview with the paid label " + tag)
                if lang == "hi": ok(not re.findall(r"(ntc|pn)\.\w+", p.inner_text("#notices-box")), "rep notices in Hindi " + tag, re.findall(r"(ntc|pn)\.\w+", p.inner_text("#notices-box"))[:5])
                ok(no_hscroll(p), "rep notices: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS.replace("createTreeWalker(document.body", "createTreeWalker(document.getElementById('notices-box')", 1)); ok(not bad, "rep notices contrast AA " + tag, bad[:5])
                small = p.evaluate(TARGET_JS); ok(not small, "rep notices targets " + tag, small[:5])
                if width == 1280: p.locator("#notices-box").screenshot(path=os.path.join(SHOTS, "93-rep-notice-form-%s-%s.png" % (lang, theme)))
                ok(not errs, "no script errors (rep notices %s)" % tag, errs)
                ctx.close()
    ctx = br.new_context(viewport={"width": 1280, "height": 900})
    ctx.add_cookies([{"name": "test_email", "value": "r2@x.in", "url": B}])
    p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e))); p.on("dialog", lambda d: d.accept())
    p.goto(B + "/rep.html"); p.wait_for_selector('.tab-btn[data-tab="notices"]'); p.click('.tab-btn[data-tab="notices"]'); p.wait_for_selector("#gn-form")
    p.fill("#gn-bodyEn", "See www.example.com")
    with p.expect_response(lambda r: "/api/notices/office" in r.url and r.request.method == "POST"): p.click("#gn-send")
    p.wait_for_selector("#gn-errs")
    ok(p.evaluate("document.activeElement.id") == "gn-errs" and p.locator("#gn-errs li").count() >= 3 and p.get_attribute("#gn-bodyEn", "aria-invalid") == "true", "rep notices: what to fix listed, focus on the list")
    p.click("#gn-errs a[data-go=category]"); ok(p.evaluate("document.activeElement.name") == "gn-category", "error link moves to the choice")
    p.check("#gn-category-CAMP"); p.fill("#gn-titleEn", "Ration card help desk"); p.fill("#gn-bodyEn", "Bring your Aadhaar card. Help with new and corrected ration cards.")
    p.fill("#gn-place", "Ward office, Tilak Nagar"); p.check("#gn-planId-week")
    with p.expect_response(lambda r: "/api/notices/office" in r.url and r.request.method == "POST"): p.click("#gn-send")
    p.wait_for_selector("#gn-msg")
    ok("Sent for approval" in p.inner_text("#gn-msg") and p.evaluate("document.activeElement.id") == "gn-msg" and p.locator(".gn-item.gn-st-PENDING").count() == 1, "rep notices: sent; listed as waiting")
    ok(not errs, "no script errors (rep notice request)", errs)
    ctx.close()
    nid = _d55.execute("SELECT id, ref FROM notices WHERE office_id = 'lu-tilak' AND status = 'PENDING'").fetchone()
    # admin: check and approve
    for lang in ("en", "hi"):
        for theme in ("dark", "light"):
            for width in (1280, 375):
                tag = "%s %s %d" % (lang, theme, width)
                ctx, p, errs = P("super@test.in", width=width, theme=theme, lang=lang)
                p.goto(B + "/admin-notices.html"); p.wait_for_selector(".nt-card")
                ok(p.locator(".nt-card .pn").count() == 1 and p.locator(".nt-checks input").count() == 5, "admin notices: the notice as citizens see it, and the checklist " + tag)
                if lang == "hi": ok(not re.findall(r"(adm|pn)\.\w+", p.inner_text("main")), "admin notices in Hindi " + tag, re.findall(r"(adm|pn)\.\w+", p.inner_text("main"))[:5])
                ok(no_hscroll(p), "admin notices: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS); ok(not bad, "admin notices contrast AA " + tag, bad[:5])
                small = p.evaluate(TARGET_JS); ok(not small, "admin notices targets " + tag, small[:5])
                if width == 1280: shot(p, "94-admin-notices-%s-%s" % (lang, theme), full=True)
                ok(not errs, "no script errors (admin notices %s)" % tag, errs)
                ctx.close()
    ctx, p, errs = P("super@test.in")
    p.goto(B + "/admin-notices.html"); p.wait_for_selector(".nt-card")
    p.click("[data-act=approve]"); p.wait_for_selector(".nt-checks .ferr")
    ok(p.evaluate("document.activeElement.getAttribute('data-check')") == "PUBLIC_INFO", "approve: every check needed; focus on the first")
    for k in ("PUBLIC_INFO", "NO_PARTY", "NO_PRAISE", "NO_CITIZEN", "TRUTHFUL"): p.check("[data-check=%s]" % k)
    with p.expect_response(lambda r: "/api/admin/notices" in r.url and r.request.method == "POST"): p.click("[data-act=approve]")
    p.wait_for_selector("#nt-msg.success")
    ok("approved" in p.inner_text("#nt-msg") and "GIQ-PAY-" in p.inner_text("#nt-msg") and p.locator(".nt-card.st-APPROVED").count() == 1, "approve: bill sent, notice waits for payment")
    # election code form
    p.fill("#np-label", "UP urban local body elections 2027"); p.select_option("#np-state", "Uttar Pradesh")
    p.fill("#np-from", today55); p.fill("#np-to", today55)
    p.click("#np-go"); p.wait_for_selector("#nt-msg.success")
    ok(p.locator(".nt-pause.now").count() == 1, "election code recorded and shown as in force")
    p.fill(".nt-pause.now input[type=text]", "Entered to test the page")
    p.click(".nt-pause.now [data-act=pause_remove]"); p.wait_for_selector(".nt-pause.gone")
    ok(not errs, "no script errors (admin approve)", errs)
    ctx.close()
    # pay it (office records, admin confirms), and three more paid notices for the ward
    bill = _d55.execute("SELECT id, total FROM bills WHERE notice_id = ?", (nid[0],)).fetchone()
    import urllib.request as _ur55
    def post55(path, body, email):
        req = _ur55.Request(B + path, method="POST", data=_json.dumps(body).encode(), headers={"content-type": "application/json", "test-email": email})
        with _ur55.urlopen(req) as r: return _json.loads(r.read())
    post55("/api/billing", {"billId": bill[0], "method": "UPI", "utr": "UIN1234567890", "amount": bill[1], "paidOn": today55}, "r2@x.in")
    pay = _d55.execute("SELECT id FROM bill_payments WHERE bill_id = ?", (bill[0],)).fetchone()[0]
    post55("/api/admin/bills", {"action": "confirm", "paymentId": pay}, "super@test.in")
    for i, (tier, oid, name, label, cat, t_en, t_hi) in enumerate([
        ("MLA", "mla-c", "Lucknow Central", "MLA", "HEALTH", "Free eye camp at the community centre", "सामुदायिक केंद्र में निःशुल्क नेत्र शिविर"),
        ("MP", "mp-lko", "Lucknow", "MP", "SCHEME", "Ayushman card enrolment drive", "आयुष्मान कार्ड पंजीकरण अभियान"),
        ("LOCAL", "lu-tilak", "Tilak Nagar-Kundari Rakabganj", "LOCAL_URBAN", "SABHA", "Ward Sabha meeting on Sunday", "रविवार को वार्ड सभा की बैठक")]):
        _d55.execute("INSERT INTO notices (id, ref, office_type, office_tier, office_id, office_name, office_label, area_id, state, category, title_en, title_hi, body_en, body_hi, event_date, event_time, place, plan_id, plan_name, days, reach, amount, gst_rate, gst, total, settings_version, start_on, starts_on, status, requested_by, requested_at) VALUES (?, ?, 'REP', ?, ?, ?, ?, 'lucknow', 'Uttar Pradesh', ?, ?, ?, 'Everyone is welcome. Bring an identity card.', 'सभी का स्वागत है। पहचान पत्र साथ लाएँ।', ?, '10 am to 2 pm', 'Community centre, Tilak Nagar', 'week', 'One week', 7, 'WIDE', 999, 18, 180, 1179, 1, ?, ?, 'PAID', 'x@x.in', ?)",
                     ("ui-n%d" % i, "GIQ-NTC-80000%d" % i, tier, oid, name, label, cat, t_en, t_hi, today55, today55, today55, today55 + "T00:00:00Z"))
    _d55.commit()
    # the citizen's ward page
    for lang in ("en", "hi"):
        for theme in ("light", "dark"):
            for width in (390, 1280):
                tag = "%s %s %d" % (lang, theme, width)
                ctx = br.new_context(viewport={"width": width, "height": 900})
                ctx.add_init_script("try{localStorage.setItem('%s','%s');localStorage.setItem('griq-theme','%s')}catch(e){}" % (LANGKEY, lang, theme))
                p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
                p.route("**/maps.googleapis.com/**", lambda rt: rt.abort())
                p.goto(B + "/lucknow/lu-tilak"); p.wait_for_selector(".pn-box")
                ok(p.locator(".pn-box > .pn").count() == 3 and p.locator("#pn-rest").is_hidden() and p.locator(".pn-more").count() == 1, "ward page: three notices, then Show all " + tag)
                txt = p.inner_text(".pn-box")
                ok(("Paid notice" in txt and "Published and paid for by" in txt) if lang == "en" else ("सशुल्क सूचना" in txt and "प्रकाशक और भुगतानकर्ता" in txt), "ward page: every notice says it is paid, and by whom " + tag)
                ok("₹" not in txt and "@" not in txt, "ward page: no prices or emails " + tag)
                p.click(".pn-more")
                ok(p.locator(".pn-box .pn:visible").count() == 4 and p.get_attribute(".pn-more", "aria-expanded") == "true" and "pn-title" in (p.evaluate("document.activeElement.className") or ""), "ward page: Show all opens the rest, focus on the next notice " + tag)
                if lang == "hi": ok(not re.findall(r"pn\.\w+", txt), "ward page notices in Hindi " + tag)
                ok(no_hscroll(p), "ward page notices: no sideways scroll " + tag)
                bad = p.evaluate(CONTRAST_JS.replace("createTreeWalker(document.body", "createTreeWalker(document.getElementById('pn-slot')", 1)); ok(not bad, "ward page notices contrast AA " + tag, bad[:5])
                small = p.evaluate(TARGET_JS.replace("document.querySelectorAll(", "document.getElementById('pn-slot').querySelectorAll(", 1)) if "document.querySelectorAll(" in TARGET_JS else p.evaluate(TARGET_JS)
                ok(not small, "ward page notices targets " + tag, small[:5])
                p.evaluate("document.getElementById('pn-slot').scrollIntoView()")
                shot(p, "95-ward-notices-%s-%s-%d" % (lang, theme, width), full=False)
                ok(not errs, "no script errors (ward page notices %s)" % tag, errs)
                ctx.close()
    # the department dashboard: Public notices screen
    ctx = br.new_context(viewport={"width": 1280, "height": 900})
    st55 = dept_signin(ctx)
    p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
    p.goto(B + "/dept#notices"); p.wait_for_selector("#gn-form, .gn-msg")
    ok(st55 == 200 and p.locator('input[name="gn-category"]').count() == 7 and "Public notices" in p.inner_text("#gn-h"), "department dashboard: the notice form", st55)
    bad = p.evaluate(CONTRAST_JS); ok(not bad, "department notices contrast AA", bad[:5])
    p.click("#back"); p.wait_for_selector("#go-notices"); ok(p.evaluate("document.activeElement.id") == "go-notices", "back: focus on Open public notices")
    ok(not errs, "no script errors (department notices)", errs)
    ctx.close()
    # grieviq-56: the authority letter
    PDF56 = {"name": "order.pdf", "mimeType": "application/pdf", "buffer": b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF"}
    for lang, theme in (("en", "light"), ("en", "dark"), ("hi", "dark")):
        tag = "%s %s" % (lang, theme)
        ctx = br.new_context(viewport={"width": 1280, "height": 900})
        ctx.add_cookies([{"name": "test_email", "value": "om@x.in", "url": B}])
        ctx.add_init_script("try{localStorage.setItem('griq-theme','%s');localStorage.setItem('%s','%s')}catch(e){}" % (theme, LANGKEY, lang))
        p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e))); p.on("dialog", lambda d: d.accept())
        p.goto(B + "/rep.html"); p.wait_for_selector('.tab-btn[data-tab="notices"]'); p.click('.tab-btn[data-tab="notices"]'); p.wait_for_selector("#gn-form")
        ok("(optional)" in p.inner_text("#gn-letter-box legend") or "(वैकल्पिक)" in p.inner_text("#gn-letter-box legend"), "letter: optional for a help desk " + tag)
        p.check("#gn-category-SERVICE"); p.wait_for_selector("#gn-letter-box")
        ok("(required)" in p.inner_text("#gn-letter-box legend") or "(आवश्यक)" in p.inner_text("#gn-letter-box legend"), "letter: required for a service interruption " + tag)
        ok(p.evaluate("document.activeElement.id") == "gn-category-SERVICE", "letter: focus stays on the choice " + tag)
        p.fill("#gn-titleEn", "Water supply off on Sunday"); p.fill("#gn-bodyEn", "Water supply will be off from 10 am to 4 pm for pipeline repair.")
        p.set_input_files("#gn-letter", files=[PDF56]); p.wait_for_selector("#gn-letter-c")
        p.fill("#gn-letterNo", "JK/SD/2026/45"); p.fill("#gn-letterDate", today55); p.fill("#gn-letterBy", "R. K. Singh, Executive Engineer, Jal Kal")
        ok(("Issued on the authority of letter no. JK/SD/2026/45" in p.inner_text("#gn-preview")) or ("प्राधिकार पर जारी" in p.inner_text("#gn-preview")), "letter: the preview shows its number and date " + tag)
        if lang == "hi": ok(not re.findall(r"(ntc|pn)\.\w+", p.inner_text("#notices-box")), "letter section in Hindi", re.findall(r"(ntc|pn)\.\w+", p.inner_text("#notices-box"))[:5])
        bad = p.evaluate(CONTRAST_JS.replace("createTreeWalker(document.body", "createTreeWalker(document.getElementById('notices-box')", 1)); ok(not bad, "letter section contrast AA " + tag, bad[:5])
        small = p.evaluate(TARGET_JS); ok(not small, "letter section targets " + tag, small[:5])
        p.locator("#gn-letter-box").screenshot(path=os.path.join(SHOTS, "96-notice-letter-%s-%s.png" % (lang, theme)))
        ok(not errs, "no script errors (letter section %s)" % tag, errs)
        ctx.close()
    ctx = br.new_context(viewport={"width": 1280, "height": 900})
    ctx.add_cookies([{"name": "test_email", "value": "om@x.in", "url": B}])
    p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e))); p.on("dialog", lambda d: d.accept())
    p.goto(B + "/rep.html"); p.wait_for_selector('.tab-btn[data-tab="notices"]'); p.click('.tab-btn[data-tab="notices"]'); p.wait_for_selector("#gn-form")
    p.check("#gn-category-SERVICE"); p.wait_for_selector("#gn-letter")
    p.fill("#gn-titleEn", "Water supply off on Sunday"); p.fill("#gn-bodyEn", "Water supply will be off from 10 am to 4 pm for pipeline repair."); p.check("#gn-planId-week")
    with p.expect_response(lambda r: "/api/notices/office" in r.url and r.request.method == "POST"): p.click("#gn-send")
    p.wait_for_selector("#gn-errs")
    ok(p.locator("#gn-errs a[data-go=letter]").count() == 1 and p.get_attribute("#gn-letter", "aria-invalid") == "true", "letter: missing letter listed and marked")
    p.set_input_files("#gn-letter", files=[PDF56]); p.wait_for_selector("#gn-letter-c")
    p.fill("#gn-letterNo", "JK/SD/2026/45"); p.fill("#gn-letterDate", today55); p.fill("#gn-letterBy", "R. K. Singh, Executive Engineer, Jal Kal")
    with p.expect_response(lambda r: "/api/notices/office" in r.url and r.request.method == "POST"): p.click("#gn-send")
    p.wait_for_selector("#gn-msg")
    ok("Sent for approval" in p.inner_text("#gn-msg") and "JK/SD/2026/45" in p.inner_text(".gn-item.gn-st-PENDING"), "letter: sent with the request; the office sees its details")
    p.click(".gn-item.gn-st-PENDING summary >> text=Send a new letter"); p.wait_for_selector("form[data-rl]")
    ok(p.locator("form[data-rl] input[type=file]").count() == 1, "letter: a new letter can be sent while it waits")
    ok(not errs, "no script errors (letter request)", errs)
    ctx.close()
    lid = _d55.execute("SELECT id, ref FROM notices WHERE office_id = 'lu-hazratganj' AND status = 'PENDING'").fetchone()
    ctx, p, errs = P("super@test.in")
    p.goto(B + "/admin-notices.html"); p.wait_for_selector(".nt-letter")
    card = "article:has(.nt-letter)"
    ok(p.get_attribute(card + " .nt-letter a", "href") == "/api/admin/notice-letter?id=" + lid[0] and p.locator(card + " [data-check=LETTER]").count() == 1, "admin: the letter to open and its own check")
    for k in ("PUBLIC_INFO", "NO_PARTY", "NO_PRAISE", "NO_CITIZEN", "TRUTHFUL", "LETTER"): p.check(card + " [data-check=%s]" % k)
    p.click(card + " [data-act=approve]"); p.wait_for_timeout(200)
    ok(p.evaluate("document.activeElement.getAttribute('data-vh')") == lid[0], "admin: asked how the letter was verified, focus there")
    p.fill(card + " [data-vh]", "Called the Jal Kal EE office on its official number from the directory")
    with p.expect_response(lambda r: "/api/admin/notices" in r.url and r.request.method == "POST"): p.click(card + " [data-act=approve]")
    p.wait_for_selector("#nt-msg.success")
    ok("approved" in p.inner_text("#nt-msg"), "admin: approved after verifying the letter")
    shot(p, "97-admin-letter-approved", full=False)
    ok(not errs, "no script errors (admin letter)", errs)
    ctx.close()
    bill = _d55.execute("SELECT id, total FROM bills WHERE notice_id = ?", (lid[0],)).fetchone()
    post55("/api/billing", {"billId": bill[0], "method": "UPI", "utr": "UIL1234567890", "amount": bill[1], "paidOn": today55}, "om@x.in")
    post55("/api/admin/bills", {"action": "confirm", "paymentId": _d55.execute("SELECT id FROM bill_payments WHERE bill_id = ?", (bill[0],)).fetchone()[0]}, "super@test.in")
    ctx = br.new_context(viewport={"width": 390, "height": 900}); p = ctx.new_page()
    p.route("**/maps.googleapis.com/**", lambda rt: rt.abort())
    p.goto(B + "/lucknow/lu-hazratganj"); p.wait_for_selector(".pn-box")
    ok("Issued on the authority of letter no. JK/SD/2026/45" in p.inner_text(".pn-box") and "Singh" not in p.inner_text(".pn-box"), "ward page: the letter's number and date, not the letter")
    p.evaluate("document.getElementById('pn-slot').scrollIntoView()"); shot(p, "98-ward-notice-authority", full=False)
    ctx.close()
    _d55.execute("DELETE FROM notices WHERE id LIKE 'ui-n%'"); _d55.commit(); _d55.close()

    # ---------- grieviq-57: wards with more than one elected member ----------
    # The API tests leave Maulvi Ganj with 4 seats: A (ma@x.in), B (mb@x.in),
    # C vacant, D empty, and complaints open to every member.
    W57 = "lu-maulvi"
    for lang, theme in (("en", "light"), ("hi", "dark")):
        ctx = br.new_context(viewport={"width": 390, "height": 900}, color_scheme=theme)
        ctx.add_init_script("try{localStorage.setItem('%s','%s');localStorage.setItem('griq-theme','%s');localStorage.setItem('griq-place',JSON.stringify({id:'%s'}))}catch(e){}" % (LANGKEY, lang, theme, W57))
        p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
        p.route("**/maps.googleapis.com/**", lambda rt: rt.abort())
        p.goto(B + "/submit"); p.wait_for_selector("#seatPick:not([hidden])")
        radios = p.locator("#seatPick input[type=radio]")
        ok(radios.count() == 3 and p.is_checked("#seat-any"), "form: any member (chosen) + the 2 members who can take complaints " + lang, radios.count())
        ok(p.locator("#seatPick legend").count() == 1 and p.is_visible("#seatConsent"), "form: a named group and the consent line " + lang)
        ok(p.locator("#seatPick label[for=seat-any] .seat-sub").count() == 1 and p.get_attribute("#seat-any", "aria-describedby") == "seat-any-sub", "form: the hint is read with the choice " + lang)
        p.focus("#seat-any"); p.keyboard.press("ArrowDown"); p.wait_for_timeout(100)
        ok(p.is_checked("#seat-0"), "form: arrow keys move between members " + lang)
        ok(no_hscroll(p), "form: no sideways scroll at 390 px " + lang)
        bad = p.evaluate(CONTRAST_JS.replace("document.body", "document.getElementById('seatPick')", 1))
        ok(not bad, "form: seat choice contrast " + lang + " " + theme, bad[:4])
        p.goto(B + "/lucknow/" + W57); p.wait_for_selector(".loc-who")
        ok(p.locator(".loc-who li").count() == 6 and ("Seat C" in p.inner_text(".loc-who") or "सीट C" in p.inner_text(".loc-who")), "ward page: every seat listed " + lang, p.inner_text(".loc-who"))
        ok(not errs, "no script errors (citizen seats) " + lang, errs)
        ctx.close()
    # a member's console: take, then return with a reason
    for lang, theme in (("en", "light"), ("hi", "dark")):
        ctx = br.new_context(viewport={"width": 1280, "height": 900}, color_scheme=theme)
        ctx.add_cookies([{"name": "test_email", "value": "mb@x.in", "url": B}])
        ctx.add_init_script("try{localStorage.setItem('%s','%s');localStorage.setItem('griq-theme','%s');localStorage.setItem('giq-welcome-mb@x.in','1')}catch(e){}" % (LANGKEY, lang, theme))
        p = ctx.new_page(); errs = []; p.on("pageerror", lambda e: errs.append(str(e)))
        p.goto(B + "/rep"); p.wait_for_selector(".take-btn")
        row = p.locator(".case-row:has(.take-btn)").first
        ok(row.locator(".ack-btn, .resolve-btn:not(.take-btn)").count() == 0 and row.locator(".seat-note").count() == 1, "console: an open-to-all case shows only Take " + lang)
        bad = p.evaluate(CONTRAST_JS.replace("document.body", "document.querySelector('.seat-note')", 1))
        ok(not bad, "console: seat note contrast " + lang + " " + theme, bad[:4])
        if lang == "en":
            cid = row.get_attribute("id").replace("case-", "")
            with p.expect_response(lambda r: "/seat" in r.url): row.locator(".take-btn").click()
            p.wait_for_selector("#case-" + cid + " .seat-return-btn")
            ok("taken the case" in p.inner_text("#resolve-status") and p.evaluate("document.activeElement.id") == "case-" + cid, "console: taken, said, focus on the case")
            p.click("#case-" + cid + " .seat-return-btn"); p.wait_for_selector("#seat-reason-" + cid)
            ok(p.evaluate("document.activeElement.id") == "seat-reason-" + cid, "console: the reason box gets focus")
            p.click("#case-" + cid + " .seat-return-go"); p.wait_for_timeout(150)
            ok(p.get_attribute("#seat-reason-" + cid, "aria-invalid") == "true" and "10 letters" in p.inner_text("#seat-reason-err-" + cid) and p.evaluate("document.activeElement.id") == "seat-reason-" + cid, "console: a reason is needed, said next to the box")
            p.click("#case-" + cid + " .seat-return-cancel"); p.wait_for_timeout(100)
            ok(p.evaluate("document.activeElement.className").find("seat-return-btn") >= 0, "console: cancel puts focus back")
            p.click("#case-" + cid + " .seat-return-btn"); p.fill("#seat-reason-" + cid, "Outside our part of the ward, please take it")
            with p.expect_response(lambda r: "/seat" in r.url): p.click("#case-" + cid + " .seat-return-go")
            p.wait_for_selector("#case-" + cid + " .take-btn")
            ok("Returned to the ward" in p.inner_text("#resolve-status"), "console: returned to the ward")
            shot(p, "99-member-console", full=False)
        ok(not errs, "no script errors (member console) " + lang, errs)
        ctx.close()
    # grieviq-58: the member's Overview and guide
    ctx, p, errs = P("mb@x.in")
    p.add_init_script("try{localStorage.setItem('giq-welcome-mb@x.in','1')}catch(e){}")
    p.goto(B + "/rep"); p.wait_for_selector('.tab-btn[data-tab="overview"]')
    p.click('.tab-btn[data-tab="overview"]'); p.wait_for_selector(".ov-ward")
    ok(p.locator(".ov-ward h3").count() == 1 and "waiting for a member" in p.inner_text(".ov-ward"), "overview: the whole ward shown apart")
    p.click("#help-btn"); p.wait_for_selector(".gq-drawer")
    ok("gqd-seats" in sop_ids(p), "guide: members of multi-member wards get the take/return guide", sop_ids(p))
    ok(not errs, "no script errors (member overview)", errs)
    ctx.close()
    # the member who took a case is shown to the others in one line
    _d57 = _sq.connect(WORK + "/test.db", timeout=10)
    gid = _d57.execute("SELECT id FROM grievances WHERE local_unit_id = ? AND seat_owner IS NULL AND status = 'OPEN' LIMIT 1", (W57,)).fetchone()[0]
    _d57.close()
    _ur55.urlopen(_ur55.Request(B + "/api/grievances/" + gid + "/seat", data=json.dumps({"action": "take"}).encode(), method="POST", headers={"content-type": "application/json", "test-email": "mb@x.in"})).read()
    ctx, p, errs = P("ma@x.in")
    p.goto(B + "/rep"); p.wait_for_selector(".ward-taken")
    ok(p.locator(".ward-taken h3").count() == 1 and "taken by Seat B" in p.inner_text(".ward-taken"), "console: other members' cases, one line each")
    ok(not errs, "no script errors (taken list)", errs)
    ctx.close()
    # admin: members in a ward
    for theme in ("light", "dark"):
        ctx, p, errs = P("super@test.in", theme=theme if theme == "light" else None)
        p.goto(B + "/admin-jurisdiction.html"); p.wait_for_selector("[data-toggle-mp]")
        p.click("[data-toggle-mp]"); p.wait_for_selector("[data-toggle-mla]"); p.click("[data-toggle-mla]"); p.wait_for_selector(".node-row.seat")
        ok(p.locator(".node-row.seat").count() == 3, "admin: seats B-D listed under the ward " + theme)
        p.click('[data-seats="lu-maulvi"]'); p.wait_for_selector("#seats-modal.open")
        ok(p.evaluate("document.activeElement.id") == "seats-count" and p.input_value("#seats-count") == "4", "admin: members box opens on the number " + theme)
        bad = p.evaluate(CONTRAST_JS.replace("document.body", "document.querySelector('#seats-modal .modal')", 1))
        ok(not bad, "admin: members box contrast " + theme, bad[:4])
        p.select_option("#seats-count", "2"); p.click("#seats-save"); p.wait_for_timeout(300)
        ok("Seat C" in p.inner_text("#seats-error") and p.get_attribute("#seats-count", "aria-invalid") == "true", "admin: a seat in use can't be removed, said in the box " + theme, p.inner_text("#seats-error"))
        p.keyboard.press("Escape"); p.wait_for_timeout(100)
        ok(not p.is_visible("#seats-modal") and p.evaluate("document.activeElement.getAttribute('data-seats')") == "lu-maulvi", "admin: Escape closes, focus back on Members " + theme)
        ok(not errs, "no script errors (admin seats) " + theme, errs)
        ctx.close()

    br.close()

print("\n%d passed, %d failed" % (passed, failed))
sys.exit(1 if failed else 0)
